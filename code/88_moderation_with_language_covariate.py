"""
Year × Contact moderation regression with out-group language ability added
as a covariate.

Tests whether the year × contact interaction holds when out-group language
proficiency is partialed out. Contact and language are conceptually related
(both index cross-group engagement), so language is a natural covariate.

Model:
    composite = B0 + B1*year + B2*contact_c + B3*(year × contact_c)
                + B4*language_c + ε

Variables:
  contact: out-group contact composite (Q51/Q52 or K4X1/K4X2), inverted so
           higher = more frequent contact, mean-centered within group
  language: out-group language ability (Q71_2/K5_2 for Estonians; Q71_1/K5_1
           for Russians), recoded as (7 − raw) so higher = more proficient,
           then mean-centered within group
  year: binary 0/1 (0 = 2020, 1 = 2023)

Outputs:
  code/_moderation_with_language.tsv
  reports/Year_x_Contact_With_Language.docx
  viz/fig_moderation_with_language_forest.jpg

The forest plot pairs the B₃ interaction from script 84 (no language control)
with the B₃ from this script (language controlled), showing how the moderation
estimate changes when language is partialed out.
"""

import math
from pathlib import Path

import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm
import matplotlib.pyplot as plt
import matplotlib.lines as mlines
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent.parent

# ---------- Load -----------------------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)
df20, _ = pyreadstat.read_sav(
    str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"), encoding="latin1"
)
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)


def clean(df, items, reverse_items=None, scale_max=None, dk_code=9):
    sub = df[items].apply(pd.to_numeric, errors="coerce")
    sub = sub.where(sub != dk_code)
    if reverse_items:
        for it in reverse_items:
            sub[it] = (scale_max + 1) - sub[it]
    return sub.mean(axis=1, skipna=True)


def to_num(s, dk=9):
    s = pd.to_numeric(s, errors="coerce")
    return s.where(s != dk).astype(float)


CONTACT_2023 = {0: [f"Q52_{i}" for i in range(1, 7)], 1: [f"Q51_{i}" for i in range(1, 7)]}
CONTACT_2020 = {0: [f"K4X2_{i}" for i in range(1, 7)], 1: [f"K4X1_{i}" for i in range(1, 7)]}
# Out-group language ability
LANG_2023 = {0: "Q71_2", 1: "Q71_1"}   # Estonians' Russian / Russians' Estonian
LANG_2020 = {0: "K5_2",  1: "K5_1"}


def build_cell(spec, df, year, grp_code):
    if isinstance(spec[f"items_{year}"], dict):
        grp_key = "E" if grp_code == 0 else "R"
        items = spec[f"items_{year}"][grp_key]
        rev = (spec[f"rev_{year}"][grp_key] if spec.get(f"rev_{year}") else None)
    else:
        items = spec[f"items_{year}"]
        rev = spec.get(f"rev_{year}")
    sub_df = df[df["ethnicity_binary"] == grp_code]
    composite_vals = clean(sub_df, items, rev, spec["smax"])
    if spec["inv"]:
        composite_vals = (spec["smax"] + 1) - composite_vals
    contact_items = (CONTACT_2023[grp_code] if year == 2023 else CONTACT_2020[grp_code])
    contact_vals = 6 - clean(sub_df, contact_items)
    lang_col = LANG_2023[grp_code] if year == 2023 else LANG_2020[grp_code]
    lang_vals = 7 - to_num(sub_df[lang_col])
    return pd.DataFrame({
        "value": composite_vals.values,
        "contact": contact_vals.values,
        "language": lang_vals.values,
        "year_2023": 1 if year == 2023 else 0,
    })


SPECS = [
    {"name": "Superordinate Identity", "smax": 4, "inv": True,
     "items_2023": ["Q67_2","Q67_4","Q67_5"], "rev_2023": ["Q67_4"],
     "items_2020": ["K6X5_2","K6X5_3","K6X5_4"], "rev_2020": ["K6X5_3"]},
    {"name": "SD: Primary Out-group", "smax": 5, "inv": False,
     "items_2023": {"E": ["Q57_1","Q58_1","Q59_1"], "R": ["Q57_2","Q58_2","Q59_2"]},
     "items_2020": {"E": ["K4X7_1","K4X8_1","K4X9_1"], "R": ["K4X7_2","K4X8_2","K4X9_2"]}},
    {"name": "SD: General Out-group", "smax": 5, "inv": False,
     "items_2023": ["Q57_4","Q57_5","Q58_4","Q58_5","Q59_4","Q59_5"],
     "items_2020": ["K4X7_3","K4X8_3","K4X9_3"]},
    {"name": "Comparative Opportunity Assessment", "smax": 5, "inv": True,
     "items_2023": [f"Q44_{i}" for i in range(1,13)],
     "items_2020": [f"K3X1_{i}" for i in range(1,13)]},
    {"name": "Belief in Inevitable Conflict", "smax": 4, "inv": False,
     "items_2023": ["Q63_1","Q63_2","Q63_3","Q63_4"], "rev_2023": ["Q63_1","Q63_2"],
     "items_2020": ["K6X1_1","K6X1_2","K6X1_3","K6X1_4"], "rev_2020": ["K6X1_1","K6X1_2"]},
    {"name": "Minority Inclusion Support", "smax": 4, "inv": True,
     "items_2023": ["Q68_1","Q68_2","Q68_3"],
     "items_2020": ["K6X6_1","K6X6_2","K6X6_3"]},
]

# Reference: B3 from no-language model (script 84)
no_lang = pd.read_csv(ROOT / "code" / "_year_x_contact_moderation.tsv", sep="\t")

rows = []
for spec in SPECS:
    for grp_code, grp_name in [(0, "Estonian"), (1, "Russian")]:
        s20 = build_cell(spec, df20, 2020, grp_code)
        s23 = build_cell(spec, df23, 2023, grp_code)
        long = pd.concat([s20, s23], ignore_index=True).dropna()
        long["contact_c"]    = long["contact"]  - long["contact"].mean()
        long["language_c"]   = long["language"] - long["language"].mean()
        long["yearXcontact"] = long["year_2023"] * long["contact_c"]
        X = sm.add_constant(long[["year_2023", "contact_c", "yearXcontact", "language_c"]])
        m = sm.OLS(long["value"], X).fit(cov_type="HC3")

        # Comparison: B3 from no-language model
        nl = no_lang[(no_lang["variable"] == spec["name"]) &
                     (no_lang["group"] == grp_name)].iloc[0]

        rows.append({
            "variable": spec["name"], "group": grp_name,
            "N": int(m.nobs),
            "B1_year":   m.params["year_2023"], "SE_year": m.bse["year_2023"], "p_year": m.pvalues["year_2023"],
            "B2_contact": m.params["contact_c"], "SE_contact": m.bse["contact_c"], "p_contact": m.pvalues["contact_c"],
            "B3_interaction": m.params["yearXcontact"],
            "SE_B3": m.bse["yearXcontact"],
            "CI_B3_low":  m.conf_int().loc["yearXcontact"][0],
            "CI_B3_high": m.conf_int().loc["yearXcontact"][1],
            "p_B3": m.pvalues["yearXcontact"],
            "B4_language": m.params["language_c"],
            "SE_language": m.bse["language_c"],
            "p_language":  m.pvalues["language_c"],
            "R2": m.rsquared,
            # Reference values (no language control)
            "B3_nolang": nl["B3_interaction"], "SE_B3_nolang": nl["SE_interaction"],
            "CI_B3_nolang_low":  nl["CI_low_interaction"],
            "CI_B3_nolang_high": nl["CI_high_interaction"],
            "p_B3_nolang":       nl["p_interaction"],
            "delta_B3": m.params["yearXcontact"] - nl["B3_interaction"],
        })


out = pd.DataFrame(rows)
tsv = ROOT / "code" / "_moderation_with_language.tsv"
out.to_csv(tsv, sep="\t", index=False, float_format="%.4f")
print(f"Saved TSV: {tsv}\n")


def stars(p):
    return "***" if p < .001 else "**" if p < .01 else "*" if p < .05 else ("⁺" if p < .10 else "ns")


def fmt_p(p):
    if pd.isna(p): return "—"
    if p < .001: return "< .001"
    return f"{p:.3f}".lstrip("0")


# Console summary
print(f"{'Composite':<37}{'Group':<10}{'B3 no-lang':>12}{'B3 +lang':>12}{'Δ B3':>9}{'B4 lang':>11}{'p (B3)':>10}")
print("=" * 102)
for r in rows:
    print(f"{r['variable']:<37}{r['group']:<10}"
          f"{r['B3_nolang']:>+12.4f}{r['B3_interaction']:>+12.4f}"
          f"{r['delta_B3']:>+9.4f}{r['B4_language']:>+11.3f}"
          f"{fmt_p(r['p_B3']):>10} {stars(r['p_B3'])}")


# ============================================================================
#                               WORD DOC
# ============================================================================
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(2.0)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)


def H(text, size=14, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    if color: r.font.color.rgb = color


def body(text, italic=False, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic


H("Year × Contact Moderation with Out-group Language Ability as Covariate",
  size=14, color=RGBColor(0x1f, 0x29, 0x37))

body(
    "The Year × Contact moderation regression (script 84) is extended with "
    "out-group language ability as a covariate. Out-group language is "
    "conceptually close to out-group contact — both index cross-group "
    "engagement — and the Year × Language regression (script 74) showed it "
    "predicts SD: Primary Out-group within each cell. Adding language tests "
    "whether the Year × Contact interaction holds *net of* language proficiency.",
    italic=True
)

eq = doc.add_paragraph()
eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = eq.add_run("composite = B₀ + B₁ × Year + B₂ × Contact_c + B₃ × (Year × Contact_c) "
                "+ B₄ × Language_c + ε")
r.italic = True; r.font.size = Pt(11)

body(
    "Contact and language are mean-centered within each group. Out-group "
    "language is reverse-coded so higher = more proficient (1 = none, 6 = native). "
    "Out-group contact is reverse-coded so higher = more frequent contact "
    "(1 = no communication, 5 = almost every day). HC3 robust SEs.",
    italic=True, size=9,
)

doc.add_paragraph()

# Table 7
H("Table 7. Year × Contact moderation regressions with language covariate", size=12)
header = ["Variable", "Group", "N",
          "B₁ (year)", "B₂ (contact)", "B₃ (interaction)",
          "B₄ (language)", "p (B₃)", "Sig", "R²"]
table = doc.add_table(rows=1, cols=len(header))
table.style = "Light Grid Accent 1"
for i, h in enumerate(header):
    c = table.rows[0].cells[i]; c.text = ""
    rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(8.5)

for r in rows:
    row = table.add_row().cells
    cells = [
        r["variable"], r["group"], str(r["N"]),
        f"{r['B1_year']:+.3f} ({r['SE_year']:.3f}) {stars(r['p_year'])}",
        f"{r['B2_contact']:+.3f} ({r['SE_contact']:.3f}) {stars(r['p_contact'])}",
        f"{r['B3_interaction']:+.4f} ({r['SE_B3']:.4f}) {stars(r['p_B3'])}",
        f"{r['B4_language']:+.3f} ({r['SE_language']:.3f}) {stars(r['p_language'])}",
        fmt_p(r["p_B3"]),
        stars(r["p_B3"]),
        f"{r['R2']:.3f}",
    ]
    for i, val in enumerate(cells):
        c = row[i]; c.text = ""
        rr = c.paragraphs[0].add_run(val); rr.font.size = Pt(8.5)
        if i >= 3:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
body(
    "Note. *** p < .001, ** p < .01, * p < .05, ⁺ p < .10, ns = not significant.",
    italic=True, size=8,
)

# Comparison narrative
doc.add_paragraph()
H("Comparison: B₃ interaction with vs. without language control", size=12)

comparison_header = ["Variable", "Group",
                     "B₃ no-lang", "B₃ +lang", "Δ B₃", "p (B₃ no-lang)", "p (B₃ +lang)"]
table_c = doc.add_table(rows=1, cols=len(comparison_header))
table_c.style = "Light Grid Accent 1"
for i, h in enumerate(comparison_header):
    c = table_c.rows[0].cells[i]; c.text = ""
    rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)
for r in rows:
    row = table_c.add_row().cells
    cells = [
        r["variable"], r["group"],
        f"{r['B3_nolang']:+.4f} {stars(r['p_B3_nolang'])}",
        f"{r['B3_interaction']:+.4f} {stars(r['p_B3'])}",
        f"{r['delta_B3']:+.4f}",
        fmt_p(r["p_B3_nolang"]),
        fmt_p(r["p_B3"]),
    ]
    for i, val in enumerate(cells):
        c = row[i]; c.text = ""
        rr = c.paragraphs[0].add_run(val); rr.font.size = Pt(9)
        if i >= 2:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
body(
    "Δ B₃ = (B₃ with language) − (B₃ without language). A small Δ means the "
    "Year × Contact moderation is robust to controlling for language; a large "
    "Δ means language partly accounted for the apparent moderation.",
    italic=True, size=9,
)

out_doc = ROOT / "reports" / "Year_x_Contact_With_Language.docx"
doc.save(out_doc)
print(f"\nSaved Word doc: {out_doc}")


# ============================================================================
#                               FOREST PLOT
# ============================================================================
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 9,
    "axes.edgecolor": "#333",
    "axes.linewidth": 0.7,
})

RUS = "#d97706"
EST = "#2563eb"

COMPOSITES = [s["name"] for s in SPECS]
PRETTY = {
    "Superordinate Identity":             "Superordinate Identity",
    "SD: Primary Out-group":              "SD: Primary Out-group",
    "SD: General Out-group":              "SD: General Out-group ⁱ",
    "Comparative Opportunity Assessment": "Comparative Opportunity",
    "Belief in Inevitable Conflict":      "Belief in Inevitable Conflict",
    "Minority Inclusion Support":         "Minority Inclusion Support",
}

fig, ax = plt.subplots(figsize=(15.0, 9.0), dpi=300)

y_positions = {}
y = 13.0
for comp in COMPOSITES:
    y_positions[(comp, "Estonian")] = y; y -= 1.0
    y_positions[(comp, "Russian")]  = y; y -= 1.7

y_min, y_max = min(y_positions.values()) - 0.6, max(y_positions.values()) + 0.7
for i, comp in enumerate(COMPOSITES):
    if i % 2 == 0:
        y_top = y_positions[(comp, "Estonian")] + 0.6
        y_bot = y_positions[(comp, "Russian")]  - 0.6
        ax.axhspan(y_bot, y_top, facecolor="#f7f7f7", zorder=0)

for r in rows:
    y_center = y_positions[(r["variable"], r["group"])]
    color = EST if r["group"] == "Estonian" else RUS
    y_no = y_center + 0.20
    y_l  = y_center - 0.20

    # No-language B3: open dot, faded
    ax.plot([r["CI_B3_nolang_low"], r["CI_B3_nolang_high"]], [y_no, y_no],
            color=color, linewidth=1.6, alpha=0.55,
            solid_capstyle="round", zorder=2)
    for xx in (r["CI_B3_nolang_low"], r["CI_B3_nolang_high"]):
        ax.plot([xx, xx], [y_no - 0.08, y_no + 0.08],
                color=color, linewidth=1.2, alpha=0.6, zorder=2)
    ax.scatter([r["B3_nolang"]], [y_no], s=140, facecolor="white",
               edgecolor=color, linewidth=2.0, zorder=3)

    # Connector
    ax.plot([r["B3_nolang"], r["B3_interaction"]], [y_no, y_l],
            color=color, linewidth=0.8, alpha=0.45,
            linestyle=(0, (2, 2)), zorder=2)

    # +Language B3: filled dot, solid
    ax.plot([r["CI_B3_low"], r["CI_B3_high"]], [y_l, y_l],
            color=color, linewidth=2.2, alpha=0.9,
            solid_capstyle="round", zorder=2)
    for xx in (r["CI_B3_low"], r["CI_B3_high"]):
        ax.plot([xx, xx], [y_l - 0.08, y_l + 0.08],
                color=color, linewidth=1.6, alpha=0.95, zorder=2)
    ax.scatter([r["B3_interaction"]], [y_l], s=160, color=color,
               edgecolor="white", linewidth=1.0, zorder=3)

    ann = (f"{r['group']:<8}  B₃ no-lang = {r['B3_nolang']:+.4f} {stars(r['p_B3_nolang'])}    "
           f"B₃ +lang = {r['B3_interaction']:+.4f} {stars(r['p_B3'])}    "
           f"Δ = {r['delta_B3']:+.4f}    "
           f"B₄ lang = {r['B4_language']:+.3f} {stars(r['p_language'])}")
    ax.text(max(r["CI_B3_nolang_high"], r["CI_B3_high"]) + 0.012, y_center,
            ann, ha="left", va="center", fontsize=8.0,
            color="#222", family="monospace")

ax.axvline(0, color="#666", linewidth=1.0, linestyle="-", zorder=1)

# Y-axis labels at composite midpoints
yticks, yticklabels = [], []
for comp in COMPOSITES:
    y_E = y_positions[(comp, "Estonian")]
    y_R = y_positions[(comp, "Russian")]
    yticks.append((y_E + y_R) / 2)
    yticklabels.append(PRETTY[comp])
ax.set_yticks(yticks)
ax.set_yticklabels(yticklabels, fontsize=10, fontweight="bold")
ax.set_ylim(y_min, y_max)
ax.invert_yaxis()

all_lo = min(min(r["CI_B3_nolang_low"], r["CI_B3_low"]) for r in rows)
all_hi = max(max(r["CI_B3_nolang_high"], r["CI_B3_high"]) for r in rows)
ax.set_xlim(all_lo - 0.02, all_hi + 1.10)
ax.set_xlabel("B₃ — Year × Contact interaction coefficient",
              fontsize=10, color="#444", labelpad=8)

for s in ("top", "right"):
    ax.spines[s].set_visible(False)
ax.spines["left"].set_color("#ccc")
ax.spines["bottom"].set_color("#666")
ax.tick_params(axis="x", length=2.5, color="#666", labelsize=9)
ax.tick_params(axis="y", length=0)

fig.text(0.04, 0.965,
         "Year × Contact Interaction — With vs. Without Out-group Language Covariate",
         fontsize=14, fontweight="bold", ha="left")
fig.text(0.04, 0.940,
         "Open dot + faded bracket = B₃ without language control (script 84). Filled dot + solid bracket = B₃ with out-group language ability as covariate.",
         fontsize=9, color="#444", ha="left")
fig.text(0.04, 0.918,
         "Dashed connector shows how much language absorbs from the moderation. When B₃ changes little, the contact moderation is independent of language.",
         fontsize=9, color="#444", ha="left")

legend_handles = [
    mlines.Line2D([], [], marker="o", color=EST, markerfacecolor="white",
                  markeredgecolor=EST, markersize=10, markeredgewidth=1.8,
                  linestyle="None", label="Estonian — no language"),
    mlines.Line2D([], [], marker="o", color=EST, markerfacecolor=EST,
                  markeredgecolor="white", markersize=10,
                  linestyle="None", label="Estonian — + language"),
    mlines.Line2D([], [], marker="o", color=RUS, markerfacecolor="white",
                  markeredgecolor=RUS, markersize=10, markeredgewidth=1.8,
                  linestyle="None", label="Russian — no language"),
    mlines.Line2D([], [], marker="o", color=RUS, markerfacecolor=RUS,
                  markeredgecolor="white", markersize=10,
                  linestyle="None", label="Russian — + language"),
]
fig.legend(handles=legend_handles, loc="upper right",
           bbox_to_anchor=(0.985, 0.94),
           frameon=False, fontsize=8.5, ncol=2,
           handlelength=1.4, columnspacing=1.5)

fig.text(0.04, 0.022,
         "Out-group language: Estonian respondents → Russian ability (Q71_2 / K5_2); Russian respondents → Estonian ability (Q71_1 / K5_1). Reverse-coded: higher = more proficient.",
         fontsize=7.5, color="#555", ha="left")
fig.text(0.04, 0.008,
         "ⁱ SD: General Out-group uses different items in 2020 (3) vs 2023 (6). HC3 robust SEs. Mean-centering within group. Significance: *** p<.001, ** p<.01, * p<.05, ⁺ p<.10.",
         fontsize=7.5, color="#555", ha="left")

plt.subplots_adjust(left=0.18, right=0.985, top=0.88, bottom=0.075)

out = ROOT / "viz" / "fig_moderation_with_language_forest.jpg"
plt.savefig(out, dpi=300, format="jpg", facecolor="white",
            pil_kwargs={"quality": 95})
print(f"Saved chart: {out}")
