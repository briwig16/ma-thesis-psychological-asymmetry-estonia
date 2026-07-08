"""
Compute Cohen's d with 95% confidence intervals for every between-group and
within-group comparison across the 11 variables (10 main + Contact: Out-group),
and emit a Word document with two formatted tables.

Cohen's d uses the root-mean-square SD denominator (CLAUDE.md):
    d = (M1 - M2) / sqrt((SD1^2 + SD2^2) / 2)

95% CI on d uses the asymptotic SE formula (Hedges & Olkin 1985):
    SE(d) = sqrt((n1 + n2)/(n1 * n2) + d^2 / (2 * (n1 + n2 - 2)))
    CI    = d ± 1.96 * SE(d)

All comparisons are independent-samples Welch's t-tests (the EIM is
cross-sectional at each wave — different respondents in 2020 vs. 2023).

Outputs:
    reports/Effect_Sizes_with_CI.docx
    code/_effect_sizes.tsv  (machine-readable copy)
"""

import math
from pathlib import Path

import pandas as pd
import pyreadstat
from scipy import stats
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

ROOT = Path(__file__).parent.parent

# ---------- Load both waves -----------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy()

df20, _ = pyreadstat.read_sav(str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"),
                              encoding="latin1")
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy()


# ---------- Helpers --------------------------------------------------------
def composite(df, items, reverse_items=None, scale_max=None, dk_code=9):
    """Return a Series of composite means using pairwise deletion.
    Coerces string-typed columns (e.g., Q57_1 in EIM23.csv with ' ' = missing)
    to numeric, with non-numeric values becoming NaN."""
    sub = df[items].apply(pd.to_numeric, errors="coerce").copy()
    sub = sub.where(sub != dk_code)        # 9 → NaN
    if reverse_items:
        for item in reverse_items:
            sub[item] = (scale_max + 1) - sub[item]
    return sub.mean(axis=1, skipna=True)

def single_item(df, var, dk_code=9):
    s = pd.to_numeric(df[var], errors="coerce")
    return s.where(s != dk_code).astype(float)

def invert_series(s, scale_max):
    return (scale_max + 1) - s.astype(float)

def m_sd_n(series):
    s = series.dropna().astype(float)
    return s.mean(), s.std(ddof=1), len(s)

def cohens_d_rms(m1, sd1, n1, m2, sd2, n2):
    """Cohen's d using root-mean-square SD denominator + asymptotic 95% CI."""
    s_rms = math.sqrt((sd1 ** 2 + sd2 ** 2) / 2)
    d = (m1 - m2) / s_rms
    se = math.sqrt((n1 + n2) / (n1 * n2) + d ** 2 / (2 * (n1 + n2 - 2)))
    return d, se, d - 1.96 * se, d + 1.96 * se

def welch_p(a, b):
    """Two-sided Welch's t-test p-value."""
    a = pd.Series(a).dropna().astype(float)
    b = pd.Series(b).dropna().astype(float)
    _, p = stats.ttest_ind(a, b, equal_var=False)
    return float(p)

def fmt_p(p):
    if p < 0.0001: return "<.0001"
    if p < 0.001:  return f"{p:.4f}"
    return f"{p:.3f}"

def stars(p):
    if p < 0.001: return "***"
    if p < 0.01:  return "**"
    if p < 0.05:  return "*"
    return "ns"


# ---------- Build per-variable group-year series --------------------------
# Each entry maps to a Series indexed by the corresponding df's index.
def build():
    out = {}

    # 1. Superordinate Identity (3-item)
    e23 = composite(df23[df23["ethnicity_binary"] == 0],
                    ["Q67_2", "Q67_4", "Q67_5"], ["Q67_4"], scale_max=4)
    r23 = composite(df23[df23["ethnicity_binary"] == 1],
                    ["Q67_2", "Q67_4", "Q67_5"], ["Q67_4"], scale_max=4)
    e20 = composite(df20[df20["ethnicity_binary"] == 0],
                    ["K6X5_2", "K6X5_3", "K6X5_4"], ["K6X5_3"], scale_max=4)
    r20 = composite(df20[df20["ethnicity_binary"] == 1],
                    ["K6X5_2", "K6X5_3", "K6X5_4"], ["K6X5_3"], scale_max=4)
    out["Superordinate Identity"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 4, "inverted": True}

    # 2. SD: Primary Out-group (group-specific items)
    e23 = composite(df23[df23["ethnicity_binary"] == 0],
                    ["Q57_1", "Q58_1", "Q59_1"])
    r23 = composite(df23[df23["ethnicity_binary"] == 1],
                    ["Q57_2", "Q58_2", "Q59_2"])
    e20 = composite(df20[df20["ethnicity_binary"] == 0],
                    ["K4X7_1", "K4X8_1", "K4X9_1"])
    r20 = composite(df20[df20["ethnicity_binary"] == 1],
                    ["K4X7_2", "K4X8_2", "K4X9_2"])
    out["SD: Primary Out-group"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 5, "inverted": False}

    # 3. SD: General Out-group  — DIFFERENT items in 2020 (3) vs 2023 (6)
    e23 = composite(df23[df23["ethnicity_binary"] == 0],
                    ["Q57_4", "Q57_5", "Q58_4", "Q58_5", "Q59_4", "Q59_5"])
    r23 = composite(df23[df23["ethnicity_binary"] == 1],
                    ["Q57_4", "Q57_5", "Q58_4", "Q58_5", "Q59_4", "Q59_5"])
    e20 = composite(df20[df20["ethnicity_binary"] == 0],
                    ["K4X7_3", "K4X8_3", "K4X9_3"])
    r20 = composite(df20[df20["ethnicity_binary"] == 1],
                    ["K4X7_3", "K4X8_3", "K4X9_3"])
    out["SD: General Out-group"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 5, "inverted": False,
        "note": "Item set differs across waves (3 items in 2020, 6 in 2023)"}

    # 4. Comparative Opportunity Assessment (Q44_1..12 / K3X1_1..12)
    items23 = [f"Q44_{i}" for i in range(1, 13)]
    items20 = [f"K3X1_{i}" for i in range(1, 13)]
    e23 = composite(df23[df23["ethnicity_binary"] == 0], items23)
    r23 = composite(df23[df23["ethnicity_binary"] == 1], items23)
    e20 = composite(df20[df20["ethnicity_binary"] == 0], items20)
    r20 = composite(df20[df20["ethnicity_binary"] == 1], items20)
    out["Comparative Opportunity Assessment"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 5, "inverted": True}

    # 5. Belief in Inevitable Conflict (Q63 / K6X1)
    # 2026-04-30: switched reverse-coding from Q63_3/Q63_4 to Q63_1/Q63_2 to fix
    # the documentation/direction mismatch. With Q63_1/Q63_2 reversed, the
    # composite has direction "higher = more conflict belief" — matching the
    # convention documented in CLAUDE.md.
    # Q63_1 and Q63_2 are conflict-positive ("conflicts inevitable" / "differences
    # divide society"), so raw 1 = strongly agree = high conflict belief. We
    # invert these so high inverted = high conflict belief. Q63_3 and Q63_4 are
    # conflict-negative ("groups can cooperate" / "immigration enriches"), so
    # raw 4 = strongly disagree = high conflict belief — already aligned, no
    # inversion needed.
    e23 = composite(df23[df23["ethnicity_binary"] == 0],
                    ["Q63_1", "Q63_2", "Q63_3", "Q63_4"], ["Q63_1", "Q63_2"], scale_max=4)
    r23 = composite(df23[df23["ethnicity_binary"] == 1],
                    ["Q63_1", "Q63_2", "Q63_3", "Q63_4"], ["Q63_1", "Q63_2"], scale_max=4)
    e20 = composite(df20[df20["ethnicity_binary"] == 0],
                    ["K6X1_1", "K6X1_2", "K6X1_3", "K6X1_4"],
                    ["K6X1_1", "K6X1_2"], scale_max=4)
    r20 = composite(df20[df20["ethnicity_binary"] == 1],
                    ["K6X1_1", "K6X1_2", "K6X1_3", "K6X1_4"],
                    ["K6X1_1", "K6X1_2"], scale_max=4)
    out["Belief in Inevitable Conflict"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 4, "inverted": False}

    # 6. Minority Support Inclusion
    e23 = composite(df23[df23["ethnicity_binary"] == 0],
                    ["Q68_1", "Q68_2", "Q68_3"])
    r23 = composite(df23[df23["ethnicity_binary"] == 1],
                    ["Q68_1", "Q68_2", "Q68_3"])
    e20 = composite(df20[df20["ethnicity_binary"] == 0],
                    ["K6X6_1", "K6X6_2", "K6X6_3"])
    r20 = composite(df20[df20["ethnicity_binary"] == 1],
                    ["K6X6_1", "K6X6_2", "K6X6_3"])
    out["Minority Support Inclusion"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 4, "inverted": True}

    # 7. Contact: Estonian Speakers
    e23 = composite(df23[df23["ethnicity_binary"] == 0],
                    [f"Q51_{i}" for i in range(1, 7)])
    r23 = composite(df23[df23["ethnicity_binary"] == 1],
                    [f"Q51_{i}" for i in range(1, 7)])
    e20 = composite(df20[df20["ethnicity_binary"] == 0],
                    [f"K4X1_{i}" for i in range(1, 7)])
    r20 = composite(df20[df20["ethnicity_binary"] == 1],
                    [f"K4X1_{i}" for i in range(1, 7)])
    out["Contact: Estonian Speakers"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 5, "inverted": True}

    # 8. Contact: Russian Speakers
    e23 = composite(df23[df23["ethnicity_binary"] == 0],
                    [f"Q52_{i}" for i in range(1, 7)])
    r23 = composite(df23[df23["ethnicity_binary"] == 1],
                    [f"Q52_{i}" for i in range(1, 7)])
    e20 = composite(df20[df20["ethnicity_binary"] == 0],
                    [f"K4X2_{i}" for i in range(1, 7)])
    r20 = composite(df20[df20["ethnicity_binary"] == 1],
                    [f"K4X2_{i}" for i in range(1, 7)])
    out["Contact: Russian Speakers"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 5, "inverted": True}

    # 9. Group ID Patterns (single item; 2020 uses code 6 for DK)
    e23 = single_item(df23[df23["ethnicity_binary"] == 0], "Q66", dk_code=9)
    r23 = single_item(df23[df23["ethnicity_binary"] == 1], "Q66", dk_code=9)
    e20 = single_item(df20[df20["ethnicity_binary"] == 0], "K6X4", dk_code=6)
    r20 = single_item(df20[df20["ethnicity_binary"] == 1], "K6X4", dk_code=6)
    out["Group ID Patterns"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 5, "inverted": False,
        "single_item": True}

    # 10. Territorial Attachment (single item)
    e23 = single_item(df23[df23["ethnicity_binary"] == 0], "Q67_1")
    r23 = single_item(df23[df23["ethnicity_binary"] == 1], "Q67_1")
    e20 = single_item(df20[df20["ethnicity_binary"] == 0], "K6X5_1")
    r20 = single_item(df20[df20["ethnicity_binary"] == 1], "K6X5_1")
    out["Territorial Attachment"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 4, "inverted": True,
        "single_item": True}

    # 11. Contact: Out-group (each group's contact with the other group)
    # Estonians → Q52/K4X2 (contact with Russian speakers)
    # Russians  → Q51/K4X1 (contact with Estonian speakers)
    e23 = composite(df23[df23["ethnicity_binary"] == 0],
                    [f"Q52_{i}" for i in range(1, 7)])
    r23 = composite(df23[df23["ethnicity_binary"] == 1],
                    [f"Q51_{i}" for i in range(1, 7)])
    e20 = composite(df20[df20["ethnicity_binary"] == 0],
                    [f"K4X2_{i}" for i in range(1, 7)])
    r20 = composite(df20[df20["ethnicity_binary"] == 1],
                    [f"K4X1_{i}" for i in range(1, 7)])
    out["Contact: Out-group"] = {
        "raw": (e20, r20, e23, r23), "scale_max": 5, "inverted": True,
        "note": "Estonians' Q52/K4X2; Russians' Q51/K4X1"}

    return out


vars_data = build()


# ---------- Compute the two tables -----------------------------------------
between_rows = []   # (variable, year, est_M, est_SD, est_N, rus_M, rus_SD, rus_N, d, ll, ul, p)
within_rows  = []   # (variable, group, M_2020, SD_2020, N_2020, M_2023, SD_2023, N_2023, d, ll, ul, p)

for varname, info in vars_data.items():
    e20_raw, r20_raw, e23_raw, r23_raw = info["raw"]
    smax = info["scale_max"]
    inv  = info["inverted"]

    # Apply visualization-direction inversion BEFORE computing d, so positive d
    # always means "more of the construct" (matching documentation tables).
    e20 = invert_series(e20_raw, smax) if inv else e20_raw.astype(float)
    r20 = invert_series(r20_raw, smax) if inv else r20_raw.astype(float)
    e23 = invert_series(e23_raw, smax) if inv else e23_raw.astype(float)
    r23 = invert_series(r23_raw, smax) if inv else r23_raw.astype(float)

    me20, sde20, ne20 = m_sd_n(e20)
    mr20, sdr20, nr20 = m_sd_n(r20)
    me23, sde23, ne23 = m_sd_n(e23)
    mr23, sdr23, nr23 = m_sd_n(r23)

    # Between-group (Estonian − Russian) for each year
    for year, (mE, sE, nE, mR, sR, nR, eS, rS) in [
        ("2020", (me20, sde20, ne20, mr20, sdr20, nr20, e20, r20)),
        ("2023", (me23, sde23, ne23, mr23, sdr23, nr23, e23, r23)),
    ]:
        d, se, ll, ul = cohens_d_rms(mE, sE, nE, mR, sR, nR)
        p = welch_p(eS, rS)
        between_rows.append((varname, year, mE, sE, nE, mR, sR, nR, d, ll, ul, p))

    # Within-group (2023 − 2020) for each group
    for group, (m20, s20, n20, m23, s23, n23, s20S, s23S) in [
        ("Estonian", (me20, sde20, ne20, me23, sde23, ne23, e20, e23)),
        ("Russian",  (mr20, sdr20, nr20, mr23, sdr23, nr23, r20, r23)),
    ]:
        d, se, ll, ul = cohens_d_rms(m23, s23, n23, m20, s20, n20)  # 2023 − 2020
        p = welch_p(s23S, s20S)
        within_rows.append((varname, group, m20, s20, n20, m23, s23, n23, d, ll, ul, p))


# ---------- Sanity-print a few sentinel values -----------------------------
print("Sample between-group rows:")
for r in between_rows[:4] + between_rows[-2:]:
    name, yr, mE, sE, nE, mR, sR, nR, d, ll, ul, p = r
    print(f"  {name:<35} {yr}  Est M={mE:.2f} SD={sE:.2f} N={nE}  "
          f"Rus M={mR:.2f} SD={sR:.2f} N={nR}  d={d:+.2f} [{ll:+.2f},{ul:+.2f}]  p={fmt_p(p)}")

print("\nSample within-group rows:")
for r in within_rows[:4] + within_rows[-2:]:
    name, gr, m20, s20, n20, m23, s23, n23, d, ll, ul, p = r
    print(f"  {name:<35} {gr:<8}  2020 M={m20:.2f} SD={s20:.2f} N={n20}  "
          f"2023 M={m23:.2f} SD={s23:.2f} N={n23}  d={d:+.2f} [{ll:+.2f},{ul:+.2f}]  p={fmt_p(p)}")


# ---------- Build the Word document ---------------------------------------
doc = Document()

# Page setup — letter, narrower margins for table room
section = doc.sections[0]
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)

# Default style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)

# Title
h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = h.add_run("Effect Sizes with 95% Confidence Intervals")
run.bold = True
run.font.size = Pt(16)

p = doc.add_paragraph()
run = p.add_run(
    "Cohen's d (root-mean-square SD denominator) for between-group "
    "(Estonian vs. Russian) and within-group (2020 → 2023) comparisons "
    "across all 11 variables. Variables marked (inv.) have been scale-inverted "
    "so that positive d indicates more of the construct."
)
run.font.size = Pt(9)
run.italic = True

doc.add_paragraph()


# ----- Table 1: Between-Group ---------------------------------------------
heading = doc.add_paragraph()
run = heading.add_run("Table 1. Between-Group Comparisons (Estonian vs. Russian)")
run.bold = True
run.font.size = Pt(12)

cols = ["Variable", "Year",
        "Estonian M (SD), N", "Russian M (SD), N",
        "Cohen's d", "95% CI", "p", "Sig"]
table = doc.add_table(rows=1, cols=len(cols))
table.style = "Light Grid Accent 1"
table.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Header
for i, col in enumerate(cols):
    cell = table.rows[0].cells[i]
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(col)
    run.bold = True
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Body
for r in between_rows:
    name, year, mE, sE, nE, mR, sR, nR, d, ll, ul, p_val = r
    is_inv = vars_data[name]["inverted"]
    label  = name + (" (inv.)" if is_inv else "")
    cells_text = [
        label,
        year,
        f"{mE:.2f} ({sE:.2f}), n = {nE}",
        f"{mR:.2f} ({sR:.2f}), n = {nR}",
        f"{d:+.2f}",
        f"[{ll:+.2f}, {ul:+.2f}]",
        fmt_p(p_val),
        stars(p_val),
    ]
    row_cells = table.add_row().cells
    for i, txt in enumerate(cells_text):
        c = row_cells[i]
        c.text = ""
        para = c.paragraphs[0]
        run = para.add_run(txt)
        run.font.size = Pt(8.5)
        if i in (4, 5, 6, 7):
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_paragraph()


# ----- Table 2: Within-Group ----------------------------------------------
heading = doc.add_paragraph()
run = heading.add_run("Table 2. Within-Group Change, 2020 → 2023")
run.bold = True
run.font.size = Pt(12)

cols2 = ["Variable", "Group",
         "2020 M (SD), N", "2023 M (SD), N",
         "Cohen's d", "95% CI", "p", "Sig"]
table2 = doc.add_table(rows=1, cols=len(cols2))
table2.style = "Light Grid Accent 1"

for i, col in enumerate(cols2):
    cell = table2.rows[0].cells[i]
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(col)
    run.bold = True
    run.font.size = Pt(9)

for r in within_rows:
    name, group, m20, s20, n20, m23, s23, n23, d, ll, ul, p_val = r
    is_inv = vars_data[name]["inverted"]
    label  = name + (" (inv.)" if is_inv else "")
    cells_text = [
        label,
        group,
        f"{m20:.2f} ({s20:.2f}), n = {n20}",
        f"{m23:.2f} ({s23:.2f}), n = {n23}",
        f"{d:+.2f}",
        f"[{ll:+.2f}, {ul:+.2f}]",
        fmt_p(p_val),
        stars(p_val),
    ]
    row_cells = table2.add_row().cells
    for i, txt in enumerate(cells_text):
        c = row_cells[i]
        c.text = ""
        para = c.paragraphs[0]
        run = para.add_run(txt)
        run.font.size = Pt(8.5)
        if i in (4, 5, 6, 7):
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_paragraph()

# Footer note
notes = doc.add_paragraph()
run = notes.add_run(
    "Notes. Cohen's d uses the root-mean-square SD denominator: "
    "d = (M₁ − M₂) / √[(SD₁² + SD₂²) / 2]. "
    "95% CIs use the asymptotic standard error: "
    "SE(d) = √[(n₁ + n₂)/(n₁n₂) + d² / (2(n₁ + n₂ − 2))]. "
    "All comparisons are independent-samples Welch's t-tests; the EIM is "
    "cross-sectional at each wave (different respondents in 2020 and 2023). "
    "Significance: *** p < .001, ** p < .01, * p < .05, ns = not significant. "
    "Variables marked (inv.) were scale-inverted (raw scale: lower = more of "
    "the construct) so that positive d in this table corresponds to more of "
    "the construct (e.g., more belonging, more contact). "
    "SD: General Out-group uses different items in 2020 (3 items, new immigrants) "
    "vs. 2023 (6 items, other Europeans + non-Europeans); within-group change "
    "estimates for that variable should be interpreted with caution."
)
run.font.size = Pt(8)
run.italic = True

out_path = ROOT / "reports" / "Effect_Sizes_with_CI.docx"
doc.save(out_path)
print(f"\nSaved: {out_path}")

# Also write a TSV copy for machine-readable use
tsv_path = ROOT / "code" / "_effect_sizes.tsv"
with open(tsv_path, "w") as f:
    f.write("table\tvariable\tcomparison\tM1\tSD1\tN1\tM2\tSD2\tN2\td\tCI_low\tCI_high\tp\n")
    for r in between_rows:
        name, year, mE, sE, nE, mR, sR, nR, d, ll, ul, p = r
        f.write(f"between\t{name}\t{year}\t{mE:.4f}\t{sE:.4f}\t{nE}\t{mR:.4f}\t{sR:.4f}\t{nR}\t{d:.4f}\t{ll:.4f}\t{ul:.4f}\t{p:.6f}\n")
    for r in within_rows:
        name, group, m20, s20, n20, m23, s23, n23, d, ll, ul, p = r
        f.write(f"within\t{name}\t{group}\t{m20:.4f}\t{s20:.4f}\t{n20}\t{m23:.4f}\t{s23:.4f}\t{n23}\t{d:.4f}\t{ll:.4f}\t{ul:.4f}\t{p:.6f}\n")
print(f"Saved: {tsv_path}")
