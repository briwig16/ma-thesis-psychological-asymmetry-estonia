"""
Build a single comprehensive PCA / reliability validation table covering all
8 multi-item composites × 2 years × 2 groups (32 rows). Recomputes every
number directly from the raw data (no reliance on prior cached values).

Per (composite × year × group), reports on the listwise-complete sample:
    N           — listwise complete cases
    Items       — number of items in the composite
    KMO         — Kaiser-Meyer-Olkin sampling adequacy
    Bartlett    — Bartlett's test of sphericity (χ², p)
    α           — Cronbach's alpha
    PC1 Var%    — % variance explained by the first principal component
    Kaiser k    — number of components with eigenvalue > 1
    Loadings    — min–max absolute loading on PC1
    Assessment  — Excellent / Good / Acceptable / Borderline / Weak based on α

Reverse-coding is applied per CLAUDE.md before computation:
    Superordinate Identity: Q67_4 / K6X5_3
    Belief in Inevitable Conflict: Q63_1, Q63_2 / K6X1_1, K6X1_2

Outputs:
    reports/PCA_Reliability_Summary.docx
    code/_pca_reliability.tsv
"""

import math
from pathlib import Path

import numpy as np
import pandas as pd
import pyreadstat
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler
from factor_analyzer.factor_analyzer import (
    calculate_kmo,
    calculate_bartlett_sphericity,
)
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

ROOT = Path(__file__).parent.parent

# ---------- Load both waves -----------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy()

df20, _ = pyreadstat.read_sav(
    str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"), encoding="latin1"
)
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy()


# ---------- Composite specifications --------------------------------------
SPECS = [
    # (label, year, group, df, items, reverse_items, scale_max)
    # Superordinate Identity (3-item, Q67_4 reversed)
    ("Superordinate Identity", "2023", "Estonian", df23, ["Q67_2", "Q67_4", "Q67_5"], ["Q67_4"], 4),
    ("Superordinate Identity", "2023", "Russian",  df23, ["Q67_2", "Q67_4", "Q67_5"], ["Q67_4"], 4),
    ("Superordinate Identity", "2020", "Estonian", df20, ["K6X5_2", "K6X5_3", "K6X5_4"], ["K6X5_3"], 4),
    ("Superordinate Identity", "2020", "Russian",  df20, ["K6X5_2", "K6X5_3", "K6X5_4"], ["K6X5_3"], 4),
    # SD: Primary Out-group (group-specific items)
    ("SD: Primary Out-group", "2023", "Estonian", df23, ["Q57_1", "Q58_1", "Q59_1"], None, 5),
    ("SD: Primary Out-group", "2023", "Russian",  df23, ["Q57_2", "Q58_2", "Q59_2"], None, 5),
    ("SD: Primary Out-group", "2020", "Estonian", df20, ["K4X7_1", "K4X8_1", "K4X9_1"], None, 5),
    ("SD: Primary Out-group", "2020", "Russian",  df20, ["K4X7_2", "K4X8_2", "K4X9_2"], None, 5),
    # SD: General Out-group (item set differs across waves)
    ("SD: General Out-group", "2023", "Estonian", df23,
        ["Q57_4", "Q57_5", "Q58_4", "Q58_5", "Q59_4", "Q59_5"], None, 5),
    ("SD: General Out-group", "2023", "Russian",  df23,
        ["Q57_4", "Q57_5", "Q58_4", "Q58_5", "Q59_4", "Q59_5"], None, 5),
    ("SD: General Out-group", "2020", "Estonian", df20,
        ["K4X7_3", "K4X8_3", "K4X9_3"], None, 5),
    ("SD: General Out-group", "2020", "Russian",  df20,
        ["K4X7_3", "K4X8_3", "K4X9_3"], None, 5),
    # Comparative Opportunity Assessment (Q44_1..12 / K3X1_1..12)
    ("Comparative Opportunity Assessment", "2023", "Estonian", df23,
        [f"Q44_{i}" for i in range(1, 13)], None, 5),
    ("Comparative Opportunity Assessment", "2023", "Russian",  df23,
        [f"Q44_{i}" for i in range(1, 13)], None, 5),
    ("Comparative Opportunity Assessment", "2020", "Estonian", df20,
        [f"K3X1_{i}" for i in range(1, 13)], None, 5),
    ("Comparative Opportunity Assessment", "2020", "Russian",  df20,
        [f"K3X1_{i}" for i in range(1, 13)], None, 5),
    # Belief in Inevitable Conflict (Q63 / K6X1) — Q63_1/Q63_2 reversed
    ("Belief in Inevitable Conflict", "2023", "Estonian", df23,
        ["Q63_1", "Q63_2", "Q63_3", "Q63_4"], ["Q63_1", "Q63_2"], 4),
    ("Belief in Inevitable Conflict", "2023", "Russian",  df23,
        ["Q63_1", "Q63_2", "Q63_3", "Q63_4"], ["Q63_1", "Q63_2"], 4),
    ("Belief in Inevitable Conflict", "2020", "Estonian", df20,
        ["K6X1_1", "K6X1_2", "K6X1_3", "K6X1_4"], ["K6X1_1", "K6X1_2"], 4),
    ("Belief in Inevitable Conflict", "2020", "Russian",  df20,
        ["K6X1_1", "K6X1_2", "K6X1_3", "K6X1_4"], ["K6X1_1", "K6X1_2"], 4),
    # Minority Support Inclusion
    ("Minority Inclusion Support", "2023", "Estonian", df23, ["Q68_1", "Q68_2", "Q68_3"], None, 4),
    ("Minority Inclusion Support", "2023", "Russian",  df23, ["Q68_1", "Q68_2", "Q68_3"], None, 4),
    ("Minority Inclusion Support", "2020", "Estonian", df20, ["K6X6_1", "K6X6_2", "K6X6_3"], None, 4),
    ("Minority Inclusion Support", "2020", "Russian",  df20, ["K6X6_1", "K6X6_2", "K6X6_3"], None, 4),
    # Contact: Estonian Speakers
    ("Contact: Estonian Speakers", "2023", "Estonian", df23, [f"Q51_{i}" for i in range(1, 7)], None, 5),
    ("Contact: Estonian Speakers", "2023", "Russian",  df23, [f"Q51_{i}" for i in range(1, 7)], None, 5),
    ("Contact: Estonian Speakers", "2020", "Estonian", df20, [f"K4X1_{i}" for i in range(1, 7)], None, 5),
    ("Contact: Estonian Speakers", "2020", "Russian",  df20, [f"K4X1_{i}" for i in range(1, 7)], None, 5),
    # Contact: Russian Speakers
    ("Contact: Russian Speakers", "2023", "Estonian", df23, [f"Q52_{i}" for i in range(1, 7)], None, 5),
    ("Contact: Russian Speakers", "2023", "Russian",  df23, [f"Q52_{i}" for i in range(1, 7)], None, 5),
    ("Contact: Russian Speakers", "2020", "Estonian", df20, [f"K4X2_{i}" for i in range(1, 7)], None, 5),
    ("Contact: Russian Speakers", "2020", "Russian",  df20, [f"K4X2_{i}" for i in range(1, 7)], None, 5),
]


# ---------- Helpers --------------------------------------------------------
def prep_matrix(df, ethnicity, items, reverse_items, scale_max, dk_code=9):
    """Return listwise-complete numeric matrix with reverse-coding applied."""
    sub = df[df["ethnicity_binary"] == ethnicity][items].apply(
        pd.to_numeric, errors="coerce"
    ).copy()
    sub = sub.where(sub != dk_code)  # DK → NaN
    if reverse_items:
        for col in reverse_items:
            sub[col] = (scale_max + 1) - sub[col]
    return sub.dropna()  # listwise


def cronbach_alpha(X):
    """Cronbach's α for a numeric DataFrame (rows = cases, cols = items)."""
    X = X.values
    k = X.shape[1]
    item_var = X.var(axis=0, ddof=1).sum()
    total_var = X.sum(axis=1).var(ddof=1)
    return (k / (k - 1)) * (1 - item_var / total_var)


def assess(alpha):
    if alpha >= 0.90: return "Excellent"
    if alpha >= 0.80: return "Good"
    if alpha >= 0.70: return "Acceptable"
    if alpha >= 0.65: return "Borderline"
    return "Weak"


def fmt_p(p):
    if p < 0.0001: return "<.0001"
    if p < 0.001:  return f"{p:.4f}"
    return f"{p:.3f}"


# ---------- Compute every row ---------------------------------------------
def ethnicity_code(group):
    return 0 if group == "Estonian" else 1


rows = []
for (label, year, group, df, items, reverse_items, scale_max) in SPECS:
    X = prep_matrix(df, ethnicity_code(group), items, reverse_items, scale_max)
    n_items = X.shape[1]
    n = X.shape[0]

    # KMO
    try:
        _, kmo_overall = calculate_kmo(X)
    except Exception:
        kmo_overall = float("nan")

    # Bartlett's test of sphericity
    try:
        chi2, p_bart = calculate_bartlett_sphericity(X)
    except Exception:
        chi2, p_bart = float("nan"), float("nan")

    # Cronbach's α
    alpha = cronbach_alpha(X)

    # PCA on z-scored items
    Xz = StandardScaler().fit_transform(X)
    pca = PCA(n_components=n_items)
    pca.fit(Xz)
    var_ratios = pca.explained_variance_ratio_
    eigenvalues = pca.explained_variance_  # variance of each component
    pc1_var_pct = var_ratios[0] * 100.0
    kaiser_k = int(np.sum(eigenvalues > 1.0))

    # PC1 loadings = eigenvector * sqrt(eigenvalue) (correlation of items w/ PC1)
    loadings_pc1 = pca.components_[0] * math.sqrt(eigenvalues[0])
    abs_loadings = np.abs(loadings_pc1)
    load_min, load_max = abs_loadings.min(), abs_loadings.max()

    rows.append({
        "Variable": label,
        "Year": year,
        "Group": group,
        "Items": n_items,
        "N": n,
        "KMO": kmo_overall,
        "Bartlett_chi2": chi2,
        "Bartlett_p": p_bart,
        "alpha": alpha,
        "PC1_Var_pct": pc1_var_pct,
        "Kaiser_k": kaiser_k,
        "Load_min": load_min,
        "Load_max": load_max,
        "Assessment": assess(alpha),
    })


# ---------- Console sanity check ------------------------------------------
print(f"{'Variable':<38} {'Yr':<5} {'Grp':<10} {'N':<5} {'k':<3} "
      f"{'KMO':<6} {'α':<7} {'PC1%':<7} {'Kaiser':<7} {'Loadings':<14} {'Assess':<12}")
print("-" * 130)
for r in rows:
    print(
        f"{r['Variable']:<38} {r['Year']:<5} {r['Group']:<10} "
        f"{r['N']:<5} {r['Items']:<3} "
        f"{r['KMO']:.3f}  {r['alpha']:.3f}  {r['PC1_Var_pct']:5.1f}%  "
        f"{r['Kaiser_k']:<7} "
        f".{int(round(r['Load_min']*100)):02d}–.{int(round(r['Load_max']*100)):02d}     "
        f"{r['Assessment']}"
    )


# ---------- Write TSV (machine-readable copy) -----------------------------
tsv_path = ROOT / "code" / "_pca_reliability.tsv"
with open(tsv_path, "w") as f:
    f.write("Variable\tYear\tGroup\tItems\tN\tKMO\tBartlett_chi2\tBartlett_p\t"
            "alpha\tPC1_Var_pct\tKaiser_k\tLoad_min\tLoad_max\tAssessment\n")
    for r in rows:
        f.write(
            f"{r['Variable']}\t{r['Year']}\t{r['Group']}\t{r['Items']}\t{r['N']}\t"
            f"{r['KMO']:.4f}\t{r['Bartlett_chi2']:.4f}\t{r['Bartlett_p']:.6g}\t"
            f"{r['alpha']:.4f}\t{r['PC1_Var_pct']:.4f}\t{r['Kaiser_k']}\t"
            f"{r['Load_min']:.4f}\t{r['Load_max']:.4f}\t{r['Assessment']}\n"
        )
print(f"\nSaved: {tsv_path}")


# ---------- Write the Word document ---------------------------------------
doc = Document()

section = doc.sections[0]
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
# Landscape for wide table
from docx.enum.section import WD_ORIENTATION
section.orientation = WD_ORIENTATION.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)

# Title
h = doc.add_paragraph()
run = h.add_run("PCA & Reliability Validation Summary")
run.bold = True
run.font.size = Pt(16)

p = doc.add_paragraph()
run = p.add_run(
    "Sampling adequacy (KMO), sphericity (Bartlett), internal consistency "
    "(Cronbach's α), and dimensionality diagnostics (PCA: PC1 variance, "
    "Kaiser components, PC1 loading range) for every multi-item composite × "
    "year × group. All values computed on the listwise-complete sample. "
    "Reverse-coding (Q67_4 / K6X5_3 for Superordinate Identity; "
    "Q63_1, Q63_2 / K6X1_1, K6X1_2 for Belief in Inevitable Conflict) is "
    "applied before all computations."
)
run.font.size = Pt(9)
run.italic = True

doc.add_paragraph()

# Table header
cols = [
    "Variable", "Year", "Group", "N", "Items",
    "KMO", "Bartlett χ² (p)", "α", "PC1 Var %", "Kaiser k",
    "PC1 Loadings |min–max|", "Assessment",
]
table = doc.add_table(rows=1, cols=len(cols))
table.style = "Light Grid Accent 1"
table.alignment = WD_ALIGN_PARAGRAPH.LEFT

for i, col in enumerate(cols):
    cell = table.rows[0].cells[i]
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(col)
    run.bold = True
    run.font.size = Pt(8.5)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Body
for r in rows:
    bart_str = f"{r['Bartlett_chi2']:.0f} ({fmt_p(r['Bartlett_p'])})"
    cells_text = [
        r["Variable"],
        r["Year"],
        r["Group"],
        str(r["N"]),
        str(r["Items"]),
        f"{r['KMO']:.3f}",
        bart_str,
        f"{r['alpha']:.3f}",
        f"{r['PC1_Var_pct']:.1f}%",
        str(r["Kaiser_k"]),
        f".{int(round(r['Load_min']*100)):02d}–.{int(round(r['Load_max']*100)):02d}",
        r["Assessment"],
    ]
    row_cells = table.add_row().cells
    for i, txt in enumerate(cells_text):
        c = row_cells[i]
        c.text = ""
        para = c.paragraphs[0]
        run = para.add_run(txt)
        run.font.size = Pt(8)
        # Bold the alpha + assessment columns for at-a-glance reading
        if i in (7, 11):
            run.bold = True
        if i in (3, 4, 5, 6, 7, 8, 9, 10):
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_paragraph()

# Footer notes
notes = doc.add_paragraph()
run = notes.add_run(
    "Notes. N = listwise-complete cases (PCA/reliability sample); composite scores "
    "in `_effect_sizes.tsv` use pairwise deletion and therefore have somewhat larger N. "
    "KMO interpretation: ≥ .90 marvelous, ≥ .80 meritorious, ≥ .70 middling, "
    "≥ .60 mediocre, ≥ .50 miserable, < .50 unacceptable (Kaiser & Rice, 1974). "
    "Bartlett's test of sphericity: significance (p < .05) indicates the correlation "
    "matrix is sufficiently distinct from an identity matrix to support factor analysis. "
    "Cronbach's α thresholds: Excellent ≥ .90, Good ≥ .80, Acceptable ≥ .70, "
    "Borderline .65–.70, Weak < .65. PC1 loadings reported as |min–max| of the "
    "absolute correlation between each item and the first principal component "
    "(eigenvector × √eigenvalue). Kaiser k = number of components with eigenvalue > 1; "
    "values > 1 indicate that a multi-factor solution may also be tenable, but the "
    "single-factor composite has been retained where α > .70 and theoretical coherence "
    "supports it (CLAUDE.md). Single-item variables (Group ID Patterns, Territorial "
    "Attachment) are excluded from this table because PCA / α are undefined for k = 1. "
    "SD: General Out-group is computed on different item sets in 2020 (3 items, "
    "new immigrants) vs. 2023 (6 items, other Europeans + non-Europeans); cross-year "
    "PCA/α comparisons for that variable should be interpreted with caution."
)
run.font.size = Pt(8)
run.italic = True

out_path = ROOT / "reports" / "PCA_Reliability_Summary.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
