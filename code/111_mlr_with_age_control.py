"""
Re-fits the MLR models from script 99 with age added as a demographic
control, per cell. Produces:

  Table 1 (additive):
    outcome = B0 + B1*Contact + B2*Language + B3*Age + ε

  Table 2 (interactive):
    outcome = B0 + B1*Contact + B2*Language + B3*Age
              + B4*(Contact × Language) + ε

All centered predictors within cell. HC3 robust SEs. 32 cells per table.

Also compares the age-controlled B1 Contact and B2 Language to their values
in script 99's no-control versions, showing how much (if anything) age
absorbs from each focal predictor.

Outputs:
  code/_mlr_with_age_additive.tsv
  code/_mlr_with_age_interaction.tsv
  reports/MLR_With_Age_Control_APA.docx
  viz/fig_mlr_with_age_forest.jpg
"""

from pathlib import Path

import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm
import matplotlib.pyplot as plt
import matplotlib.lines as mlines
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent.parent

# ---------- Load data + helpers (same as script 99/106) ------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)
df20, _ = pyreadstat.read_sav(
    str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"), encoding="latin1"
)
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)

# Age columns (per CLAUDE.md mapping)
df23["age"] = pd.to_numeric(df23["T3"], errors="coerce")
df20["age"] = pd.to_numeric(df20["vanus"], errors="coerce")


def clean(df, items, reverse_items=None, scale_max=None, dk_code=9):
    sub = df[items].apply(pd.to_numeric, errors="coerce")
    sub = sub.where(sub != dk_code)
    if reverse_items:
        for it in reverse_items:
            sub[it] = (scale_max + 1) - sub[it]
    return sub.mean(axis=1, skipna=True)


def single_item(df, var, dk_code=9):
    s = pd.to_numeric(df[var], errors="coerce")
    return s.where(s != dk_code).astype(float)


def to_num(s, dk=9):
    s = pd.to_numeric(s, errors="coerce")
    return s.where(s != dk).astype(float)


CONTACT_2023 = {0: [f"Q52_{i}" for i in range(1, 7)], 1: [f"Q51_{i}" for i in range(1, 7)]}
CONTACT_2020 = {0: [f"K4X2_{i}" for i in range(1, 7)], 1: [f"K4X1_{i}" for i in range(1, 7)]}
LANG_2023 = {0: "Q71_2", 1: "Q71_1"}
LANG_2020 = {0: "K5_2",  1: "K5_1"}


def build_cell(spec, year, grp_code):
    df = df23 if year == 2023 else df20
    sub_df = df[df["ethnicity_binary"] == grp_code]
    if spec.get("single_item"):
        item = spec[f"items_{year}"]
        dk = spec.get(f"dk_{year}", 9)
        outcome = single_item(sub_df, item, dk_code=dk)
    elif isinstance(spec[f"items_{year}"], dict):
        grp_key = "E" if grp_code == 0 else "R"
        items = spec[f"items_{year}"][grp_key]
        rev = (spec[f"rev_{year}"][grp_key] if spec.get(f"rev_{year}") else None)
        outcome = clean(sub_df, items, rev, spec["smax"])
    else:
        items = spec[f"items_{year}"]
        rev = spec.get(f"rev_{year}")
        outcome = clean(sub_df, items, rev, spec["smax"])
    if spec["inv"]:
        outcome = (spec["smax"] + 1) - outcome
    contact_items = CONTACT_2023[grp_code] if year == 2023 else CONTACT_2020[grp_code]
    contact = 6 - clean(sub_df, contact_items)
    lang_col = LANG_2023[grp_code] if year == 2023 else LANG_2020[grp_code]
    lang = 7 - to_num(sub_df[lang_col])
    age = sub_df["age"].values
    return pd.DataFrame({
        "value": outcome.values, "contact": contact.values,
        "language": lang.values, "age": age,
    }).dropna()


SPECS = [
    {"name": "Superordinate Identity", "smax": 4, "inv": True,
     "items_2023": ["Q67_2","Q67_4","Q67_5"], "rev_2023": ["Q67_4"],
     "items_2020": ["K6X5_2","K6X5_3","K6X5_4"], "rev_2020": ["K6X5_3"]},
    {"name": "SD: Primary Out-group", "smax": 5, "inv": False,
     "items_2023": {"E": ["Q57_1","Q58_1","Q59_1"], "R": ["Q57_2","Q58_2","Q59_2"]},
     "items_2020": {"E": ["K4X7_1","K4X8_1","K4X9_1"], "R": ["K4X7_2","K4X8_2","K4X9_2"]}},
    {"name": "SD: General Out-group", "smax": 5, "inv": False,
     "items_2023": ["Q57_4","Q57_5","Q58_4","Q58_5","Q59_4","Q59_5"],
     "items_2020": ["K4X7_3","K4X8_3","K4X9_3"]},
    {"name": "Comparative Opportunity Assessment", "smax": 5, "inv": True,
     "items_2023": [f"Q44_{i}" for i in range(1,13)],
     "items_2020": [f"K3X1_{i}" for i in range(1,13)]},
    {"name": "Belief in Inevitable Conflict", "smax": 4, "inv": False,
     "items_2023": ["Q63_1","Q63_2","Q63_3","Q63_4"], "rev_2023": ["Q63_1","Q63_2"],
     "items_2020": ["K6X1_1","K6X1_2","K6X1_3","K6X1_4"], "rev_2020": ["K6X1_1","K6X1_2"]},
    {"name": "Minority Inclusion Support", "smax": 4, "inv": True,
     "items_2023": ["Q68_1","Q68_2","Q68_3"],
     "items_2020": ["K6X6_1","K6X6_2","K6X6_3"]},
    {"name": "Group ID Patterns", "smax": 5, "inv": False, "single_item": True,
     "items_2023": "Q66", "items_2020": "K6X4",
     "dk_2023": 9, "dk_2020": 6},
    {"name": "Territorial Attachment", "smax": 4, "inv": True, "single_item": True,
     "items_2023": "Q67_1", "items_2020": "K6X5_1"},
]


def stars(p):
    if pd.isna(p): return ""
    if p < .001: return "***"
    if p < .01:  return "**"
    if p < .05:  return "*"
    if p < .10:  return "⁺"
    return ""


def fmt_p(p):
    if pd.isna(p): return "—"
    if p < .001: return "< .001"
    return f"{p:.3f}".lstrip("0")


# Load the no-control versions for comparison
no_age_add = pd.read_csv(ROOT / "code" / "_mlr_contact_language_additive.tsv",
                          sep="\t")
no_age_int = pd.read_csv(ROOT / "code" / "_mlr_contact_language_interaction.tsv",
                          sep="\t")


# Fit both models per cell, with age centered as a control
add_rows = []; int_rows = []
for spec in SPECS:
    for grp_code, grp_name in [(0, "Estonian"), (1, "Russian")]:
        for year in (2020, 2023):
            data = build_cell(spec, year, grp_code)
            if len(data) < 10: continue
            data["contact_c"]  = data["contact"]  - data["contact"].mean()
            data["language_c"] = data["language"] - data["language"].mean()
            data["age_c"]      = data["age"]      - data["age"].mean()
            data["cxl"] = data["contact_c"] * data["language_c"]

            sd_y = data["value"].std(ddof=1)
            sd_c = data["contact_c"].std(ddof=1)
            sd_l = data["language_c"].std(ddof=1)
            sd_a = data["age_c"].std(ddof=1)

            # ADDITIVE with age control
            X_a = sm.add_constant(data[["contact_c", "language_c", "age_c"]])
            m_a = sm.OLS(data["value"], X_a).fit(cov_type="HC3")
            B0a = m_a.params["const"]
            B1a = m_a.params["contact_c"];  SE1a = m_a.bse["contact_c"]; p1a = m_a.pvalues["contact_c"]
            B2a = m_a.params["language_c"]; SE2a = m_a.bse["language_c"]; p2a = m_a.pvalues["language_c"]
            B3a = m_a.params["age_c"];      SE3a = m_a.bse["age_c"];     p3a = m_a.pvalues["age_c"]
            beta1a = B1a * sd_c / sd_y
            beta2a = B2a * sd_l / sd_y
            beta3a = B3a * sd_a / sd_y

            # Compare to no-age additive
            prior = no_age_add[(no_age_add["variable"] == spec["name"]) &
                               (no_age_add["group"] == grp_name) &
                               (no_age_add["year"] == year)].iloc[0]
            add_rows.append({
                "variable": spec["name"], "group": grp_name, "year": year,
                "N": int(m_a.nobs), "B0": B0a,
                "B1_contact": B1a, "SE1": SE1a, "p1": p1a, "beta1": beta1a,
                "B2_language": B2a, "SE2": SE2a, "p2": p2a, "beta2": beta2a,
                "B3_age": B3a, "SE3": SE3a, "p3": p3a, "beta3": beta3a,
                "R2": m_a.rsquared,
                "B1_no_age": prior["B1_contact"], "B2_no_age": prior["B2_language"],
                "delta_B1": B1a - prior["B1_contact"],
                "delta_B2": B2a - prior["B2_language"],
                "R2_no_age": prior["R2"],
                "delta_R2": m_a.rsquared - prior["R2"],
            })

            # INTERACTIVE with age control
            X_i = sm.add_constant(data[["contact_c", "language_c", "age_c", "cxl"]])
            m_i = sm.OLS(data["value"], X_i).fit(cov_type="HC3")
            B0i = m_i.params["const"]
            B1i = m_i.params["contact_c"];  SE1i = m_i.bse["contact_c"]; p1i = m_i.pvalues["contact_c"]
            B2i = m_i.params["language_c"]; SE2i = m_i.bse["language_c"]; p2i = m_i.pvalues["language_c"]
            B3i = m_i.params["age_c"];      SE3i = m_i.bse["age_c"];     p3i = m_i.pvalues["age_c"]
            B4i = m_i.params["cxl"];        SE4i = m_i.bse["cxl"];       p4i = m_i.pvalues["cxl"]
            ll4, ul4 = m_i.conf_int().loc["cxl"]

            prior_i = no_age_int[(no_age_int["variable"] == spec["name"]) &
                                  (no_age_int["group"] == grp_name) &
                                  (no_age_int["year"] == year)].iloc[0]
            int_rows.append({
                "variable": spec["name"], "group": grp_name, "year": year,
                "N": int(m_i.nobs), "B0": B0i,
                "B1_contact": B1i, "SE1": SE1i, "p1": p1i,
                "B2_language": B2i, "SE2": SE2i, "p2": p2i,
                "B3_age": B3i, "SE3": SE3i, "p3": p3i,
                "B4_cxl": B4i, "SE4": SE4i, "p4": p4i,
                "CI4_low": ll4, "CI4_high": ul4,
                "R2": m_i.rsquared,
                "R2_additive_with_age": m_a.rsquared,
                "delta_R2_interaction": m_i.rsquared - m_a.rsquared,
                "B4_no_age": prior_i["B3_cxl"],
                "delta_B4": B4i - prior_i["B3_cxl"],
                "p4_no_age": prior_i["p3"],
            })

df_add = pd.DataFrame(add_rows)
df_int = pd.DataFrame(int_rows)
df_add.to_csv(ROOT / "code" / "_mlr_with_age_additive.tsv", sep="\t",
              index=False, float_format="%.4f")
df_int.to_csv(ROOT / "code" / "_mlr_with_age_interaction.tsv", sep="\t",
              index=False, float_format="%.4f")
print("Saved TSVs.\n")


# Console summary
print(f"{'Composite':<37}{'Group':<10}{'Year':>6}"
      f"{'B1 no age':>12}{'B1 +age':>10}{'Δ':>9}"
      f"{'B2 no age':>12}{'B2 +age':>10}{'Δ':>9}{'B age':>10}{'p age':>10}")
print("=" * 137)
for r in add_rows[:8]:
    print(f"{r['variable']:<37}{r['group']:<10}{r['year']:>6}"
          f"{r['B1_no_age']:>+12.4f}{r['B1_contact']:>+10.4f}{r['delta_B1']:>+9.4f}"
          f"{r['B2_no_age']:>+12.4f}{r['B2_language']:>+10.4f}{r['delta_B2']:>+9.4f}"
          f"{r['B3_age']:>+10.4f}{fmt_p(r['p3']):>10}")
print("... (and 24 more rows)")


# ============================================================================
#                          APA WORD DOC
# ============================================================================
def _set_cell_border(cell, **kwargs):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders"); tcPr.append(tcBorders)
    for side in ("top","left","bottom","right"):
        spec_ = kwargs.get(side)
        elem = tcBorders.find(qn(f"w:{side}"))
        if elem is None:
            elem = OxmlElement(f"w:{side}"); tcBorders.append(elem)
        if spec_ is None:
            elem.set(qn("w:val"), "nil")
        else:
            elem.set(qn("w:val"), spec_.get("val","single"))
            elem.set(qn("w:sz"),  str(spec_.get("sz", 8)))
            elem.set(qn("w:color"), spec_.get("color","auto"))


def apply_apa_borders(table, header_rows=1):
    n = len(table.rows)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            top = bottom = None
            if r_idx == 0:
                top = {"sz": 12, "val":"single", "color":"000000"}
            if r_idx == header_rows - 1:
                bottom = {"sz": 6, "val":"single", "color":"000000"}
            if r_idx == n - 1:
                bottom = {"sz": 12, "val":"single", "color":"000000"}
            _set_cell_border(cell, top=top, bottom=bottom, left=None, right=None)


def remove_table_style(table):
    tblPr = table._tbl.find(qn("w:tblPr"))
    if tblPr is not None:
        tblStyle = tblPr.find(qn("w:tblStyle"))
        if tblStyle is not None: tblPr.remove(tblStyle)
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.find(qn("w:shd"))
            if shd is not None: tcPr.remove(shd)


doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.2)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.orientation = 1
sec.page_width  = Cm(29.7); sec.page_height = Cm(21.0)
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)


def set_cell(cell, text, *, bold=False, italic=False, size=10,
              align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]; para.alignment = align
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        r = para.add_run(line)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic
        if i < len(lines) - 1: r.add_break()


def set_cell_italics(cell, segments, *, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]; para.alignment = align
    for text, bold, italic in segments:
        r = para.add_run(text)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic


def merge_v(table, rows, col):
    cells = [table.rows[r].cells[col] for r in rows]
    merged = cells[0]
    for c in cells[1:]: merged = merged.merge(c)
    return merged


def H_table_number(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.name = "Times New Roman"


def H_table_title(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.size = Pt(11); r.font.name = "Times New Roman"


def table_note(segments, size=9):
    p = doc.add_paragraph()
    for text, bold, italic in segments:
        r = p.add_run(text)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic


# Title
p = doc.add_paragraph()
r = p.add_run("MLR with Age as a Control — Contact + Language + Age")
r.bold = True; r.font.size = Pt(16); r.font.name = "Times New Roman"

p = doc.add_paragraph()
r = p.add_run(
    "Demographic-controlled multiple-regression models for each outcome × "
    "group × wave cell. Age (continuous, in years) is added to the additive "
    "MLR (Table 1) and the interactive MLR (Table 2) from script 99. All "
    "predictors mean-centered within cell. HC3 robust standard errors. "
    "Each table includes a comparison column showing the focal coefficients "
    "from the no-age version so age's effect on the substantive predictors "
    "is visible at a glance."
)
r.italic = True; r.font.size = Pt(10); r.font.name = "Times New Roman"
doc.add_paragraph()


# -------------- TABLE 1: Additive with age control -----------------------
H_table_number("Table 1")
H_table_title(
    "Additive MLR with age control "
    "(outcome = B₀ + B₁ × Contact + B₂ × Language + B₃ × Age + ε)."
)

header1 = ["Outcome", "Group", "Wave", "n",
           "B₁ Contact (SE)", "B₂ Language (SE)", "B₃ Age (SE)",
           "R²",
           "B₁ no-age", "Δ B₁", "B₂ no-age", "Δ B₂"]
tbl1 = doc.add_table(rows=1, cols=len(header1))
for i, h in enumerate(header1):
    if h == "n":
        set_cell_italics(tbl1.rows[0].cells[i], [("n", True, True)], size=9,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
    elif h == "R²":
        set_cell_italics(tbl1.rows[0].cells[i],
                          [("R", True, True), ("²", True, False)], size=9,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        set_cell(tbl1.rows[0].cells[i], h, bold=True, size=9,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

ordered_a = []
for spec in SPECS:
    for grp in ("Estonian", "Russian"):
        for year in (2020, 2023):
            m = [r for r in add_rows if r["variable"] == spec["name"]
                                       and r["group"] == grp and r["year"] == year]
            if m: ordered_a.append(m[0])

row_idx = 1
outcome_starts1 = {}
for r in ordered_a:
    row = tbl1.add_row().cells
    set_cell(row[0], r["variable"], size=9)
    set_cell(row[1], r["group"], size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[2], str(r["year"]), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[3], str(r["N"]), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[4], f"{r['B1_contact']:+.3f} ({r['SE1']:.3f}){stars(r['p1'])}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[5], f"{r['B2_language']:+.3f} ({r['SE2']:.3f}){stars(r['p2'])}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[6], f"{r['B3_age']:+.4f} ({r['SE3']:.4f}){stars(r['p3'])}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[7], f"{r['R2']:.3f}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[8], f"{r['B1_no_age']:+.3f}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[9], f"{r['delta_B1']:+.4f}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[10], f"{r['B2_no_age']:+.3f}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[11], f"{r['delta_B2']:+.4f}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    if r["variable"] not in outcome_starts1:
        outcome_starts1[r["variable"]] = row_idx
    row_idx += 1

for spec in SPECS:
    start = outcome_starts1.get(spec["name"])
    if start is None: continue
    end = start + 3
    if end < len(tbl1.rows):
        merge_v(tbl1, list(range(start, end + 1)), 0)
        set_cell(tbl1.rows[start].cells[0], spec["name"], size=9)
remove_table_style(tbl1); apply_apa_borders(tbl1)

table_note([
    ("Note. ", False, True),
    ("Age is centered within cell; coefficient is on years. ", False, False),
    ("B", False, True), ("₃ Age is small in raw units because age is in "
                          "years (so each year of age changes the outcome "
                          "by ~0.001–0.01 points). Δ ", False, False),
    ("B", False, True), ("₁ and Δ ", False, False),
    ("B", False, True), ("₂ show how much contact and language coefficients "
                          "change when age enters the model. Small Δ → focal "
                          "predictors are independent of age. *** ", False, False),
    ("p", False, True), (" < .001, ** ", False, False),
    ("p", False, True), (" < .01, * ", False, False),
    ("p", False, True), (" < .05, ⁺ ", False, False),
    ("p", False, True), (" < .10.", False, False),
])

doc.add_paragraph(); doc.add_page_break()


# -------------- TABLE 2: Interactive with age control --------------------
H_table_number("Table 2")
H_table_title(
    "Interactive MLR with age control — adds Contact × Language interaction "
    "(outcome = B₀ + B₁ × Contact + B₂ × Language + B₃ × Age "
    "+ B₄ × (Contact × Language) + ε)."
)

header2 = ["Outcome", "Group", "Wave", "n",
           "B₁ Contact (SE)", "B₂ Language (SE)", "B₃ Age (SE)",
           "B₄ C × L (SE)", "p (B₄)", "R²",
           "Δ R²\n(interaction − additive)",
           "B₄ no-age", "Δ B₄"]
tbl2 = doc.add_table(rows=1, cols=len(header2))
for i, h in enumerate(header2):
    if h == "n":
        set_cell_italics(tbl2.rows[0].cells[i], [("n", True, True)], size=9,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
    elif h == "R²":
        set_cell_italics(tbl2.rows[0].cells[i],
                          [("R", True, True), ("²", True, False)], size=9,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
    elif h.startswith("Δ R²"):
        set_cell_italics(tbl2.rows[0].cells[i],
                          [("Δ ", True, False), ("R", True, True),
                           ("²", True, False),
                           ("\n(interaction − additive)", True, False)],
                          size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    elif "p (B₄)" == h:
        set_cell_italics(tbl2.rows[0].cells[i],
                          [("p", True, True), (" (B₄)", True, False)],
                          size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        set_cell(tbl2.rows[0].cells[i], h, bold=True, size=9,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

ordered_i = []
for spec in SPECS:
    for grp in ("Estonian", "Russian"):
        for year in (2020, 2023):
            m = [r for r in int_rows if r["variable"] == spec["name"]
                                       and r["group"] == grp and r["year"] == year]
            if m: ordered_i.append(m[0])

row_idx = 1
outcome_starts2 = {}
for r in ordered_i:
    row = tbl2.add_row().cells
    set_cell(row[0], r["variable"], size=9)
    set_cell(row[1], r["group"], size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[2], str(r["year"]), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[3], str(r["N"]), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[4], f"{r['B1_contact']:+.3f} ({r['SE1']:.3f}){stars(r['p1'])}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[5], f"{r['B2_language']:+.3f} ({r['SE2']:.3f}){stars(r['p2'])}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[6], f"{r['B3_age']:+.4f} ({r['SE3']:.4f}){stars(r['p3'])}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[7], f"{r['B4_cxl']:+.4f} ({r['SE4']:.4f}){stars(r['p4'])}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[8], fmt_p(r["p4"]), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[9], f"{r['R2']:.3f}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[10], f"{r['delta_R2_interaction']:+.4f}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[11], f"{r['B4_no_age']:+.4f}{stars(r['p4_no_age'])}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[12], f"{r['delta_B4']:+.4f}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    if r["variable"] not in outcome_starts2:
        outcome_starts2[r["variable"]] = row_idx
    row_idx += 1

for spec in SPECS:
    start = outcome_starts2.get(spec["name"])
    if start is None: continue
    end = start + 3
    if end < len(tbl2.rows):
        merge_v(tbl2, list(range(start, end + 1)), 0)
        set_cell(tbl2.rows[start].cells[0], spec["name"], size=9)
remove_table_style(tbl2); apply_apa_borders(tbl2)

table_note([
    ("Note. ", False, True),
    ("B", False, True), ("₄ Contact × Language is the focal moderation term. "
                          "Δ ", False, False),
    ("R", False, True), ("² (interaction − additive) shows how much variance "
                          "the Contact × Language term adds beyond the "
                          "additive Contact + Language + Age model — i.e., "
                          "the incremental variance explained by the "
                          "moderation. Δ ", False, False),
    ("B", False, True), ("₄ shows how much B₄ changes from the no-age "
                          "version (script 99 Table 2). Significance "
                          "as above.", False, False),
])

out_doc = ROOT / "reports" / "MLR_With_Age_Control_APA.docx"
doc.save(out_doc)
print(f"Saved Word doc: {out_doc}")


# ============================================================================
#                       FOREST PLOT: B1 Contact with vs. without age
# ============================================================================
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 9,
    "axes.edgecolor": "#333",
    "axes.linewidth": 0.7,
})

EST_COLOR = "#2563eb"
RUS_COLOR = "#d97706"

OUTCOMES = [s["name"] for s in SPECS]
PRETTY = {
    "Superordinate Identity":             "Superordinate Identity",
    "SD: Primary Out-group":              "SD: Primary Out-group",
    "SD: General Out-group":              "SD: General Out-group",
    "Comparative Opportunity Assessment": "Comparative Opportunity",
    "Belief in Inevitable Conflict":      "Belief in Inevitable Conflict",
    "Minority Inclusion Support":         "Minority Inclusion Support",
    "Group ID Patterns":                  "Group ID Patterns",
    "Territorial Attachment":             "Territorial Attachment",
}

fig, ax = plt.subplots(figsize=(15.0, 13.0), dpi=300)

y_positions = {}
y = len(OUTCOMES) * 6.5
section_gap = 0.8
within_gap = 0.9
for outcome in OUTCOMES:
    for cell in [("Estonian", 2020), ("Estonian", 2023),
                  ("Russian", 2020), ("Russian", 2023)]:
        y_positions[(outcome, *cell)] = y
        y -= within_gap
    y -= section_gap

y_min = y + 0.5
y_max = len(OUTCOMES) * 6.5 + 0.5

# Section banding
for i, outcome in enumerate(OUTCOMES):
    if i % 2 == 0:
        y_top = y_positions[(outcome, "Estonian", 2020)] + 0.5
        y_bot = y_positions[(outcome, "Russian", 2023)] - 0.5
        ax.axhspan(y_bot, y_top, facecolor="#f7f7f7", zorder=0)

# Plot two points per cell: B1 without age (open dot) and B1 with age (filled dot)
for r in add_rows:
    key = (r["variable"], r["group"], r["year"])
    if key not in y_positions: continue
    y_pos = y_positions[key]
    color = EST_COLOR if r["group"] == "Estonian" else RUS_COLOR

    # No-age B1 (open dot)
    ax.scatter([r["B1_no_age"]], [y_pos + 0.18], s=110, facecolor="white",
               edgecolor=color, linewidth=1.6, zorder=3, alpha=0.85)
    # With-age B1 (filled dot)
    ax.scatter([r["B1_contact"]], [y_pos - 0.18], s=130, color=color,
               edgecolor="white", linewidth=1.0, zorder=3)
    # Connector
    ax.plot([r["B1_no_age"], r["B1_contact"]], [y_pos + 0.18, y_pos - 0.18],
            color=color, linewidth=0.6, alpha=0.4, linestyle=(0, (2, 2)),
            zorder=2)

    # Annotation
    ann = (f"{r['group'][0]} {r['year']}  "
           f"B₁ no-age={r['B1_no_age']:+.3f}  →  +age={r['B1_contact']:+.3f}{stars(r['p1'])}    "
           f"Δ={r['delta_B1']:+.3f}")
    ax.text(max(r["B1_no_age"], r["B1_contact"]) + 0.005, y_pos, ann,
            ha="left", va="center", fontsize=7.8,
            color="#222", family="monospace")

ax.axvline(0, color="#666", linewidth=1.0, linestyle="-", zorder=1)

yticks, yticklabels = [], []
for outcome in OUTCOMES:
    y_top = y_positions[(outcome, "Estonian", 2020)]
    y_bot = y_positions[(outcome, "Russian", 2023)]
    yticks.append((y_top + y_bot) / 2)
    yticklabels.append(PRETTY[outcome])
ax.set_yticks(yticks)
ax.set_yticklabels(yticklabels, fontsize=10, fontweight="bold")
ax.set_ylim(y_min, y_max)

all_lo = min(df_add["B1_contact"].min(), df_add["B1_no_age"].min())
all_hi = max(df_add["B1_contact"].max(), df_add["B1_no_age"].max())
ax.set_xlim(all_lo - 0.02, all_hi + 0.35)
ax.set_xlabel("B₁ Contact — additive MLR coefficient",
              fontsize=10, color="#444", labelpad=8)

for s in ("top", "right"):
    ax.spines[s].set_visible(False)
ax.spines["left"].set_color("#ccc")
ax.spines["bottom"].set_color("#666")
ax.tick_params(axis="x", length=2.5, color="#666", labelsize=9)
ax.tick_params(axis="y", length=0)

# Title
fig.text(0.04, 0.975,
         "Effect of Age Control on B₁ Contact — Additive MLR",
         fontsize=14, fontweight="bold", ha="left")
fig.text(0.04, 0.957,
         "Open dot = B₁ Contact without age control (script 99). Filled dot = B₁ Contact with age control (script 111). Dashed line connects the two estimates per cell.",
         fontsize=9, color="#444", ha="left")
fig.text(0.04, 0.942,
         "If open and filled dots overlap, age does not affect the contact coefficient. Larger separation → age was absorbing or revealing more contact variance.",
         fontsize=9, color="#444", ha="left")

handles = [
    mlines.Line2D([], [], marker="o", color=EST_COLOR, markerfacecolor="white",
                  markeredgecolor=EST_COLOR, markersize=10, markeredgewidth=1.6,
                  linestyle="None", label="Estonian — no age"),
    mlines.Line2D([], [], marker="o", color=EST_COLOR, markerfacecolor=EST_COLOR,
                  markeredgecolor="white", markersize=10,
                  linestyle="None", label="Estonian — with age"),
    mlines.Line2D([], [], marker="o", color=RUS_COLOR, markerfacecolor="white",
                  markeredgecolor=RUS_COLOR, markersize=10, markeredgewidth=1.6,
                  linestyle="None", label="Russian — no age"),
    mlines.Line2D([], [], marker="o", color=RUS_COLOR, markerfacecolor=RUS_COLOR,
                  markeredgecolor="white", markersize=10,
                  linestyle="None", label="Russian — with age"),
]
fig.legend(handles=handles, loc="upper right", bbox_to_anchor=(0.99, 0.955),
           frameon=False, fontsize=9, ncol=4,
           handlelength=1.4, columnspacing=1.4)

fig.text(0.04, 0.012,
         "All centered predictors. HC3 robust standard errors. Significance: *** p<.001, ** p<.01, * p<.05, ⁺ p<.10.",
         fontsize=7.5, color="#555", ha="left")

plt.subplots_adjust(left=0.18, right=0.985, top=0.93, bottom=0.04)

out_jpg = ROOT / "viz" / "fig_mlr_with_age_forest.jpg"
plt.savefig(out_jpg, dpi=300, format="jpg", facecolor="white",
            pil_kwargs={"quality": 95})
print(f"Saved chart: {out_jpg}")
