"""
Build Word document summarizing latent-mean analysis for all 8 composites.

Reads the four TSVs produced by code/79_all_composites_latent_means.R and
emits one section per composite with four tables corresponding to the
four steps:

  Step 1: Cronbach's alpha per group × year cell
  Step 2: Multi-group CFA invariance (configural / metric / scalar) per year
  Step 3: Latent-factor intercepts (Estonian fixed at 0)
  Step 4: Estonian − Russian gap, 2020 vs 2023, plus change in gap

Output: reports/Latent_Means_All_Composites.docx
"""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path("/Users/brianwiggins/Desktop/Claude Code/EIM2")

ALPHA = pd.read_csv(ROOT / "code" / "_latent_means_alpha.tsv", sep="\t")
FIT = pd.read_csv(ROOT / "code" / "_latent_means_fit.tsv", sep="\t")
INT = pd.read_csv(ROOT / "code" / "_latent_means_intercepts.tsv", sep="\t")
GAP = pd.read_csv(ROOT / "code" / "_latent_means_gaps.tsv", sep="\t")

OUT_PATH = ROOT / "reports" / "Latent_Means_All_Composites.docx"

COMPOSITES = [
    "Superordinate Identity",
    "SD: Primary Out-group",
    "SD: General Out-group",
    "Comparative Opportunity Assessment",
    "Belief in Inevitable Conflict",
    "Minority Support Inclusion",
    "Contact: Estonian Speakers",
    "Contact: Russian Speakers",
]

NOTES = {
    "Superordinate Identity":
        "3 items (Q67_2, Q67_4_inv, Q67_5 in 2023 / K6X5_2, K6X5_3_inv, K6X5_4 in 2020). "
        "Q67_4 / K6X5_3 reverse-coded. Lower = stronger belonging.",
    "SD: Primary Out-group":
        "3 items (live near, work with, be friends with). GROUP-SPECIFIC referent: "
        "Estonians rate Russian-speakers; Russians rate Estonian-speakers. "
        "Higher = more social distance.",
    "SD: General Out-group":
        "Item set differs across waves. 2023 = 6 items (other Europeans + non-Europeans, "
        "3 contexts); 2020 = 3 items (new immigrants, 3 contexts). "
        "Multi-group CFA tested separately within each year; cross-year invariance not testable. "
        "Higher = more social distance.",
    "Comparative Opportunity Assessment":
        "12 items. Lower = perceives Estonian advantage; higher = perceives advantage "
        "for other nationalities.",
    "Belief in Inevitable Conflict":
        "4 items (Q63_1_inv, Q63_2_inv, Q63_3, Q63_4 in 2023). "
        "Q63_1 + Q63_2 reverse-coded (Option B). Higher = stronger belief in conflict.",
    "Minority Support Inclusion":
        "3 items (Q68_1, Q68_2, Q68_3). Lower = more supportive of inclusion.",
    "Contact: Estonian Speakers":
        "6 items (Q51_1–Q51_6) covering 6 contexts. Lower = more contact.",
    "Contact: Russian Speakers":
        "6 items (Q52_1–Q52_6) covering 6 contexts. Lower = more contact.",
}


# -----------------------------------------------------------------------------
# Document setup
# -----------------------------------------------------------------------------
doc = Document()
for section in doc.sections:
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)


def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def add_para(text, bold=False, italic=False, size=10):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def fmt_num(x, dec=3):
    if pd.isna(x) or x == "":
        return "—"
    try:
        v = float(x)
    except (TypeError, ValueError):
        return str(x)
    return f"{v:.{dec}f}"


def fmt_p(x):
    if pd.isna(x) or x == "":
        return "—"
    try:
        v = float(x)
    except (TypeError, ValueError):
        return str(x)
    if v < 0.001:
        return "<.001"
    return f"{v:.3f}"


def add_table(headers, rows, col_widths=None, header_bold=True):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        cell = tbl.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = header_bold
        run.font.size = Pt(9)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
    if col_widths:
        for row in tbl.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = w
    return tbl


# -----------------------------------------------------------------------------
# Title page
# -----------------------------------------------------------------------------
title = doc.add_heading("Latent Means via Multi-Group CFA — All Composites", level=0)

intro = doc.add_paragraph()
intro.add_run(
    "This document reports a four-step latent-mean analysis applied uniformly to each "
    "of the eight multi-item composites in the EIM2 project. For each composite, the "
    "procedure is:"
).font.size = Pt(10)

steps = doc.add_paragraph()
steps.style = "List Number"
steps.add_run("Reliability check — Cronbach's α for the items, separately for each "
              "group × year cell.").font.size = Pt(10)
doc.add_paragraph(
    "Multi-group confirmatory factor analysis (Estonian vs Russian) within each year. "
    "Configural, metric, and scalar invariance models fit via lavaan with MLR estimator "
    "and FIML missing-data handling.",
    style="List Number"
).runs[0].font.size = Pt(10)
doc.add_paragraph(
    "Latent-factor intercepts (means) extracted from the scalar model. Estonian is "
    "fixed as the reference group with latent mean = 0; the Russian latent mean is "
    "estimated relative to it.",
    style="List Number"
).runs[0].font.size = Pt(10)
doc.add_paragraph(
    "Gap calculation: Estonian − Russian latent mean, by year. Change in gap "
    "(2023 − 2020) reported.",
    style="List Number"
).runs[0].font.size = Pt(10)

doc.add_paragraph()

caveats_h = doc.add_heading("Caveats", level=2)
add_para(
    "• SD: Primary Out-group uses GROUP-SPECIFIC items (Estonians rate Russian-speakers; "
    "Russians rate Estonian-speakers). Multi-group CFA is computable but conceptually "
    "asymmetric — the latent factor represents \"social distance toward primary out-group\" "
    "where the referent differs by respondent group."
)
add_para(
    "• SD: General Out-group uses different items in 2020 (3 items, new immigrants) vs. "
    "2023 (6 items, other Europeans + non-Europeans). Multi-group CFA was run separately "
    "within each year; cross-wave latent-mean comparison would require shared items."
)
add_para(
    "• Latent-mean comparisons assume scalar invariance. Where scalar invariance does not "
    "hold (ΔCFI < −0.01 or ΔRMSEA > 0.015), the latent means in the scalar model are still "
    "computed but include intercept variation that may not be purely substantive. The "
    "Step-2 invariance tables let the reader see when this caveat applies."
)
add_para(
    "• Sign of the gap follows the scale direction of each composite (described in the "
    "header note for each composite). A positive Estonian − Russian gap means Estonians "
    "score higher on that latent factor."
)

doc.add_paragraph()
doc.add_page_break()


# -----------------------------------------------------------------------------
# Per-composite sections
# -----------------------------------------------------------------------------
for comp in COMPOSITES:
    add_heading(comp, level=1)
    add_para(NOTES[comp], italic=True)
    doc.add_paragraph()

    # ----- STEP 1: Alpha -----
    add_heading("Step 1 — Reliability (Cronbach's α)", level=2)
    a = ALPHA[ALPHA["composite"] == comp].copy()
    rows = []
    for _, r in a.iterrows():
        rows.append([
            r["year"],
            r["group"],
            int(r["n_items"]),
            fmt_num(r["alpha"], 3),
            int(r["n"]),
        ])
    add_table(["Year", "Group", "k items", "α", "N"], rows)
    doc.add_paragraph()

    # ----- STEP 2: Invariance -----
    add_heading("Step 2 — Multi-Group CFA Invariance", level=2)
    f = FIT[FIT["composite"] == comp].copy()
    for yr in [2023, 2020]:
        add_para(f"{yr}", bold=True, size=10)
        fy = f[f["year"] == yr].copy()
        rows = []
        for _, r in fy.iterrows():
            rows.append([
                r["level"],
                fmt_num(r["chisq"], 2),
                "" if pd.isna(r["df"]) else int(r["df"]),
                fmt_p(r["p"]),
                fmt_num(r["cfi"], 3),
                fmt_num(r["tli"], 3),
                fmt_num(r["rmsea"], 3),
                fmt_num(r["srmr"], 3),
                fmt_num(r["d_cfi"], 3),
                fmt_num(r["d_rmsea"], 3),
            ])
        add_table(
            ["Level", "χ²", "df", "p", "CFI", "TLI", "RMSEA", "SRMR", "ΔCFI", "ΔRMSEA"],
            rows
        )
        # LRT chi-square diff row
        met_row = fy[fy["level"] == "metric"].iloc[0]
        sca_row = fy[fy["level"] == "scalar"].iloc[0]
        lrt_text = (
            f"LRT (Satorra-Bentler scaled): "
            f"configural→metric Δχ²={fmt_num(met_row['d_chisq'], 2)}, "
            f"df={'' if pd.isna(met_row['d_df']) else int(met_row['d_df'])}, "
            f"p={fmt_p(met_row['d_p'])}; "
            f"metric→scalar Δχ²={fmt_num(sca_row['d_chisq'], 2)}, "
            f"df={'' if pd.isna(sca_row['d_df']) else int(sca_row['d_df'])}, "
            f"p={fmt_p(sca_row['d_p'])}"
        )
        add_para(lrt_text, italic=True, size=9)
        doc.add_paragraph()

    # ----- STEP 3: Latent intercepts -----
    add_heading("Step 3 — Latent-Factor Intercepts (Estonian fixed at 0)", level=2)
    i_df = INT[INT["composite"] == comp].copy()
    rows = []
    for yr in [2023, 2020]:
        sub = i_df[i_df["year"] == yr]
        for g in ["Estonian", "Russian"]:
            r = sub[sub["group_label"] == g]
            if len(r) == 0:
                continue
            r = r.iloc[0]
            est = r["est"]
            se = r["se"]
            z = r["z"]
            p = r["p"]
            ci_lo = r["ci_lower"]
            ci_hi = r["ci_upper"]
            if g == "Estonian":
                rows.append([yr, g, fmt_num(est, 3), "—", "—", "—", "—"])
            else:
                ci_str = f"[{fmt_num(ci_lo, 3)}, {fmt_num(ci_hi, 3)}]"
                rows.append([
                    yr, g, fmt_num(est, 3),
                    fmt_num(se, 3), fmt_num(z, 2), fmt_p(p), ci_str
                ])
    add_table(
        ["Year", "Group", "Latent mean", "SE", "z", "p", "95% CI"],
        rows
    )
    doc.add_paragraph()

    # ----- STEP 4: Gap -----
    add_heading("Step 4 — Gap (Estonian − Russian Latent Mean)", level=2)
    g_df = GAP[GAP["composite"] == comp].copy()
    rows = []
    for yr in [2023, 2020]:
        gr = g_df[g_df["year"] == yr]
        if len(gr) == 0:
            continue
        gr = gr.iloc[0]
        rows.append([
            yr,
            fmt_num(gr["estonian_lm"], 3),
            fmt_num(gr["russian_lm"], 3),
            fmt_num(gr["gap_est_minus_rus"], 3),
        ])
    add_table(
        ["Year", "Estonian", "Russian", "Gap (Est − Rus)"],
        rows
    )
    g23 = g_df[g_df["year"] == 2023]["gap_est_minus_rus"].iloc[0]
    g20 = g_df[g_df["year"] == 2020]["gap_est_minus_rus"].iloc[0]
    delta = g23 - g20
    add_para(
        f"Change in gap (2023 − 2020): {delta:+.3f}    "
        f"|Gap 2023| = {abs(g23):.3f}    |Gap 2020| = {abs(g20):.3f}",
        bold=True
    )

    doc.add_page_break()


# -----------------------------------------------------------------------------
# Summary table — all gaps in one place
# -----------------------------------------------------------------------------
add_heading("Summary — All Composites", level=1)
add_para(
    "Estonian − Russian latent-mean gap by year, plus change. Sign follows the "
    "scale direction of each composite; see per-composite header notes for "
    "interpretation.",
    italic=True
)
doc.add_paragraph()

rows = []
for comp in COMPOSITES:
    g_df = GAP[GAP["composite"] == comp]
    g23 = g_df[g_df["year"] == 2023]["gap_est_minus_rus"].iloc[0]
    g20 = g_df[g_df["year"] == 2020]["gap_est_minus_rus"].iloc[0]
    rows.append([
        comp,
        fmt_num(g20, 3),
        fmt_num(g23, 3),
        fmt_num(g23 - g20, 3),
        fmt_num(abs(g23) - abs(g20), 3),
    ])
add_table(
    ["Composite", "Gap 2020", "Gap 2023", "Δ Gap (signed)", "Δ |Gap| (magnitude)"],
    rows
)


# -----------------------------------------------------------------------------
# Save
# -----------------------------------------------------------------------------
OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")
