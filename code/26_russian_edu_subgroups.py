"""
Russian-respondent Subgroup Robustness — by Education Language
================================================================
Blindspot FLAG #2 — addresses the convenient absence of a within-Russian
moderator. Asks: do Russian-school graduates differ from Russian-speaking
respondents who had ANY exposure to Estonian-language schooling?

Methodology
-----------
- 2023 sample only (the 2020 SPSS file does not carry a clean equivalent
  of T19_1..T19_5 from which `edu_language` is derived; reported as a
  limitation).
- Two subgroups within the 2023 Russian sample:
    A. Russian-school only (`edu_language` == 2)        N ~ 484
    B. Any Estonian-school exposure (`edu_language` in [1, 3, 4])  N ~ 38
- Within Group A, recompute every Round-3 outcome (composite means + SDs)
  and compare to Group B. Cohen's d (RMS-SD) and 95% CI are reported per
  variable; with the small N in Group B the CIs are wide and only
  directional inference is appropriate.

Outputs
-------
  reports/Russian_Edu_Subgroup_Robustness.docx
"""

import math
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from scipy import stats

ROOT = Path(__file__).parent.parent

# ---------- Load and filter ------------------------------------------------
df = pd.read_csv(ROOT / "data" / "EIM23.csv")
df = df[df["ethnicity_binary"] == 1].copy()    # Russian respondents only

# Subgroups
df_rus_only = df[df["edu_language"] == 2].copy()
df_est_any  = df[df["edu_language"].isin([1, 3, 4])].copy()

print(f"2023 Russian respondents: N = {len(df)}")
print(f"  Russian-school only (edu_language == 2):           N = {len(df_rus_only)}")
print(f"  Any Estonian-school exposure (1, 3, or 4):         N = {len(df_est_any)}")
print(f"  Excluded (other / missing):                        N = "
      f"{len(df) - len(df_rus_only) - len(df_est_any)}")


# ---------- Helpers ---------------------------------------------------------
def to_num(s, dk=9):
    s = pd.to_numeric(s, errors="coerce")
    return s.where(s != dk)

def composite(df, items, reverse_items=None, scale_max=None):
    sub = df[items].apply(pd.to_numeric, errors="coerce")
    sub = sub.where(sub != 9)
    if reverse_items:
        for it in reverse_items:
            sub[it] = (scale_max + 1) - sub[it]
    return sub.mean(axis=1, skipna=True)

def invert(s, scale_max):
    return (scale_max + 1) - s.astype(float)

def m_sd_n(s):
    s = s.dropna().astype(float)
    return s.mean(), s.std(ddof=1), len(s)

def cohens_d_ci(m1, sd1, n1, m2, sd2, n2):
    s = math.sqrt((sd1**2 + sd2**2) / 2)
    d = (m1 - m2) / s
    se = math.sqrt((n1 + n2)/(n1*n2) + d**2/(2*(n1+n2-2)))
    return d, d - 1.96*se, d + 1.96*se

def welch_p(a, b):
    a = pd.Series(a).dropna().astype(float)
    b = pd.Series(b).dropna().astype(float)
    return float(stats.ttest_ind(a, b, equal_var=False).pvalue)

def stars(p):
    if p < 0.001: return "***"
    if p < 0.01:  return "**"
    if p < 0.05:  return "*"
    return "ns"


# ---------- Build all 2023 outcomes for both subgroups --------------------
def build(d):
    """Return dict of {variable_name: inverted_series} for one subgroup."""
    out = {}
    out["Superordinate Identity"] = invert(
        composite(d, ["Q67_2","Q67_4","Q67_5"], ["Q67_4"], scale_max=4), 4)
    out["SD: Primary Out-group (vs Estonian spkrs)"] = composite(
        d, ["Q57_2","Q58_2","Q59_2"])
    out["SD: General Out-group"] = composite(
        d, ["Q57_4","Q57_5","Q58_4","Q58_5","Q59_4","Q59_5"])
    out["Comparative Opportunity Assessment"] = invert(
        composite(d, [f"Q44_{i}" for i in range(1,13)]), 5)
    out["Belief in Inevitable Conflict"] = composite(
        d, ["Q63_1","Q63_2","Q63_3","Q63_4"], ["Q63_3","Q63_4"], scale_max=4)
    out["Minority Support Inclusion"] = invert(
        composite(d, ["Q68_1","Q68_2","Q68_3"]), 4)
    out["Contact: Estonian Speakers"] = invert(
        composite(d, [f"Q51_{i}" for i in range(1,7)]), 5)
    out["Contact: Russian Speakers"] = invert(
        composite(d, [f"Q52_{i}" for i in range(1,7)]), 5)
    out["Group ID Patterns"] = to_num(d["Q66"], dk=9)
    out["Territorial Attachment"] = invert(to_num(d["Q67_1"]), 4)
    return out

A = build(df_rus_only)
B = build(df_est_any)


# ---------- Compare each outcome -------------------------------------------
print("\n" + "="*98)
print(f"{'Variable':<42} {'A: Rus-only':>22} {'B: Est-any':>22} {'d [95% CI]':>22}")
print("-"*98)

results = []
for var in A.keys():
    mA, sA, nA = m_sd_n(A[var])
    mB, sB, nB = m_sd_n(B[var])
    if nA < 5 or nB < 5:
        continue
    d, lo, hi = cohens_d_ci(mB, sB, nB, mA, sA, nA)   # B - A so d>0 means Est-school higher
    p = welch_p(B[var], A[var])
    print(f"{var:<42} {mA:>5.2f} ({sA:.2f}), n={nA:<3d}    "
          f"{mB:>5.2f} ({sB:.2f}), n={nB:<3d}    "
          f"d = {d:+.2f} [{lo:+.2f}, {hi:+.2f}]  {stars(p)}")
    results.append({
        "Variable": var,
        "A_M": mA, "A_SD": sA, "A_N": nA,
        "B_M": mB, "B_SD": sB, "B_N": nB,
        "d": d, "CI_lo": lo, "CI_hi": hi, "p": p,
        "sig": stars(p),
    })

# ---------- Word doc ------------------------------------------------------
doc = Document()
section = doc.sections[0]
section.left_margin = section.right_margin = Cm(2.0)
section.top_margin = section.bottom_margin = Cm(2.0)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("Russian Respondents — Subgroup Robustness by Education Language")
r.bold = True; r.font.size = Pt(15)

p = doc.add_paragraph()
r = p.add_run("Within-Russian comparison: schooling in Russian-only vs. any Estonian-language exposure (2023 wave). "
              "This addresses the Blindspot Round 2 audit's flag that the headline ethnicity comparison may "
              "obscure heterogeneity within the Russian sample by educational pathway.")
r.font.size = Pt(9); r.italic = True

doc.add_paragraph()

p = doc.add_paragraph()
p.add_run("Subgroups (2023 Russian respondents only — 2020 SPSS file does not carry an equivalent of T19_1..T19_5)").bold = True

p = doc.add_paragraph()
p.add_run(f"  • A. Russian-only schooling (edu_language = 2):  N = {len(df_rus_only)}\n"
          f"  • B. Any Estonian-school exposure (1, 3, or 4):  N = {len(df_est_any)}\n"
          f"  • Excluded (other / missing): N = {len(df) - len(df_rus_only) - len(df_est_any)}")

doc.add_paragraph()

p = doc.add_paragraph()
r = p.add_run("Comparison Table — d is (B − A) / RMS-SD, so positive d means the Estonian-school-exposure subgroup scores higher on the construct.")
r.bold = True; r.font.size = Pt(11)

cols = ["Variable",
        "A. Russian-only M (SD), N",
        "B. Estonian-exposure M (SD), N",
        "Cohen's d", "95% CI", "p", "Sig"]
table = doc.add_table(rows=1, cols=len(cols))
table.style = "Light Grid Accent 1"
for i, c in enumerate(cols):
    cell = table.rows[0].cells[i]
    cell.text = ""
    pp = cell.paragraphs[0]
    rr = pp.add_run(c); rr.bold = True; rr.font.size = Pt(8.5)

for r_ in results:
    cells_text = [
        r_["Variable"],
        f"{r_['A_M']:.2f} ({r_['A_SD']:.2f}), n = {r_['A_N']}",
        f"{r_['B_M']:.2f} ({r_['B_SD']:.2f}), n = {r_['B_N']}",
        f"{r_['d']:+.2f}",
        f"[{r_['CI_lo']:+.2f}, {r_['CI_hi']:+.2f}]",
        f"{r_['p']:.3f}" if r_['p'] >= .001 else "<.001",
        r_["sig"],
    ]
    row = table.add_row().cells
    for i, txt in enumerate(cells_text):
        row[i].text = ""
        pp = row[i].paragraphs[0]
        rr = pp.add_run(txt); rr.font.size = Pt(8.5)
        if i in (3, 4, 5, 6):
            pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_paragraph()
note = doc.add_paragraph()
nr = note.add_run(
    "Notes. The Estonian-school-exposure subgroup (B) is small (N ≈ 38) so confidence intervals are wide; "
    "interpretation is restricted to the direction of effects, not their precise magnitude. A negative d indicates "
    "that Russian-school-only respondents scored higher on the construct than Russian respondents with any "
    "Estonian-language schooling exposure. "
    "Cohen's d uses root-mean-square SD: d = (M₁ − M₂) / √[(SD₁² + SD₂²) / 2]. "
    "Significance: *** p<.001, ** p<.01, * p<.05, ns = not significant."
)
nr.font.size = Pt(8); nr.italic = True

out = ROOT / "reports" / "Russian_Edu_Subgroup_Robustness.docx"
doc.save(out)
print(f"\nSaved: {out}")
