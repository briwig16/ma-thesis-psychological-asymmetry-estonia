"""
Bivariate regression of each attitudinal outcome on out-group contact,
run separately for each group × year cell.

Model (per cell):
    outcome = B0 + B1 * contact + ε

Where contact = out-group contact composite (inverted so higher = more contact):
    Estonian respondents → contact with Russian speakers (Q52 / K4X2)
    Russian respondents  → contact with Estonian speakers (Q51 / K4X1)

8 outcomes × 4 cells (Est 2020, Est 2023, Rus 2020, Rus 2023) = 32 regressions.

This is the most direct cross-sectional test of contact's predictive power on
each attitudinal outcome — no year covariate, no moderation, no language. Just:
"At this point in time, does contact predict this attitude in this group?"

Outputs:
  code/_bivariate_contact_per_cell.tsv
  reports/Bivariate_Contact_Per_Cell_APA.docx
"""

import math
from pathlib import Path

import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent.parent

df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)
df20, _ = pyreadstat.read_sav(
    str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"), encoding="latin1"
)
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)


def clean(df, items, reverse_items=None, scale_max=None, dk_code=9):
    sub = df[items].apply(pd.to_numeric, errors="coerce")
    sub = sub.where(sub != dk_code)
    if reverse_items:
        for it in reverse_items:
            sub[it] = (scale_max + 1) - sub[it]
    return sub.mean(axis=1, skipna=True)


def single_item(df, var, dk_code=9):
    s = pd.to_numeric(df[var], errors="coerce")
    return s.where(s != dk_code).astype(float)


CONTACT_2023 = {0: [f"Q52_{i}" for i in range(1, 7)], 1: [f"Q51_{i}" for i in range(1, 7)]}
CONTACT_2020 = {0: [f"K4X2_{i}" for i in range(1, 7)], 1: [f"K4X1_{i}" for i in range(1, 7)]}


def build_cell(spec, year, grp_code):
    df = df23 if year == 2023 else df20
    sub_df = df[df["ethnicity_binary"] == grp_code]

    if spec.get("single_item"):
        item = spec[f"items_{year}"]
        dk = spec.get(f"dk_{year}", 9)
        outcome = single_item(sub_df, item, dk_code=dk)
    elif isinstance(spec[f"items_{year}"], dict):
        grp_key = "E" if grp_code == 0 else "R"
        items = spec[f"items_{year}"][grp_key]
        rev = (spec[f"rev_{year}"][grp_key] if spec.get(f"rev_{year}") else None)
        outcome = clean(sub_df, items, rev, spec["smax"])
    else:
        items = spec[f"items_{year}"]
        rev = spec.get(f"rev_{year}")
        outcome = clean(sub_df, items, rev, spec["smax"])

    if spec["inv"]:
        outcome = (spec["smax"] + 1) - outcome

    contact_items = CONTACT_2023[grp_code] if year == 2023 else CONTACT_2020[grp_code]
    contact = 6 - clean(sub_df, contact_items)

    return pd.DataFrame({"value": outcome.values, "contact": contact.values}).dropna()


SPECS = [
    {"name": "Superordinate Identity", "smax": 4, "inv": True,
     "items_2023": ["Q67_2","Q67_4","Q67_5"], "rev_2023": ["Q67_4"],
     "items_2020": ["K6X5_2","K6X5_3","K6X5_4"], "rev_2020": ["K6X5_3"]},
    {"name": "SD: Primary Out-group", "smax": 5, "inv": False,
     "items_2023": {"E": ["Q57_1","Q58_1","Q59_1"], "R": ["Q57_2","Q58_2","Q59_2"]},
     "items_2020": {"E": ["K4X7_1","K4X8_1","K4X9_1"], "R": ["K4X7_2","K4X8_2","K4X9_2"]}},
    {"name": "SD: General Out-group", "smax": 5, "inv": False,
     "items_2023": ["Q57_4","Q57_5","Q58_4","Q58_5","Q59_4","Q59_5"],
     "items_2020": ["K4X7_3","K4X8_3","K4X9_3"]},
    {"name": "Comparative Opportunity Assessment", "smax": 5, "inv": True,
     "items_2023": [f"Q44_{i}" for i in range(1,13)],
     "items_2020": [f"K3X1_{i}" for i in range(1,13)]},
    {"name": "Belief in Inevitable Conflict", "smax": 4, "inv": False,
     "items_2023": ["Q63_1","Q63_2","Q63_3","Q63_4"], "rev_2023": ["Q63_1","Q63_2"],
     "items_2020": ["K6X1_1","K6X1_2","K6X1_3","K6X1_4"], "rev_2020": ["K6X1_1","K6X1_2"]},
    {"name": "Minority Inclusion Support", "smax": 4, "inv": True,
     "items_2023": ["Q68_1","Q68_2","Q68_3"],
     "items_2020": ["K6X6_1","K6X6_2","K6X6_3"]},
    {"name": "Group ID Patterns", "smax": 5, "inv": False, "single_item": True,
     "items_2023": "Q66", "items_2020": "K6X4",
     "dk_2023": 9, "dk_2020": 6},
    {"name": "Territorial Attachment", "smax": 4, "inv": True, "single_item": True,
     "items_2023": "Q67_1", "items_2020": "K6X5_1"},
]


def stars(p):
    if pd.isna(p): return ""
    if p < .001: return "***"
    if p < .01:  return "**"
    if p < .05:  return "*"
    if p < .10:  return "⁺"
    return ""


def fmt_p(p):
    if pd.isna(p): return "—"
    if p < .001:   return "< .001"
    return f"{p:.3f}".lstrip("0")


# Run regressions
rows = []
for spec in SPECS:
    for grp_code, grp_name in [(0, "Estonian"), (1, "Russian")]:
        for year in (2020, 2023):
            data = build_cell(spec, year, grp_code)
            if len(data) < 10:
                continue
            X = sm.add_constant(data["contact"])
            m = sm.OLS(data["value"], X).fit(cov_type="HC3")
            B0 = m.params["const"]; B1 = m.params["contact"]
            SE = m.bse["contact"]; p = m.pvalues["contact"]
            ll, ul = m.conf_int().loc["contact"]
            sd_x = data["contact"].std(ddof=1)
            sd_y = data["value"].std(ddof=1)
            beta = B1 * sd_x / sd_y if sd_y > 0 else float("nan")
            rows.append({
                "variable": spec["name"], "group": grp_name, "year": year,
                "N": int(m.nobs),
                "B0": B0, "B1": B1, "SE": SE,
                "CI_low": ll, "CI_high": ul, "p": p,
                "beta": beta, "R2": m.rsquared,
            })


out = pd.DataFrame(rows)
tsv = ROOT / "code" / "_bivariate_contact_per_cell.tsv"
out.to_csv(tsv, sep="\t", index=False, float_format="%.4f")
print(f"Saved TSV: {tsv}\n")

# Console summary
print(f"{'Outcome':<37}{'Group':<10}{'Year':>6}{'N':>6}{'B1':>10}{'β':>9}{'p':>11}{'R²':>8}")
print("=" * 97)
for r in rows:
    print(f"{r['variable']:<37}{r['group']:<10}{r['year']:>6}{r['N']:>6}"
          f"{r['B1']:>+10.3f}{r['beta']:>+9.3f}{fmt_p(r['p']):>11} {stars(r['p'])}{r['R2']:>8.3f}")


# ============================================================================
#                          APA WORD DOC
# ============================================================================
def _set_cell_border(cell, **kwargs):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders"); tcPr.append(tcBorders)
    for side in ("top","left","bottom","right"):
        spec_ = kwargs.get(side)
        elem = tcBorders.find(qn(f"w:{side}"))
        if elem is None:
            elem = OxmlElement(f"w:{side}"); tcBorders.append(elem)
        if spec_ is None:
            elem.set(qn("w:val"), "nil")
        else:
            elem.set(qn("w:val"), spec_.get("val","single"))
            elem.set(qn("w:sz"),  str(spec_.get("sz", 8)))
            elem.set(qn("w:color"), spec_.get("color","auto"))


def apply_apa_borders(table, header_rows=1):
    n = len(table.rows)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            top = bottom = None
            if r_idx == 0:
                top = {"sz": 12, "val":"single", "color":"000000"}
            if r_idx == header_rows - 1:
                bottom = {"sz": 6, "val":"single", "color":"000000"}
            if r_idx == n - 1:
                bottom = {"sz": 12, "val":"single", "color":"000000"}
            _set_cell_border(cell, top=top, bottom=bottom, left=None, right=None)


def remove_table_style(table):
    tblPr = table._tbl.find(qn("w:tblPr"))
    if tblPr is not None:
        tblStyle = tblPr.find(qn("w:tblStyle"))
        if tblStyle is not None: tblPr.remove(tblStyle)
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.find(qn("w:shd"))
            if shd is not None: tcPr.remove(shd)


doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(2.0)
sec.orientation = 1
sec.page_width  = Cm(29.7); sec.page_height = Cm(21.0)
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)


def set_cell(cell, text, *, bold=False, italic=False, size=10,
              align=WD_ALIGN_PARAGRAPH.LEFT, vcenter=False):
    cell.text = ""
    para = cell.paragraphs[0]; para.alignment = align
    if vcenter: cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        r = para.add_run(line)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic
        if i < len(lines) - 1: r.add_break()


def set_cell_italics(cell, segments, *, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]; para.alignment = align
    for text, bold, italic in segments:
        r = para.add_run(text)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic


def merge_v(table, rows, col):
    cells = [table.rows[r].cells[col] for r in rows]
    merged = cells[0]
    for c in cells[1:]: merged = merged.merge(c)
    return merged


def H_table_number(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.name = "Times New Roman"


def H_table_title(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.size = Pt(11); r.font.name = "Times New Roman"


def table_note(segments, size=9):
    p = doc.add_paragraph()
    for text, bold, italic in segments:
        r = p.add_run(text)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic


# Title
p = doc.add_paragraph()
r = p.add_run("Bivariate Regression: Out-group Contact as Predictor of Attitudinal Outcomes")
r.bold = True; r.font.size = Pt(15); r.font.name = "Times New Roman"

p = doc.add_paragraph()
r = p.add_run(
    "Bivariate OLS regressions of each attitudinal outcome on out-group "
    "contact frequency. Contact is the only predictor — no year covariate, "
    "no moderation, no demographic controls. Each row represents one "
    "outcome × ethnic group × wave combination, providing a cross-sectional "
    "estimate of contact's association with the outcome in that specific "
    "cell. Out-group contact: Estonian respondents = contact with Russian-"
    "speakers (Q52 / K4X2); Russian respondents = contact with Estonian-"
    "speakers (Q51 / K4X1). Inverted: higher = more frequent contact."
)
r.italic = True; r.font.size = Pt(10); r.font.name = "Times New Roman"

doc.add_paragraph()


# --- TABLE 1 ---------------------------------------------------------------
H_table_number("Table 1")
H_table_title(
    "Bivariate regression of each attitudinal outcome on out-group contact, "
    "per ethnic group × wave (outcome = B₀ + B₁ × Contact + ε)."
)

header = ["Outcome", "Group", "Wave", "n",
          "B₀ (Intercept)", "B₁ Contact (SE)", "95% CI on B₁",
          "β", "p", "R²"]
table = doc.add_table(rows=1, cols=len(header))
for i, h in enumerate(header):
    if h == "n":
        set_cell_italics(table.rows[0].cells[i], [("n", True, True)], size=10,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
    elif h == "R²":
        set_cell_italics(table.rows[0].cells[i],
                          [("R", True, True), ("²", True, False)], size=10,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
    elif h == "β":
        set_cell_italics(table.rows[0].cells[i], [("β", True, True)],
                          size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    elif h == "p":
        set_cell_italics(table.rows[0].cells[i], [("p", True, True)],
                          size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        set_cell(table.rows[0].cells[i], h, bold=True, size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

# Build rows grouped by outcome, with merged outcome column
# First collect rows by outcome to determine where to merge
ordered = []
for spec in SPECS:
    for grp_name in ("Estonian", "Russian"):
        for year in (2020, 2023):
            match = [r for r in rows if r["variable"] == spec["name"]
                                       and r["group"] == grp_name
                                       and r["year"] == year]
            if match:
                ordered.append(match[0])

row_idx = 1
outcome_starts = {}
for r in ordered:
    table_row = table.add_row().cells
    set_cell(table_row[0], r["variable"], size=10)
    set_cell(table_row[1], r["group"], size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table_row[2], str(r["year"]), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table_row[3], str(r["N"]), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table_row[4], f"{r['B0']:+.3f}", size=10,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table_row[5], f"{r['B1']:+.3f} ({r['SE']:.3f}){stars(r['p'])}",
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table_row[6], f"[{r['CI_low']:+.3f}, {r['CI_high']:+.3f}]",
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table_row[7], f"{r['beta']:+.3f}{stars(r['p'])}", size=10,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table_row[8], fmt_p(r["p"]), size=10,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table_row[9], f"{r['R2']:.3f}", size=10,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    if r["variable"] not in outcome_starts:
        outcome_starts[r["variable"]] = row_idx
    row_idx += 1

# Merge outcome column across rows
for spec in SPECS:
    start = outcome_starts.get(spec["name"])
    if start is None: continue
    end = start + 3  # 4 rows per outcome (Est 2020, Est 2023, Rus 2020, Rus 2023)
    if end >= len(table.rows): end = len(table.rows) - 1
    merge_v(table, list(range(start, end + 1)), 0)
    set_cell(table.rows[start].cells[0], spec["name"], size=10)

remove_table_style(table); apply_apa_borders(table)

table_note([
    ("Note. ", False, True),
    ("Each row is one bivariate OLS regression fit within a single "
     "ethnic group × wave cell. ", False, False),
    ("B", False, True), ("₀ = predicted outcome at contact = 0 (often not "
                          "substantively meaningful given the 1–5 contact "
                          "scale). ", False, False),
    ("B", False, True), ("₁ = unstandardized slope of contact on the outcome, "
                          "in original Likert units of the outcome per 1-unit "
                          "of contact. ", False, False),
    ("β", False, True), (" = standardized slope, in SD-units of the outcome per "
                          "SD-unit of contact. HC3 robust standard errors. "
                          "*** ", False, False),
    ("p", False, True), (" < .001, ** ", False, False),
    ("p", False, True), (" < .01, * ", False, False),
    ("p", False, True), (" < .05, ⁺ ", False, False),
    ("p", False, True), (" < .10.", False, False),
])

doc.add_paragraph()
table_note([
    ("Reading note. ", False, True),
    ("For positive-direction outcomes (Superordinate Identity, Comparative "
     "Opportunity, Minority Inclusion Support, Territorial Attachment), "
     "a positive ", False, False),
    ("β", False, True), (" supports classical contact theory (more contact → "
                          "more of the positive construct). For distance / "
                          "conflict outcomes (SD: Primary Out-group, SD: "
                          "General Out-group, Belief in Inevitable Conflict, "
                          "Group ID Patterns), a negative ", False, False),
    ("β", False, True), (" supports classical contact theory (more contact → "
                          "less distance / conflict / out-group identity).", False, False),
])

out = ROOT / "reports" / "Bivariate_Contact_Per_Cell_APA.docx"
doc.save(out)
print(f"\nSaved: {out}")
