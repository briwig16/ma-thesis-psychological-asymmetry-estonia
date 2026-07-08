"""
Measurement Invariance Report — Word Doc Generator
====================================================
Reads code/_invariance_results.tsv produced by code/36_measurement_invariance.R
and renders a formatted Word document summarizing configural / metric / scalar
invariance test results for the 8 multi-item composites, with verdicts and
interpretation notes.

Output: reports/Measurement_Invariance.docx
"""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent.parent
TSV  = ROOT / "code" / "_invariance_results.tsv"

df = pd.read_csv(TSV, sep="\t")

# Split composite into name + scope (e.g., "Superordinate Identity — multi-group 2023"
# → name="Superordinate Identity", scope="multi-group 2023")
df[["name", "scope"]] = df["composite"].str.split(" — ", n=1, expand=True)


# Whether a composite is "just-identified" at the configural level (3 items)
THREE_ITEM = {
    "Superordinate Identity",
    "SD: Primary Out-group",
    "Minority Support Inclusion",
}


# -------- Build the Word doc --------------------------------------------
doc = Document()
section = doc.sections[0]
section.left_margin = section.right_margin = Cm(1.7)
section.top_margin  = section.bottom_margin = Cm(2.0)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)

p = doc.add_paragraph()
r = p.add_run("Measurement Invariance Testing — Multi-Group and Longitudinal CFA")
r.bold = True; r.font.size = Pt(15)

p = doc.add_paragraph()
r = p.add_run(
    "This report tests measurement invariance for each of the 8 multi-item composites "
    "used in the analysis, across two dimensions:"
); r.font.size = Pt(10)

bullets = [
    "Multi-group invariance (Estonian vs. Russian respondents) — within each wave",
    "Longitudinal invariance (2020 vs. 2023) — within each ethnic group",
]
for b in bullets:
    pp = doc.add_paragraph(b, style="List Bullet")
    pp.runs[0].font.size = Pt(10)

p = doc.add_paragraph()
r = p.add_run(
    "Each test runs a hierarchy of confirmatory factor analyses with progressively "
    "tighter equality constraints across groups/waves: "
); r.font.size = Pt(10)

p = doc.add_paragraph()
for line in [
    ("Configural", " — same factor structure (no equality constraints)"),
    ("Metric", " — factor loadings constrained equal across groups/waves"),
    ("Scalar", " — loadings AND item intercepts constrained equal"),
]:
    rr = p.add_run(line[0]); rr.bold = True; rr.font.size = Pt(10)
    rr = p.add_run(line[1] + "  "); rr.font.size = Pt(10)

p = doc.add_paragraph()
r = p.add_run(
    "Decision criteria (Cheung & Rensvold 2002; Chen 2007): invariance level HOLDS if "
    "ΔCFI ≥ −0.01 and ΔRMSEA ≤ +0.015. PARTIAL if ΔCFI ≥ −0.02 and ΔRMSEA ≤ +0.030. "
    "Otherwise FAILS."
); r.font.size = Pt(9.5); r.italic = True

doc.add_paragraph()


# -------- Headline Summary Table --------------------------------------
p = doc.add_paragraph()
r = p.add_run("Summary — Verdicts at Each Invariance Level")
r.bold = True; r.font.size = Pt(12)

cols = ["Composite", "Comparison", "Configural fit", "Metric", "Scalar"]
t = doc.add_table(rows=1, cols=len(cols))
t.style = "Light Grid Accent 1"
for i, c in enumerate(cols):
    cell = t.rows[0].cells[i]; cell.text = ""
    pp = cell.paragraphs[0]
    rr = pp.add_run(c); rr.bold = True; rr.font.size = Pt(9)

# Order composites in canonical order
COMP_ORDER = [
    "Superordinate Identity",
    "SD: Primary Out-group",
    "SD: General Out-group",
    "Comparative Opportunity Assessment",
    "Belief in Inevitable Conflict",
    "Minority Support Inclusion",
    "Contact: Estonian Speakers",
    "Contact: Russian Speakers",
]

def fit_label(row):
    if pd.isna(row["cfi"]) or pd.isna(row["rmsea"]):
        return "—"
    if row["df"] == 0:
        return "just-identified"
    cfi = row["cfi"]; rmsea = row["rmsea"]
    if cfi >= 0.95 and rmsea < 0.06:
        return f"CFI={cfi:.3f}, RMSEA={rmsea:.3f} ✓"
    elif cfi >= 0.90 and rmsea < 0.08:
        return f"CFI={cfi:.3f}, RMSEA={rmsea:.3f} (acceptable)"
    else:
        return f"CFI={cfi:.3f}, RMSEA={rmsea:.3f} (poor)"

for cmp_name in COMP_ORDER:
    sub = df[df["name"] == cmp_name]
    for scope in sub["scope"].unique():
        rows = sub[sub["scope"] == scope]
        cfg = rows[rows["level"] == "configural"]
        met = rows[rows["level"] == "metric"]
        sca = rows[rows["level"] == "scalar"]
        cells = t.add_row().cells

        cells[0].text = ""
        pp = cells[0].paragraphs[0]
        rr = pp.add_run(cmp_name); rr.font.size = Pt(8.5); rr.bold = True
        if cmp_name in THREE_ITEM:
            rr = pp.add_run("\n3-item composite (just-identified at configural)")
            rr.font.size = Pt(7); rr.italic = True

        cells[1].text = ""
        pp = cells[1].paragraphs[0]
        rr = pp.add_run(scope); rr.font.size = Pt(8.5)

        cfg_label = fit_label(cfg.iloc[0]) if len(cfg) else "—"
        cells[2].text = ""
        pp = cells[2].paragraphs[0]
        rr = pp.add_run(cfg_label); rr.font.size = Pt(8)

        for c_i, lvl_df in [(3, met), (4, sca)]:
            cells[c_i].text = ""
            pp = cells[c_i].paragraphs[0]
            if len(lvl_df) == 0:
                rr = pp.add_run("—")
            else:
                v = lvl_df.iloc[0]["verdict"]
                rr = pp.add_run(v if pd.notna(v) else "—")
                rr.bold = True
                if v == "HOLDS":
                    rr.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
                elif v == "PARTIAL":
                    rr.font.color.rgb = RGBColor(0xb4, 0x53, 0x09)
                elif v == "FAILS":
                    rr.font.color.rgb = RGBColor(0x99, 0x1b, 0x1e)
            rr.font.size = Pt(8.5)

doc.add_paragraph()


# -------- Per-composite detailed tables -------------------------------
p = doc.add_paragraph()
r = p.add_run("Detailed Fit Indices")
r.bold = True; r.font.size = Pt(12)

doc.add_paragraph()

for cmp_name in COMP_ORDER:
    sub = df[df["name"] == cmp_name]
    if len(sub) == 0: continue

    p = doc.add_paragraph()
    r = p.add_run(cmp_name)
    r.bold = True; r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(0x1f, 0x29, 0x37)
    if cmp_name in THREE_ITEM:
        p2 = doc.add_paragraph()
        rr = p2.add_run(
            "Note: 3-indicator model. The configural model is just-identified "
            "(df = 0), so configural fit indices are uninformative. Invariance "
            "tests via χ² difference still work but ΔCFI and ΔRMSEA at the "
            "metric/scalar levels can behave artifactually."
        )
        rr.font.size = Pt(8.5); rr.italic = True

    cols2 = ["Comparison", "Level", "χ²", "df", "p", "CFI", "TLI",
             "RMSEA", "SRMR", "Δχ² (Δdf, p)", "ΔCFI", "ΔRMSEA", "Verdict"]
    tt = doc.add_table(rows=1, cols=len(cols2))
    tt.style = "Light Grid Accent 1"
    for i, c in enumerate(cols2):
        cell = tt.rows[0].cells[i]; cell.text = ""
        pp = cell.paragraphs[0]
        rr = pp.add_run(c); rr.bold = True; rr.font.size = Pt(8)

    for scope in sub["scope"].unique():
        rows = sub[sub["scope"] == scope]
        for lvl in ("configural", "metric", "scalar"):
            r = rows[rows["level"] == lvl]
            if len(r) == 0: continue
            r = r.iloc[0]
            cells = tt.add_row().cells

            txts = [
                scope if lvl == "configural" else "",
                lvl,
                f"{r['chisq']:.2f}" if pd.notna(r['chisq']) else "—",
                f"{int(r['df'])}" if pd.notna(r['df']) else "—",
                (("<.001" if r["p"] < .001 else f"{r['p']:.3f}")
                 if pd.notna(r['p']) else "—"),
                f"{r['cfi']:.3f}" if pd.notna(r['cfi']) else "—",
                f"{r['tli']:.3f}" if pd.notna(r['tli']) else "—",
                f"{r['rmsea']:.3f}" if pd.notna(r['rmsea']) else "—",
                f"{r['srmr']:.3f}" if pd.notna(r['srmr']) else "—",
                ((f"{r['d_chisq']:.2f} ({int(r['d_df'])}, "
                  + (("<.001" if r['d_p'] < .001 else f"{r['d_p']:.3f}")
                     if pd.notna(r['d_p']) else "—") + ")")
                 if pd.notna(r['d_chisq']) else "—"),
                f"{r['d_cfi']:+.3f}" if pd.notna(r['d_cfi']) else "—",
                f"{r['d_rmsea']:+.3f}" if pd.notna(r['d_rmsea']) else "—",
                r["verdict"] if pd.notna(r["verdict"]) else "—",
            ]
            for i, txt in enumerate(txts):
                cells[i].text = ""
                pp = cells[i].paragraphs[0]
                rr = pp.add_run(txt); rr.font.size = Pt(7.5)

    doc.add_paragraph()


# -------- Interpretation note ------------------------------------------
p = doc.add_paragraph()
r = p.add_run("Interpretation and Implications for the Thesis")
r.bold = True; r.font.size = Pt(12)

interp = [
    "Configural invariance is met where models converge — the same single-factor "
    "structure exists for both groups and both waves on every composite. This is the "
    "minimum precondition for treating the constructs as comparable.",
    "",
    "Metric invariance (equal factor loadings) holds for some composites and not others. "
    "Where it holds, it means the items load on the underlying construct with the same "
    "weight across groups/waves — i.e., a one-unit change on Item X means the same thing "
    "in terms of the latent construct for both Estonians and Russians. Where metric "
    "invariance fails, group-level correlations involving the construct (regression "
    "coefficients, interactions) are not strictly comparable across groups.",
    "",
    "Scalar invariance (equal loadings AND intercepts) is the standard required for "
    "directly comparing latent means — i.e., for treating Cohen's d on the composite "
    "as reflecting only the underlying construct difference rather than measurement-"
    "level intercept shifts. Scalar invariance fails for most comparisons in this "
    "dataset, which is consistent with cross-cultural and pre/post-shock survey "
    "research generally.",
    "",
    "The pragmatic interpretation: the d values reported in the canonical TSV "
    "(reports/Effect_Sizes_with_CI.docx) are useful descriptions of group-level "
    "differences in the observed composite scores, but readers should be aware that "
    "some portion of those differences may reflect measurement-level variation in how "
    "items are calibrated rather than substantive attitude differences. This is "
    "explicitly noted in the methods chapter.",
    "",
    "Where partial invariance was found, future analyses could relax constraints on "
    "specific items (using lavaan's partial-invariance procedures) to recover "
    "comparable latent means while acknowledging item-specific measurement "
    "differences.",
]

for line in interp:
    pp = doc.add_paragraph()
    if line:
        rr = pp.add_run(line); rr.font.size = Pt(10)

# Footer note
p = doc.add_paragraph()
r = p.add_run(
    "Notes. Models estimated using lavaan 0.6.20 with the MLR estimator and full-"
    "information maximum likelihood (FIML) for missing data. Reverse-coded items "
    "were inverted prior to estimation so all loadings face the same direction. "
    "SD: Primary Out-group is excluded from multi-group invariance tests because "
    "Estonians and Russians answer different items (Q57/Q58/Q59 _1 vs _2). SD: "
    "General Out-group is excluded from longitudinal invariance because the item "
    "set differs between waves (3 items in 2020, 6 in 2023). "
    "Belief in Inevitable Conflict uses the Option B reverse-coding (Q63_1 and Q63_2 "
    "reversed) so the composite direction matches the documentation."
)
r.font.size = Pt(8); r.italic = True


out = ROOT / "reports" / "Measurement_Invariance.docx"
doc.save(out)
print(f"Saved: {out}")
