"""
Extends the 2 × 2 factorial ANOVA in script 94 to include both single-item
variables (Group ID Patterns and Territorial Attachment) alongside the eight
composite outcomes.

Single-item variables are documented in CLAUDE.md:
  Group ID Patterns: Q66 (2023) / K6X4 (2020), 1–5 scale, higher = more
    out-group identity. DK code 9 (Q66) and 6 (K6X4).
  Territorial Attachment: Q67_1 (2023) / K6X5_1 (2020), 1–4 scale,
    lower = stronger attachment. Inverted (5 − raw) for display so
    higher = stronger attachment.

Output:
  code/_anova_2x2_all_variables.tsv (10 variables)
  reports/ANOVA_2x2_APA_all_variables.docx
"""

import math
from pathlib import Path

import pandas as pd
import pyreadstat
import statsmodels.api as sm
import statsmodels.formula.api as smf
from scipy import stats
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent.parent

df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)
df20, _ = pyreadstat.read_sav(
    str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"), encoding="latin1"
)
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)


def composite(df, items, reverse_items=None, scale_max=None, dk_code=9):
    sub = df[items].apply(pd.to_numeric, errors="coerce").copy()
    sub = sub.where(sub != dk_code)
    if reverse_items:
        for it in reverse_items:
            sub[it] = (scale_max + 1) - sub[it]
    return sub.mean(axis=1, skipna=True)


def single_item(df, var, dk_code=9):
    s = pd.to_numeric(df[var], errors="coerce")
    return s.where(s != dk_code).astype(float)


def build_long(spec):
    """Build long-format DataFrame for a composite or single-item variable."""
    pieces = []
    for df, year, item_spec, rev_spec, dk in [
        (df23, 2023, spec.get("items_23"), spec.get("reverse_items_23"),
         spec.get("dk_23", 9)),
        (df20, 2020, spec.get("items_20"), spec.get("reverse_items_20"),
         spec.get("dk_20", 9)),
    ]:
        if spec.get("single_item"):
            # Single-item path
            for grp_code in (0, 1):
                sub_df = df[df["ethnicity_binary"] == grp_code]
                vals = single_item(sub_df, item_spec, dk_code=dk)
                if spec["inverted"]:
                    vals = (spec["scale_max"] + 1) - vals
                pieces.append(pd.DataFrame({
                    "group": "Estonian" if grp_code == 0 else "Russian",
                    "year": year, "value": vals.values,
                }))
        elif isinstance(item_spec, dict):
            # Group-specific composite items
            for grp_key, grp_code in (("E", 0), ("R", 1)):
                sub_df = df[df["ethnicity_binary"] == grp_code]
                vals = composite(sub_df, item_spec[grp_key],
                                 rev_spec[grp_key] if rev_spec else None,
                                 spec["scale_max"])
                if spec["inverted"]:
                    vals = (spec["scale_max"] + 1) - vals
                pieces.append(pd.DataFrame({
                    "group": "Estonian" if grp_code == 0 else "Russian",
                    "year": year, "value": vals.values,
                }))
        else:
            vals = composite(df, item_spec, rev_spec, spec["scale_max"])
            if spec["inverted"]:
                vals = (spec["scale_max"] + 1) - vals
            df_year = df[["ethnicity_binary"]].copy()
            df_year["value"] = vals.values
            df_year["group"] = df_year["ethnicity_binary"].map({0: "Estonian", 1: "Russian"})
            df_year["year"] = year
            pieces.append(df_year[["group", "year", "value"]])
    return pd.concat(pieces, ignore_index=True).dropna(subset=["value"]).reset_index(drop=True)


SPECS = [
    {"name": "Superordinate Identity", "scale_max": 4, "inverted": True,
     "items_23": ["Q67_2", "Q67_4", "Q67_5"], "reverse_items_23": ["Q67_4"],
     "items_20": ["K6X5_2", "K6X5_3", "K6X5_4"], "reverse_items_20": ["K6X5_3"]},
    {"name": "SD: Primary Out-group", "scale_max": 5, "inverted": False,
     "items_23": {"E": ["Q57_1","Q58_1","Q59_1"], "R": ["Q57_2","Q58_2","Q59_2"]},
     "items_20": {"E": ["K4X7_1","K4X8_1","K4X9_1"], "R": ["K4X7_2","K4X8_2","K4X9_2"]}},
    {"name": "SD: General Out-group", "scale_max": 5, "inverted": False,
     "items_23": ["Q57_4","Q57_5","Q58_4","Q58_5","Q59_4","Q59_5"],
     "items_20": ["K4X7_3","K4X8_3","K4X9_3"]},
    {"name": "Comparative Opportunity Assessment", "scale_max": 5, "inverted": True,
     "items_23": [f"Q44_{i}" for i in range(1, 13)],
     "items_20": [f"K3X1_{i}" for i in range(1, 13)]},
    {"name": "Belief in Inevitable Conflict", "scale_max": 4, "inverted": False,
     "items_23": ["Q63_1","Q63_2","Q63_3","Q63_4"], "reverse_items_23": ["Q63_1","Q63_2"],
     "items_20": ["K6X1_1","K6X1_2","K6X1_3","K6X1_4"], "reverse_items_20": ["K6X1_1","K6X1_2"]},
    {"name": "Minority Inclusion Support", "scale_max": 4, "inverted": True,
     "items_23": ["Q68_1","Q68_2","Q68_3"],
     "items_20": ["K6X6_1","K6X6_2","K6X6_3"]},
    {"name": "Contact: Estonian Speakers", "scale_max": 5, "inverted": True,
     "items_23": [f"Q51_{i}" for i in range(1, 7)],
     "items_20": [f"K4X1_{i}" for i in range(1, 7)]},
    {"name": "Contact: Russian Speakers", "scale_max": 5, "inverted": True,
     "items_23": [f"Q52_{i}" for i in range(1, 7)],
     "items_20": [f"K4X2_{i}" for i in range(1, 7)]},
    # Single-item variables
    {"name": "Group ID Patterns", "scale_max": 5, "inverted": False,
     "single_item": True,
     "items_23": "Q66", "items_20": "K6X4",
     "dk_23": 9, "dk_20": 6,
     "note": "Single item. 1-5 scale, higher = more out-group identity. "
              "K6X4 (2020) uses DK code 6; Q66 (2023) uses DK code 9."},
    {"name": "Territorial Attachment", "scale_max": 4, "inverted": True,
     "single_item": True,
     "items_23": "Q67_1", "items_20": "K6X5_1",
     "dk_23": 9, "dk_20": 9,
     "note": "Single item. 1-4 scale (raw: lower = stronger attachment). "
              "Inverted (5 - raw) so higher = stronger attachment."},
]


def run_anova(df_long):
    df_long = df_long.copy()
    df_long["group_c"] = df_long["group"].astype("category")
    df_long["year_c"]  = df_long["year"].astype("category")
    m = smf.ols("value ~ C(group_c, Sum) * C(year_c, Sum)", data=df_long).fit()
    tbl = sm.stats.anova_lm(m, typ=3)
    ss_resid = tbl.loc["Residual", "sum_sq"]
    out = {}
    for label, term in [
        ("Ethnicity",      "C(group_c, Sum)"),
        ("Year",           "C(year_c, Sum)"),
        ("Ethnicity × Year","C(group_c, Sum):C(year_c, Sum)"),
    ]:
        ss = tbl.loc[term, "sum_sq"]
        out[label] = {"F": tbl.loc[term, "F"], "p": tbl.loc[term, "PR(>F)"],
                      "eta2p": ss / (ss + ss_resid)}
    out["__N__"] = int(m.nobs)
    return out


def cell_stats(df_long):
    out = {}
    for g in ("Estonian", "Russian"):
        for y in (2020, 2023):
            s = df_long[(df_long["group"] == g) & (df_long["year"] == y)]["value"].dropna()
            out[(g, y)] = (s.mean(), s.std(ddof=1), len(s))
    return out


def levene_within(df_long, group):
    s20 = df_long[(df_long["group"] == group) & (df_long["year"] == 2020)]["value"].dropna()
    s23 = df_long[(df_long["group"] == group) & (df_long["year"] == 2023)]["value"].dropna()
    F, p = stats.levene(s20, s23, center="median")
    return F, p, s20.var(ddof=1), s23.var(ddof=1)


results = []
for spec in SPECS:
    long = build_long(spec)
    anova = run_anova(long)
    cells = cell_stats(long)
    lev_E = levene_within(long, "Estonian")
    lev_R = levene_within(long, "Russian")
    results.append({"spec": spec, "anova": anova, "cells": cells,
                    "lev_E": lev_E, "lev_R": lev_R, "N_total": anova["__N__"]})

# TSV
tsv_rows = []
for r in results:
    name = r["spec"]["name"]
    cE20 = r["cells"][("Estonian", 2020)]; cE23 = r["cells"][("Estonian", 2023)]
    cR20 = r["cells"][("Russian", 2020)];  cR23 = r["cells"][("Russian", 2023)]
    tsv_rows.append({
        "variable": name, "N_total": r["N_total"],
        "Est_2020_M": cE20[0], "Est_2020_SD": cE20[1], "Est_2020_N": cE20[2],
        "Est_2023_M": cE23[0], "Est_2023_SD": cE23[1], "Est_2023_N": cE23[2],
        "Rus_2020_M": cR20[0], "Rus_2020_SD": cR20[1], "Rus_2020_N": cR20[2],
        "Rus_2023_M": cR23[0], "Rus_2023_SD": cR23[1], "Rus_2023_N": cR23[2],
        "F_ethnicity": r["anova"]["Ethnicity"]["F"], "p_ethnicity": r["anova"]["Ethnicity"]["p"],
        "eta2p_ethnicity": r["anova"]["Ethnicity"]["eta2p"],
        "F_year": r["anova"]["Year"]["F"], "p_year": r["anova"]["Year"]["p"],
        "eta2p_year": r["anova"]["Year"]["eta2p"],
        "F_interaction": r["anova"]["Ethnicity × Year"]["F"],
        "p_interaction": r["anova"]["Ethnicity × Year"]["p"],
        "eta2p_interaction": r["anova"]["Ethnicity × Year"]["eta2p"],
        "Levene_Est_F": r["lev_E"][0], "Levene_Est_p": r["lev_E"][1],
        "Levene_Est_var_2020": r["lev_E"][2], "Levene_Est_var_2023": r["lev_E"][3],
        "Levene_Rus_F": r["lev_R"][0], "Levene_Rus_p": r["lev_R"][1],
        "Levene_Rus_var_2020": r["lev_R"][2], "Levene_Rus_var_2023": r["lev_R"][3],
    })
tsv = ROOT / "code" / "_anova_2x2_all_variables.tsv"
pd.DataFrame(tsv_rows).to_csv(tsv, sep="\t", index=False, float_format="%.4f")
print(f"Saved TSV: {tsv}\n")

# Console summary for the two new rows
for r in results[-2:]:
    name = r["spec"]["name"]
    print(f"=== {name}  (N = {r['N_total']}) ===")
    for g in ("Estonian", "Russian"):
        for y in (2020, 2023):
            M, SD, N = r["cells"][(g, y)]
            print(f"  {g:<9} {y}: M = {M:.3f}, SD = {SD:.3f}, n = {N}")
    for eff in ("Ethnicity", "Year", "Ethnicity × Year"):
        e = r["anova"][eff]
        print(f"  {eff:<18}: F = {e['F']:.2f}, p = {e['p']:.4g}, η²ₚ = {e['eta2p']:.3f}")
    fE, pE, vE20, vE23 = r["lev_E"]; fR, pR, vR20, vR23 = r["lev_R"]
    print(f"  Levene Est: F = {fE:.2f}, p = {pE:.4g}  ({vE20:.3f} → {vE23:.3f})")
    print(f"  Levene Rus: F = {fR:.2f}, p = {pR:.4g}  ({vR20:.3f} → {vR23:.3f})")
    print()


# ============================================================================
#                       APA WORD DOC (same as script 94 but with 10 rows)
# ============================================================================
es = pd.read_csv(ROOT / "code" / "_effect_sizes.tsv", sep="\t")
es_within = es[es["table"] == "within"].copy()
es_between = es[es["table"] == "between"].copy()

ANOVA_TO_ES = {
    "Superordinate Identity":             "Superordinate Identity",
    "SD: Primary Out-group":              "SD: Primary Out-group",
    "SD: General Out-group":              "SD: General Out-group",
    "Comparative Opportunity Assessment": "Comparative Opportunity Assessment",
    "Belief in Inevitable Conflict":      "Belief in Inevitable Conflict",
    "Minority Inclusion Support":         "Minority Support Inclusion",
    "Contact: Estonian Speakers":         "Contact: Estonian Speakers",
    "Contact: Russian Speakers":          "Contact: Russian Speakers",
    "Group ID Patterns":                  "Group ID Patterns",
    "Territorial Attachment":             "Territorial Attachment",
}


def stars(p):
    if pd.isna(p): return ""
    if p < .001: return "***"
    if p < .01:  return "**"
    if p < .05:  return "*"
    if p < .10:  return "⁺"
    return ""


def fmt_p_apa(p):
    if pd.isna(p): return "—"
    if p < .001:   return "< .001"
    return f"{p:.3f}".lstrip("0")


def fmt_F(F): return "—" if pd.isna(F) else f"{F:.2f}"


def fmt_eta(eta):
    if pd.isna(eta): return "—"
    return f"{eta:.3f}".lstrip("0")


def within_p(varname, group):
    es_name = ANOVA_TO_ES[varname]
    row = es_within[(es_within["variable"] == es_name) &
                    (es_within["comparison"] == group)]
    return float(row["p"].iloc[0]) if len(row) else None


def superscript(year, group_p):
    if group_p is None: return ""
    if group_p < .05: return "a" if year == 2020 else "b"
    if group_p < .10: return "⁺"
    return ""


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for side in ("top", "left", "bottom", "right"):
        spec_ = kwargs.get(side)
        elem = tcBorders.find(qn(f"w:{side}"))
        if elem is None:
            elem = OxmlElement(f"w:{side}")
            tcBorders.append(elem)
        if spec_ is None:
            elem.set(qn("w:val"), "nil")
        else:
            elem.set(qn("w:val"), spec_.get("val", "single"))
            elem.set(qn("w:sz"),  str(spec_.get("sz", 8)))
            elem.set(qn("w:color"), spec_.get("color", "auto"))


def apply_apa_borders(table, header_rows=1):
    n_rows = len(table.rows)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            top = bottom = None
            if r_idx == 0:
                top = {"sz": 12, "val": "single", "color": "000000"}
            if r_idx == header_rows - 1:
                bottom = {"sz": 6, "val": "single", "color": "000000"}
            if r_idx == n_rows - 1:
                bottom = {"sz": 12, "val": "single", "color": "000000"}
            _set_cell_border(cell, top=top, bottom=bottom, left=None, right=None)


def remove_table_style(table):
    tblPr = table._tbl.find(qn("w:tblPr"))
    if tblPr is not None:
        tblStyle = tblPr.find(qn("w:tblStyle"))
        if tblStyle is not None:
            tblPr.remove(tblStyle)
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.find(qn("w:shd"))
            if shd is not None:
                tcPr.remove(shd)


doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(2.0)
sec.orientation = 1
sec.page_width  = Cm(29.7); sec.page_height = Cm(21.0)
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)


def set_cell(cell, text, *, bold=False, italic=False, size=10,
              align=WD_ALIGN_PARAGRAPH.LEFT, vcenter=False):
    cell.text = ""
    para = cell.paragraphs[0]; para.alignment = align
    if vcenter: cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        r = para.add_run(line)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic
        if i < len(lines) - 1: r.add_break()


def set_cell_italics(cell, segments, *, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]; para.alignment = align
    for tup in segments:
        if len(tup) == 3:
            text, bold, italic = tup; sup = False
        else:
            text, bold, italic, sup = tup
        r = para.add_run(text)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic
        if sup: r.font.superscript = True


def merge_v(table, rows, col):
    cells = [table.rows[r].cells[col] for r in rows]
    merged = cells[0]
    for c in cells[1:]: merged = merged.merge(c)
    return merged


def merge_h(table, row, cols):
    cells = [table.rows[row].cells[c] for c in cols]
    merged = cells[0]
    for c in cells[1:]: merged = merged.merge(c)
    return merged


def H_table_number(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.name = "Times New Roman"


def H_table_title(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.size = Pt(11); r.font.name = "Times New Roman"


def table_note(segments, size=9):
    p = doc.add_paragraph()
    for text, bold, italic in segments:
        r = p.add_run(text)
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.bold = bold; r.italic = italic


p = doc.add_paragraph()
r = p.add_run("2 × 2 Factorial ANOVA Results — Ethnicity × Year (All 10 variables)")
r.bold = True; r.font.size = Pt(16); r.font.name = "Times New Roman"

p = doc.add_paragraph()
r = p.add_run(
    "ANOVA tables for all ten outcome variables: six multi-item attitudinal "
    "composites, two contact composites (behavioral indicators), and two "
    "single-item variables (Group ID Patterns and Territorial Attachment). "
    "All inverted variables have been transformed so higher = more of the "
    "construct (more belonging, more support, more contact, stronger "
    "attachment). APA-7 formatting throughout."
)
r.italic = True; r.font.size = Pt(10); r.font.name = "Times New Roman"
doc.add_paragraph()


# --- TABLE 1 --------------------------------------------------------------
df_an = pd.DataFrame(tsv_rows)

H_table_number("Table 1")
H_table_title(
    "Means and standard deviations of all outcome measures and the results "
    "of 2 × 2 univariate analyses of variance (Ethnicity × Year)."
)

ncols = 12
nrows = 2 + len(df_an) * 2
tbl1 = doc.add_table(rows=nrows, cols=ncols)
super_header = [
    ("Outcomes", 0, 1), ("Wave", 1, 1),
    ("Estonian", 2, 2), ("Russian", 4, 2),
    ("Ethnicity", 6, 2), ("Year", 8, 2), ("Ethnicity × Year", 10, 2),
]
for label, start, span in super_header:
    if span > 1:
        merged = merge_h(tbl1, 0, range(start, start + span))
        set_cell(merged, label, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        set_cell(tbl1.rows[0].cells[start], label, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
sub_header = ["", "", "M", "SD", "M", "SD", "F", "p", "F", "p", "F", "p"]
for i, label in enumerate(sub_header):
    if label in ("M", "SD", "F", "p"):
        set_cell_italics(tbl1.rows[1].cells[i], [(label, True, True)],
                          size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    elif label:
        set_cell(tbl1.rows[1].cells[i], label, bold=True, size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
merge_v(tbl1, [0, 1], 0)
set_cell(tbl1.rows[0].cells[0], "Outcomes", bold=True, size=10)
merge_v(tbl1, [0, 1], 1)
set_cell(tbl1.rows[0].cells[1], "Wave", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

row_idx = 2
for _, r in df_an.iterrows():
    varname = r["variable"]
    pE = within_p(varname, "Estonian"); pR = within_p(varname, "Russian")
    merge_v(tbl1, [row_idx, row_idx + 1], 0)
    set_cell(tbl1.rows[row_idx].cells[0], varname, size=10)
    for col_start in (6, 8, 10):
        F_val = {6: r["F_ethnicity"], 8: r["F_year"], 10: r["F_interaction"]}[col_start]
        p_val = {6: r["p_ethnicity"], 8: r["p_year"], 10: r["p_interaction"]}[col_start]
        eta_val = {6: r["eta2p_ethnicity"], 8: r["eta2p_year"],
                   10: r["eta2p_interaction"]}[col_start]
        merge_v(tbl1, [row_idx, row_idx + 1], col_start)
        merge_v(tbl1, [row_idx, row_idx + 1], col_start + 1)
        set_cell(tbl1.rows[row_idx].cells[col_start], fmt_F(F_val), size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(tbl1.rows[row_idx].cells[col_start + 1],
                  f"{fmt_p_apa(p_val)}\nη²ₚ = {fmt_eta(eta_val)}",
                  size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    for offset, year in enumerate([2020, 2023]):
        rr = row_idx + offset
        set_cell(tbl1.rows[rr].cells[1], str(year), size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        m_est = r[f"Est_{year}_M"]; sd_est = r[f"Est_{year}_SD"]
        sup_est = superscript(year, pE)
        set_cell_italics(tbl1.rows[rr].cells[2],
                          [(f"{m_est:.2f}", False, False, False)]
                          + ([(sup_est, False, False, True)] if sup_est else []),
                          size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(tbl1.rows[rr].cells[3], f"{sd_est:.2f}", size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        m_rus = r[f"Rus_{year}_M"]; sd_rus = r[f"Rus_{year}_SD"]
        sup_rus = superscript(year, pR)
        set_cell_italics(tbl1.rows[rr].cells[4],
                          [(f"{m_rus:.2f}", False, False, False)]
                          + ([(sup_rus, False, False, True)] if sup_rus else []),
                          size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(tbl1.rows[rr].cells[5], f"{sd_rus:.2f}", size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    row_idx += 2

remove_table_style(tbl1); apply_apa_borders(tbl1, header_rows=2)
table_note([
    ("Note. ", False, True),
    ("Superscripts (a, b) within a group's column indicate within-group "
     "wave means differ at ", False, False),
    ("p", False, True), (" < .05; ⁺ indicates ", False, False),
    ("p", False, True), (" < .10. Group ID Patterns and Territorial Attachment "
                          "are single items. Group ID Patterns ranges 1-5, raw "
                          "(higher = more out-group identity). Territorial "
                          "Attachment ranges 1-4 (inverted so higher = stronger "
                          "attachment).  Contact composites inverted so higher "
                          "= more frequent contact. η²ₚ = partial eta-squared "
                          "(small ≈ .01, medium ≈ .06, large ≈ .14).", False, False),
])

doc.add_paragraph(); doc.add_page_break()


# --- TABLE 2 — d ---------------------------------------------------------
H_table_number("Table 2")
H_table_title(
    "Simple-effect Cohen's d with 95% confidence intervals decomposing the "
    "2 × 2 design into between-group (within wave) and within-group "
    "(across wave) comparisons."
)

header2 = ["Outcomes",
           "Between-group 2020\n(Estonian − Russian)",
           "Between-group 2023\n(Estonian − Russian)",
           "Estonian respondents\n(2023 − 2020)",
           "Russian respondents\n(2023 − 2020)"]
tbl2 = doc.add_table(rows=1, cols=len(header2))
for i, h in enumerate(header2):
    set_cell(tbl2.rows[0].cells[i], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)


def d_lookup(t, v, c):
    sub = es_between if t == "between" else es_within
    row = sub[(sub["variable"] == v) & (sub["comparison"] == c)]
    if row.empty: return None
    r = row.iloc[0]
    return float(r["d"]), float(r["CI_low"]), float(r["CI_high"]), float(r["p"])


def fmt_d(d, ll, ul, p):
    return f"{d:+.3f}{stars(p)}\n[{ll:+.3f}, {ul:+.3f}]"


for _, r in df_an.iterrows():
    varname = r["variable"]; es_name = ANOVA_TO_ES[varname]
    row = tbl2.add_row().cells
    set_cell(row[0], varname, size=10)
    for i, (t, c) in enumerate([("between", "2020"), ("between", "2023"),
                                  ("within", "Estonian"), ("within", "Russian")], start=1):
        res = d_lookup(t, es_name, c)
        if res:
            set_cell(row[i], fmt_d(*res), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            set_cell(row[i], "—", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

remove_table_style(tbl2); apply_apa_borders(tbl2)
table_note([
    ("Note. ", False, True),
    ("d", False, True), (" = Cohen's d with root-mean-square SD denominator. "
                          "Significance: *** ", False, False),
    ("p", False, True), (" < .001, ** ", False, False),
    ("p", False, True), (" < .01, * ", False, False),
    ("p", False, True), (" < .05, ⁺ ", False, False),
    ("p", False, True), (" < .10.", False, False),
])

doc.add_paragraph(); doc.add_page_break()


# --- TABLE 3 — Levene's --------------------------------------------------
H_table_number("Table 3")
H_table_title(
    "Levene's tests (Brown–Forsythe variant) of variance homogeneity across "
    "waves, conducted separately within each ethnic group."
)

header3 = ["Outcomes",
           "Estonian respondents\nF, p, var(2020) → var(2023)",
           "Russian respondents\nF, p, var(2020) → var(2023)"]
tbl3 = doc.add_table(rows=1, cols=len(header3))
for i, h in enumerate(header3):
    set_cell(tbl3.rows[0].cells[i], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)


def fmt_lev(F, p, v20, v23):
    arrow = "↑" if v23 > v20 else "↓"
    return f"F = {F:.2f}, p = {fmt_p_apa(p)}{stars(p)}\n{v20:.3f} → {v23:.3f} {arrow}"


for _, r in df_an.iterrows():
    row = tbl3.add_row().cells
    set_cell(row[0], r["variable"], size=10)
    set_cell(row[1],
              fmt_lev(r["Levene_Est_F"], r["Levene_Est_p"],
                      r["Levene_Est_var_2020"], r["Levene_Est_var_2023"]),
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row[2],
              fmt_lev(r["Levene_Rus_F"], r["Levene_Rus_p"],
                      r["Levene_Rus_var_2020"], r["Levene_Rus_var_2023"]),
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

remove_table_style(tbl3); apply_apa_borders(tbl3)
table_note([
    ("Note. ", False, True),
    ("Median-centered (Brown–Forsythe) Levene's test. var(2020) → var(2023) "
     "shows within-group variance at each wave. ↑ = increased dispersion "
     "(polarization); ↓ = decreased dispersion (convergence). *** ", False, False),
    ("p", False, True), (" < .001, ** ", False, False),
    ("p", False, True), (" < .01, * ", False, False),
    ("p", False, True), (" < .05, ⁺ ", False, False),
    ("p", False, True), (" < .10.", False, False),
])

out = ROOT / "reports" / "ANOVA_2x2_APA_all_variables.docx"
doc.save(out)
print(f"Saved: {out}")
