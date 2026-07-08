"""
Full test-statistics table for the 8 composite variables.

Reads canonical M/SD/N/d/CI/p from code/_effect_sizes.tsv and derives Welch's
t and df from M/SD/N (no data re-load required):

    SE  = sqrt(SD1^2/N1 + SD2^2/N2)
    t   = (M1 - M2) / SE
    df  = (SD1^2/N1 + SD2^2/N2)^2
          / [ (SD1^2/N1)^2 / (N1 - 1) + (SD2^2/N2)^2 / (N2 - 1) ]

Emits:
    reports/Test_Statistics_Table.docx  (two formatted tables: between, within)
    code/_test_statistics.tsv           (machine-readable copy with t and df added)
"""

import csv
import math
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

ROOT = Path(__file__).parent.parent
TSV  = Path(__file__).parent / "_effect_sizes.tsv"
OUT_DOCX = ROOT / "reports" / "Test_Statistics_Table.docx"
OUT_TSV  = Path(__file__).parent / "_test_statistics.tsv"

COMPOSITES = [
    "Superordinate Identity",
    "SD: Primary Out-group",
    "SD: General Out-group",
    "Comparative Opportunity Assessment",
    "Belief in Inevitable Conflict",
    "Minority Support Inclusion",
    "Contact: Estonian Speakers",
    "Contact: Russian Speakers",
]

# ---------- Load canonical TSV --------------------------------------------
records = []
with open(TSV) as f:
    for row in csv.DictReader(f, delimiter="\t"):
        if row["variable"] not in COMPOSITES:
            continue
        records.append(row)

# ---------- Compute Welch t + df ------------------------------------------
def welch_t_df(M1, SD1, N1, M2, SD2, N2):
    v1 = SD1 * SD1 / N1
    v2 = SD2 * SD2 / N2
    se = math.sqrt(v1 + v2)
    t  = (M1 - M2) / se
    num = (v1 + v2) ** 2
    den = (v1 ** 2) / (N1 - 1) + (v2 ** 2) / (N2 - 1)
    df  = num / den
    return t, df

def stars(p):
    p = float(p)
    if p < 0.001: return "***"
    if p < 0.01:  return "**"
    if p < 0.05:  return "*"
    return "ns"

def fmt_p(p):
    p = float(p)
    if p < 0.001:
        return "<.001"
    return f"{p:.3f}".lstrip("0")

enriched = []
for r in records:
    M1, SD1, N1 = float(r["M1"]), float(r["SD1"]), int(r["N1"])
    M2, SD2, N2 = float(r["M2"]), float(r["SD2"]), int(r["N2"])
    t, df = welch_t_df(M1, SD1, N1, M2, SD2, N2)
    enriched.append({
        **r,
        "t":  t,
        "df": df,
        "stars": stars(r["p"]),
    })

# ---------- Write enriched TSV --------------------------------------------
with open(OUT_TSV, "w", newline="") as f:
    fieldnames = list(records[0].keys()) + ["t", "df", "stars"]
    w = csv.DictWriter(f, fieldnames=fieldnames, delimiter="\t")
    w.writeheader()
    for row in enriched:
        out = dict(row)
        out["t"]  = f"{row['t']:.3f}"
        out["df"] = f"{row['df']:.1f}"
        w.writerow(out)
print(f"Wrote: {OUT_TSV}")

# ---------- Word doc ------------------------------------------------------
doc = Document()

# Page setup — landscape, narrow margins (lots of columns)
section = doc.sections[0]
section.page_height, section.page_width = section.page_width, section.page_height
section.left_margin   = Cm(1.5)
section.right_margin  = Cm(1.5)
section.top_margin    = Cm(1.5)
section.bottom_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)

# Title
h = doc.add_heading("Test Statistics for Composite Variables", level=1)
for run in h.runs:
    run.font.size = Pt(16)

p = doc.add_paragraph()
r = p.add_run(
    "Welch's independent-samples t-test for every between-group (Estonian vs. "
    "Russian) and within-group (2020 vs. 2023) comparison on the eight "
    "multi-item composites. Pairwise deletion; Cohen's d uses the root-mean-"
    "square SD denominator. CIs are 95%."
)
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

# ----- Helpers -----
def add_table_header(table, headers):
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        para = hdr[i].paragraphs[0]
        run = para.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def add_data_row(table, cells, align_left_idx=(0, 1)):
    row = table.add_row().cells
    for i, val in enumerate(cells):
        row[i].text = ""
        para = row[i].paragraphs[0]
        run = para.add_run(str(val))
        run.font.size = Pt(8.5)
        para.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT if i in align_left_idx
            else WD_ALIGN_PARAGRAPH.CENTER
        )
        row[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def make_section(title, subtitle, group_label, rows):
    h = doc.add_heading(title, level=2)
    for run in h.runs:
        run.font.size = Pt(13)
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    headers = [
        "Variable", group_label,
        "M₁", "SD₁", "n₁",
        "M₂", "SD₂", "n₂",
        "t", "df", "p",
        "Cohen's d", "95% CI",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    add_table_header(table, headers)

    for r_ in rows:
        d = float(r_["d"])
        ci = f"[{float(r_['CI_low']):+.2f}, {float(r_['CI_high']):+.2f}]"
        p_str = fmt_p(r_["p"]) + " " + r_["stars"]
        add_data_row(table, [
            r_["variable"], r_["comparison"],
            f"{float(r_['M1']):.2f}", f"{float(r_['SD1']):.2f}", r_["N1"],
            f"{float(r_['M2']):.2f}", f"{float(r_['SD2']):.2f}", r_["N2"],
            f"{r_['t']:+.2f}", f"{r_['df']:.1f}", p_str,
            f"{d:+.2f}", ci,
        ])

# ----- Between-group (M1 = Estonian, M2 = Russian) -----
between = [r for r in enriched if r["table"] == "between"
           and r["variable"] in COMPOSITES]
# Order by COMPOSITES, then year
between.sort(key=lambda r: (COMPOSITES.index(r["variable"]), r["comparison"]))
make_section(
    "Between-Group Comparisons",
    "Estonian (group 1) vs. Russian (group 2) respondents within each survey wave.",
    "Year",
    between,
)

doc.add_paragraph()

# ----- Within-group (M1 = 2020, M2 = 2023) -----
within = [r for r in enriched if r["table"] == "within"
          and r["variable"] in COMPOSITES]
within.sort(key=lambda r: (COMPOSITES.index(r["variable"]), r["comparison"]))
make_section(
    "Within-Group Comparisons (2020 → 2023)",
    "Group means in 2020 (group 1) vs. 2023 (group 2) for Estonian and Russian respondents separately.",
    "Group",
    within,
)

# ----- Footnote -----
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run(
    "Note. Welch's independent-samples t-test (separate variances; df fractional). "
    "Significance: *** p < .001, ** p < .01, * p < .05, ns = not significant. "
    "Cohen's d uses the root-mean-square SD denominator: "
    "d = (M₁ − M₂) / √[(SD₁² + SD₂²)/2]. "
    "95% CI on d via Hedges & Olkin (1985) asymptotic SE."
)
r.font.size = Pt(8.5)
r.italic = True
r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.save(OUT_DOCX)
print(f"Wrote: {OUT_DOCX}")
