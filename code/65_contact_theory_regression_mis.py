"""
Contact-Theory Regression Check — Minority Inclusion Support
==============================================================
Mirror of script 63, but with Minority Inclusion Support as the DV instead
of SD: Primary Out-group. Same four cells (Russian / Estonian × 2020 / 2023),
same single-equation linear model framework.

Tests whether within-group contact with the primary out-group predicts
greater support for minority inclusion.

Direction handling:
  - Contact (Q51 / Q52, K4X1 / K4X2): inverted via (6 − raw) → higher = more contact
  - MIS (Q68 / K6X6): inverted via (5 − raw) → higher = more support for inclusion
  - Contact theory predicts POSITIVE slope: more contact → more support
  - Sign convention is the OPPOSITE of script 63 (SD regression) because
    SD has higher = more distance while inverted MIS has higher = more support.
    Both signs represent the same theoretical direction.

Outputs:
  Console:  coefficient, SE, 95% CI, t, p, R², N for each regression
  reports/Contact_Theory_Regression_MIS.docx — formatted Word doc
  viz/fig_contact_theory_regression_mis.jpg — 2×2 scatter+fit panels
"""

from pathlib import Path
import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, RGBColor

ROOT = Path(__file__).parent.parent

# ---------- Load data --------------------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy()

df20, _ = pyreadstat.read_sav(str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"),
                              encoding="latin1")
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy()


def composite_pairwise(df, items, dk=9, invert_max=None):
    """Pairwise-deletion composite. If invert_max is set, returns (invert_max + 1) − mean."""
    sub = df[items].apply(pd.to_numeric, errors="coerce").where(lambda x: x != dk)
    m = sub.mean(axis=1, skipna=True)
    return ((invert_max + 1) - m) if invert_max is not None else m


# ---------- Item codes ------------------------------------------------------
MIS_2023 = ["Q68_1", "Q68_2", "Q68_3"]
MIS_2020 = ["K6X6_1", "K6X6_2", "K6X6_3"]
CONTACT_RUS_2023 = [f"Q51_{i}" for i in range(1, 7)]   # Russians' contact with Estonian-speakers
CONTACT_EST_2023 = [f"Q52_{i}" for i in range(1, 7)]   # Estonians' contact with Russian-speakers
CONTACT_RUS_2020 = [f"K4X1_{i}" for i in range(1, 7)]
CONTACT_EST_2020 = [f"K4X2_{i}" for i in range(1, 7)]


def build_pair(df, eth_val, contact_items, year):
    sub = df[df["ethnicity_binary"] == eth_val]
    mis_items = MIS_2020 if year == 2020 else MIS_2023
    mis_inv     = composite_pairwise(sub, mis_items, invert_max=4)        # 1-4 scale, invert
    contact_inv = composite_pairwise(sub, contact_items, invert_max=5)    # 1-5 scale, invert
    return pd.DataFrame({"MIS": mis_inv, "Contact": contact_inv}).dropna()


CELLS = [
    ("Russian 2020",  build_pair(df20, 1, CONTACT_RUS_2020, 2020),
                      "Contact with Estonian-speakers"),
    ("Russian 2023",  build_pair(df23, 1, CONTACT_RUS_2023, 2023),
                      "Contact with Estonian-speakers"),
    ("Estonian 2020", build_pair(df20, 0, CONTACT_EST_2020, 2020),
                      "Contact with Russian-speakers"),
    ("Estonian 2023", build_pair(df23, 0, CONTACT_EST_2023, 2023),
                      "Contact with Russian-speakers"),
]


def stars(p):
    if p < 0.001: return "***"
    if p < 0.01:  return "**"
    if p < 0.05:  return "*"
    return "ns"


# ---------- Run regressions -------------------------------------------------
print("=" * 92)
print("CONTACT-THEORY REGRESSION CHECK — Minority Inclusion Support")
print("Model: MIS (inverted: higher = more support) ~ Contact (inverted: higher = more contact)")
print("Contact theory predicts POSITIVE slope: more contact → more support")
print("=" * 92)

results = []
for label, df_pair, contact_label in CELLS:
    X = sm.add_constant(df_pair["Contact"])
    y = df_pair["MIS"]
    model = sm.OLS(y, X).fit()
    b0, b1 = model.params["const"], model.params["Contact"]
    se1 = model.bse["Contact"]
    ci_lo, ci_hi = model.conf_int().loc["Contact"]
    t = model.tvalues["Contact"]
    p = model.pvalues["Contact"]
    r2 = model.rsquared
    n = int(model.nobs)
    results.append({
        "label": label, "contact_label": contact_label,
        "b0": b0, "b1": b1, "se1": se1,
        "ci_lo": ci_lo, "ci_hi": ci_hi,
        "t": t, "p": p, "r2": r2, "n": n,
        "model": model,
    })
    print(f"\n  {label} ({contact_label} → MIS)")
    print(f"    n = {n}")
    print(f"    Intercept (b0)  = {b0:+.3f}")
    print(f"    Slope (b1)      = {b1:+.4f}   SE = {se1:.4f}   "
          f"95% CI [{ci_lo:+.4f}, {ci_hi:+.4f}]")
    print(f"    t = {t:+.3f}   p = {p:.4g}   {stars(p)}")
    print(f"    R² = {r2:.4f}")


# ============================================================================
#  Word doc
# ============================================================================
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(2.0)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)


def H(text, size=14, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    if color: r.font.color.rgb = color


def body(text, italic=False, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size)
    if italic: r.italic = True


H("Contact-Theory Regression Check — Minority Inclusion Support",
  size=14, color=RGBColor(0x1f, 0x29, 0x37))

body(
    "Targeted bivariate regression of Minority Inclusion Support on frequency of contact "
    "with the primary out-group, run separately for each group × year cell. Companion to "
    "the SD: Primary Out-group regression (Contact_Theory_Regression.docx). Tests whether "
    "more contact predicts greater support for involving other nationalities in economic, "
    "governance, and civic life.",
    italic=True
)
doc.add_paragraph()

header = ["Group × Year", "n", "b₁ (slope)", "SE", "95% CI", "t", "p", "R²"]
table = doc.add_table(rows=1, cols=len(header))
table.style = "Light Grid Accent 1"
for i, h in enumerate(header):
    cell = table.rows[0].cells[i]; cell.text = ""
    rr = cell.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)


def fmt_p(p):
    if p < 0.001: return "< .001"
    return f"{p:.3f}".lstrip("0") if p < 1 else f"{p:.3f}"


def fmt_signed(x, dec=4):
    s = f"{x:+.{dec}f}"
    return s.replace("-", "−")


for r in results:
    cells = table.add_row().cells
    cells[0].text = r["label"]
    cells[1].text = str(r["n"])
    cells[2].text = f"{fmt_signed(r['b1'])} {stars(r['p'])}"
    cells[3].text = f"{r['se1']:.4f}"
    cells[4].text = f"[{fmt_signed(r['ci_lo'])}, {fmt_signed(r['ci_hi'])}]"
    cells[5].text = fmt_signed(r["t"], dec=2)
    cells[6].text = fmt_p(r["p"])
    cells[7].text = f"{r['r2']:.3f}"
    for c in cells:
        for para in c.paragraphs:
            for run in para.runs:
                run.font.size = Pt(9)


doc.add_paragraph()
body(
    "Note. Each row reports an OLS regression of the Minority Inclusion Support composite "
    "(items inverted via 5 − raw, so higher = more support) on the Contact composite "
    "(items inverted via 6 − raw, so higher = more frequent contact). For Russian "
    "respondents, contact is with Estonian-speakers (Q51 / K4X1 in 2023 / 2020); for "
    "Estonian respondents, contact is with Russian-speakers (Q52 / K4X2). Contact theory "
    "predicts a POSITIVE slope (more contact → more support). Listwise deletion: respondents "
    "missing on either composite are excluded. Cross-sectional design: associations are not "
    "causal. Significance: *** p < .001, ** p < .01, * p < .05, ns = not significant.",
    italic=True, size=8
)

out_doc = ROOT / "reports" / "Contact_Theory_Regression_MIS.docx"
doc.save(out_doc)
print(f"\nSaved Word doc: {out_doc}")


# ============================================================================
#  Figure — 2×2 scatter+fit panels
# ============================================================================
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size":   9,
    "axes.edgecolor": "#333",
    "axes.linewidth": 0.6,
})

RUS = "#d97706"; EST = "#2563eb"

fig, axes = plt.subplots(2, 2, figsize=(11.0, 9.0), dpi=300, sharex=True, sharey=True)

panel_config = [
    (results[0], CELLS[0][1], RUS, axes[0, 0]),
    (results[1], CELLS[1][1], RUS, axes[0, 1]),
    (results[2], CELLS[2][1], EST, axes[1, 0]),
    (results[3], CELLS[3][1], EST, axes[1, 1]),
]

for r, df_pair, color, ax in panel_config:
    rng = np.random.default_rng(42)
    jx = rng.uniform(-0.04, 0.04, size=len(df_pair))
    jy = rng.uniform(-0.04, 0.04, size=len(df_pair))
    ax.scatter(df_pair["Contact"] + jx, df_pair["MIS"] + jy,
               s=12, color=color, alpha=0.30, edgecolor="none", zorder=2)

    x_grid = np.linspace(1, 5, 50)
    y_pred = r["b0"] + r["b1"] * x_grid
    ax.plot(x_grid, y_pred, color="#111", linewidth=1.6, zorder=3)

    pred = r["model"].get_prediction(sm.add_constant(pd.Series(x_grid, name="Contact")))
    ci = pred.summary_frame(alpha=0.05)
    ax.fill_between(x_grid, ci["mean_ci_lower"], ci["mean_ci_upper"],
                    color="#111", alpha=0.10, zorder=1)

    ax.set_title(r["label"], fontsize=11, fontweight="bold", loc="left", pad=8)
    ax.text(0.02, 0.97,
            f"b₁ = {r['b1']:+.3f}  ({stars(r['p'])})\n"
            f"95% CI [{r['ci_lo']:+.3f}, {r['ci_hi']:+.3f}]\n"
            f"R² = {r['r2']:.3f}   n = {r['n']}",
            transform=ax.transAxes,
            ha="left", va="top",
            fontsize=8.5, color="#222",
            family="monospace",
            bbox=dict(boxstyle="round,pad=0.4", facecolor="white",
                      edgecolor="#888", linewidth=0.5, alpha=0.92))

    ax.set_xlim(0.7, 5.3)
    ax.set_ylim(0.7, 4.3)            # MIS is 1-4 scale, not 1-5
    ax.set_xticks([1, 2, 3, 4, 5])
    ax.set_yticks([1, 2, 3, 4])
    for s_ in ("top", "right"):
        ax.spines[s_].set_visible(False)
    ax.tick_params(axis="both", labelsize=8, color="#888")

axes[1, 0].set_xlabel("Contact composite (inverted: higher = more frequent contact)",
                       fontsize=9, color="#444")
axes[1, 1].set_xlabel("Contact composite (inverted: higher = more frequent contact)",
                       fontsize=9, color="#444")
axes[0, 0].set_ylabel("Minority Inclusion Support (inverted: higher = more support)",
                       fontsize=9, color="#444")
axes[1, 0].set_ylabel("Minority Inclusion Support (inverted: higher = more support)",
                       fontsize=9, color="#444")

fig.text(0.04, 0.965,
         "Contact-Theory Regression Check — Minority Inclusion Support ~ Contact",
         fontsize=13, fontweight="bold", ha="left")
fig.text(0.04, 0.940,
         "Top row: Russian respondents (Contact with Estonian-speakers, orange).  "
         "Bottom row: Estonian respondents (Contact with Russian-speakers, blue).",
         fontsize=9, color="#444")
fig.text(0.04, 0.922,
         "Contact theory predicts a POSITIVE slope: more contact → more support for minority inclusion.",
         fontsize=9, color="#444")

fig.text(0.04, 0.018,
         "Each dot = one respondent (jittered ±0.04 for legibility). Black line = OLS regression fit. Shaded band = 95% CI on the conditional mean.",
         fontsize=7, color="#555")
fig.text(0.04, 0.005,
         "Cross-sectional associations; no causal interpretation. EIM is a repeated cross-section, not a panel.",
         fontsize=7, color="#555")

plt.subplots_adjust(left=0.07, right=0.985, top=0.88, bottom=0.085,
                    wspace=0.10, hspace=0.30)

out_jpg = ROOT / "viz" / "fig_contact_theory_regression_mis.jpg"
plt.savefig(out_jpg, dpi=300, format="jpg",
            facecolor="white", pil_kwargs={"quality": 95})
print(f"Saved chart: {out_jpg}")
