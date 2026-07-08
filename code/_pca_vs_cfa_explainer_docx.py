"""
Generate a Word doc explaining the difference between PCA and multi-group CFA,
and the distinction between slope-comparison and mean-comparison (metric vs
scalar invariance). Reference document for the thesis methods chapter.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).parent.parent
out_path = ROOT / "reports" / "PCA_vs_MultiGroup_CFA.docx"

doc = Document()

# Page setup
section = doc.sections[0]
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)

# Default style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def add_title(text, size=18):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)


def add_heading(text, size=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)


def add_subheading(text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.italic = True
    r.font.size = Pt(size)


def add_para(text, italic=False, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    return p


def add_para_runs(runs):
    """runs = list of (text, bold, italic) tuples."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    for text, bold, italic in runs:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(11)
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.runs[0] if p.runs else p.add_run("")
    p.runs[0].font.size = Pt(11)
    # Replace the text — docx adds an empty run when style applied with no text
    p.text = text
    for run in p.runs:
        run.font.size = Pt(11)


def add_bullet_runs(runs):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    for text, bold, italic in runs:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(11)


def add_table(headers, rows, col_widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(val)
            r.font.size = Pt(10)
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)


# =====================================================================
# Title
# =====================================================================
add_title("PCA vs. Multi-Group CFA")
add_para(
    "A reference note on what each method models, what its score represents, "
    "and which kinds of group comparisons each one supports. Written for the "
    "EIM2 thesis methods chapter.",
    italic=True, size=10,
)

# =====================================================================
# Section 1: What they're modeling
# =====================================================================
add_heading("1. What they're modeling")

add_para_runs([
    ("PCA", True, False),
    (" is descriptive variance decomposition. It asks: ", False, False),
    ("what linear combination of these items captures the most variance in the observed data?", False, True),
    (" PC1 is whatever weighted average of the items has the highest variance. It makes no assumption about why the items are correlated.", False, False),
])

add_para_runs([
    ("CFA", True, False),
    (" is a confirmatory measurement model. It asks: ", False, False),
    ("if I assume one latent construct generates these items, with each item being the construct plus some item-specific noise, do the data fit?", False, True),
    (" CFA explicitly posits the measurement equation:", False, False),
])

# Equation paragraph (centered, italic)
eq = doc.add_paragraph()
eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
eq.paragraph_format.space_before = Pt(4)
eq.paragraph_format.space_after = Pt(4)
r = eq.add_run("x_ij = ν_i + λ_i · η_j + ε_ij")
r.italic = True
r.font.size = Pt(12)

add_para(
    "— each observed item value x is its intercept ν, plus a loading λ times the "
    "latent score η, plus measurement error ε. The latent isn't a weighted sum of "
    "the items; it's a separate variable inferred from them."
)

# =====================================================================
# Section 2: What the score represents
# =====================================================================
add_heading("2. What the \"score\" represents")

add_para_runs([
    ("PC1 score", True, False),
    (" = a weighted sum of the items, where weights maximize variance explained. "
     "The score has whatever variance falls out of the math (often normalized to "
     "mean 0, SD 1). It contains both true construct variance ", False, False),
    ("and", False, True),
    (" item-specific noise that happens to covary across items.", False, False),
])

add_para_runs([
    ("CFA factor score", True, False),
    (" = an estimate of the latent η, with item-specific noise partialed out. The factor score is ", False, False),
    ("disattenuated", False, True),
    (" — measurement error has been (partially) removed. This is why Cohen's d on CFA factor scores comes out larger than d on PC1 or composite means: the same true-score difference now sits on a less noisy denominator.", False, False),
])

add_para(
    "Concrete example from the EIM2 data: Belief in Inevitable Conflict, between-"
    "group 2023, went from d = +0.78 (composite mean) → +0.90 (PC1) → +1.18 (CFA). "
    "The composite includes all item-specific noise. PC1 sharpens by weighting "
    "items by their loadings but still includes residual noise. CFA strips the "
    "noise — the d climbs further."
)

# =====================================================================
# Section 3: Single-group vs. multi-group
# =====================================================================
add_heading("3. Single-group vs. multi-group")
add_para("This is the key difference for the EIM2 project.")

add_para_runs([
    ("PCA", True, False),
    (" is intrinsically single-group. Even when fit on a pooled sample of Estonians "
     "and Russians, the loadings are forced to be identical across groups by "
     "construction — there's no concept of \"let loadings differ by group.\" If "
     "Estonians and Russians actually use the items differently, PCA cannot detect "
     "that; it averages over them.", False, False),
])

add_para_runs([
    ("Multi-group CFA", True, False),
    (" estimates the same measurement model in each group simultaneously, with "
     "explicit choices about what to constrain equal across groups:", False, False),
])

add_table(
    headers=["Invariance level", "What's equal across groups",
             "What it lets you compare"],
    rows=[
        ["Configural", "Just the model structure (which items load on which factor)",
         "Existence of the construct in each group"],
        ["Metric", "Loadings (λ)",
         "Slopes, correlations, regression paths involving the latent"],
        ["Scalar", "Loadings AND intercepts (ν)",
         "Latent means (i.e., between-group Cohen's d)"],
        ["Strict", "Loadings, intercepts, AND residual variances",
         "Full equivalence (rarely required in practice)"],
    ],
    col_widths=[3.0, 5.5, 7.0],
)

add_para(
    "You then test whether each level fits adequately (CFI drop < .01, RMSEA "
    "increase < .015 — Cheung & Rensvold 2002; Chen 2007). If scalar invariance "
    "fails, that's a finding: the two groups answer items in systematically "
    "different ways even conditional on their underlying construct level. "
    "Comparing means becomes partially confounded.",
    size=11,
)

add_para_runs([
    ("PCA gives you a single number per respondent and no diagnostics about "
     "whether comparison is defensible. CFA gives you the score ", False, False),
    ("and", False, True),
    (" the verdict.", False, False),
])

# =====================================================================
# Section 4: What you can test
# =====================================================================
add_heading("4. What you can test")

add_para(
    "PCA's only diagnostic is variance explained. You can ask \"does one "
    "component capture enough?\" (Kaiser criterion, scree plot) but you cannot "
    "ask \"do these items behave the same way in two populations?\""
)

add_para("CFA gives you:")
add_bullet_runs([("Fit indices ", False, False), ("(CFI, TLI, RMSEA, SRMR) ", True, False), ("for whether the one-factor model is adequate", False, False)])
add_bullet_runs([("Invariance tests ", True, False), ("for whether the model is the same across groups or waves", False, False)])
add_bullet_runs([("Modification indices ", True, False), ("suggesting where the model is misspecified", False, False)])
add_bullet_runs([("Standard errors ", True, False), ("on loadings and means", False, False)])

add_para(
    "This is why measurement-invariance testing is a CFA exercise, not a PCA "
    "exercise. code/36_measurement_invariance.R is only possible because CFA "
    "decomposes the model into testable pieces; PCA has no equivalent."
)

# =====================================================================
# Section 5: Why the three methods showed similar results
# =====================================================================
add_heading("5. Why composite means, PC1, and CFA factor scores gave similar d's")

add_para(
    "Because every composite in the EIM2 project has alpha > .7 and a clean "
    "one-factor structure, the differences between the three methods are small "
    "(typically |Δd| < 0.20). When items are highly inter-correlated and load "
    "consistently on one factor:"
)
add_bullet("The unit-weighted composite (mean) ≈ the variance-weighted PC1 ≈ the model-implied CFA factor")
add_bullet("They diverge mainly when (a) some items load much more weakly than others, or (b) measurement invariance fails")

add_para(
    "This is why the Belief in Inevitable Conflict and Minority Inclusion Support "
    "rows showed the biggest CFA-vs-composite gaps — those are exactly the "
    "composites where the invariance tests flagged problems."
)

# =====================================================================
# Section 6: Slope-comparison vs. mean-comparison
# =====================================================================
add_heading("6. Slope-comparison vs. mean-comparison")
add_para(
    "Metric invariance is enough to compare \"relationships between the latent "
    "and other variables.\" Scalar invariance is required to compare \"latent "
    "means.\" These are not the same kind of comparison."
)

add_subheading("Two different things you might compare")

add_para_runs([
    ("\"Relationships between latent and other variables\"", True, False),
    (" = comparing ", False, False),
    ("correlations, regression coefficients, or structural paths", False, True),
    (" that involve the latent. For example:", False, False),
])
add_bullet("Does Belief in Inevitable Conflict predict Minority Inclusion Support more strongly for Russians than for Estonians? (a slope comparison)")
add_bullet("Is the correlation between Superordinate Identity and Comparative Opportunity the same in both groups? (a correlation comparison)")

add_para_runs([
    ("\"Latent means\"", True, False),
    (" = comparing ", False, False),
    ("the average level of the construct", False, True),
    (" between groups. For example:", False, False),
])
add_bullet("Do Russians have more Belief in Inevitable Conflict than Estonians? (a Cohen's d on the latent)")
add_bullet("Did Minority Inclusion Support decrease from 2020 to 2023? (a mean change)")

add_para(
    "These are not the same. Two groups can have identical slopes but very "
    "different means, or identical means but very different slopes."
)

add_subheading("Why metric invariance is enough for one but not the other")

add_para("Recall the measurement equation:")
eq = doc.add_paragraph()
eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = eq.add_run("x_ij = ν_i + λ_i · η_j + ε_ij")
r.italic = True
r.font.size = Pt(12)

add_bullet("λ (loading) controls how a one-unit change in the latent translates into a change in the observed item.")
add_bullet("ν (intercept) controls where the item is anchored — what value the item takes when the latent is at zero.")

add_para_runs([
    ("Slope comparisons", True, False),
    (" depend on λ being equal across groups. If a one-unit increase in the "
     "latent produces a one-unit change in the item for Estonians but a 0.5-unit "
     "change for Russians, then any regression involving that item is on a "
     "different scale per group. Once λ is constrained equal (= metric "
     "invariance), the latent is \"stretched\" the same way in both groups, and "
     "you can compare how it relates to other things. Differences in ν don't "
     "matter for slopes — adding a constant to a regressor shifts the intercept "
     "but doesn't change the slope.", False, False),
])

add_para_runs([
    ("Mean comparisons", True, False),
    (" additionally depend on ν being equal across groups. The mean of the "
     "observed item is ν + λ · mean(η). If Russians score lower than Estonians "
     "on item x, that could be because their latent mean is lower (mean(η) "
     "differs) — ", False, False),
    ("or", False, True),
    (" because their intercept ν is lower (they answer item x lower regardless "
     "of where they are on the construct). Until ν is constrained equal "
     "(= scalar invariance), the two explanations are confounded, and any "
     "\"latent mean difference\" includes intercept variation.", False, False),
])

add_subheading("Intuition: what an intercept non-invariance looks like")

add_para(
    "Take Q67_4 — \"I feel like a second-class citizen\" — as part of "
    "Superordinate Identity. Suppose the loading λ is the same for both groups: "
    "a one-unit increase in belonging produces the same drop in agreement with "
    "this item, for everyone."
)

add_para(
    "But suppose Russians have a higher baseline tendency to endorse that "
    "statement — perhaps because of differential interpretation of \"second-"
    "class,\" language framing, or a response-style effect — independent of "
    "their actual sense of belonging. That's an intercept (ν) difference. "
    "Metric invariance still holds (the slope is the same), but scalar "
    "invariance fails (the intercepts differ)."
)

add_para("Consequence:")
add_bullet("A slope comparison (\"does the relationship between belonging and trust differ by group?\") is fine — both groups translate latent units into item units identically.")
add_bullet("A mean comparison (\"do Russians have lower belonging?\") is biased — the observed item-mean gap is part true-belonging difference, part intercept artifact, and the two cannot be separated.")

# =====================================================================
# Section 7: How this maps to the project
# =====================================================================
add_heading("7. How this maps to the EIM2 project")

add_para(
    "This is why the invariance results matter for the thesis claims. Each "
    "kind of claim requires a different level of invariance:"
)

add_table(
    headers=["Claim you might make", "Invariance level needed"],
    rows=[
        ["\"Belonging correlates with conflict belief similarly in both groups\"", "Metric"],
        ["\"Russians have lower belonging than Estonians (d = +1.08)\"", "Scalar"],
        ["\"Belonging declined more for Estonians than Russians, 2020 → 2023\"", "Scalar (longitudinal)"],
        ["\"The factor structure is interpretable in both groups\"", "Configural"],
    ],
    col_widths=[10.0, 5.5],
)

add_para(
    "Most of the EIM2 headline claims are between-group or within-group mean "
    "comparisons — d values. Those depend on scalar invariance, which fails "
    "for most composites in the invariance results "
    "(reports/Measurement_Invariance.docx). That's why the methods chapter "
    "notes the d values describe observed-group differences without claiming "
    "the gap is purely substantive: part may reflect intercept-level item "
    "functioning differences. This is the standard caveat in cross-cultural "
    "and pre/post-shock survey research."
)

add_para(
    "Slope-comparison claims (metric-invariance-only) do not currently feature "
    "in the analysis since regression was dropped from thesis scope. If a "
    "multivariate model is revisited later — e.g., \"does contact predict "
    "belonging more strongly for one group?\" — metric invariance is the "
    "relevant standard, and the bar is much lower."
)

# =====================================================================
# Section 8: Practical guidance
# =====================================================================
add_heading("8. Practical guidance for the thesis")

add_bullet_runs([
    ("Composite mean ", True, False),
    ("is the canonical reporting metric. It is transparent, intuitive, and "
     "defensible. All values in code/_effect_sizes.tsv use composite means.", False, False),
])
add_bullet_runs([
    ("PC1 ", True, False),
    ("is a validation check that items hang together as expected. Used in "
     "code/04, 05, 07, 08, 09 for exactly that purpose.", False, False),
])
add_bullet_runs([
    ("CFA ", True, False),
    ("is the right tool for asking \"is comparison legitimate?\" — that is the "
     "job of code/36_measurement_invariance.R, not a CFA-factor-score "
     "replacement of composite means.", False, False),
])

add_para(
    "The composite means are the right primary metric. CFA serves as the "
    "methodological footnote explaining when group/year comparisons are clean "
    "(scalar invariance HOLDS) and when they are partially confounded by "
    "measurement non-invariance (scalar invariance FAILS). The invariance "
    "results already in the project are the right way to use CFA, not as an "
    "alternative scoring method."
)


# =====================================================================
# Save
# =====================================================================
doc.save(out_path)
print(f"Saved: {out_path}")
