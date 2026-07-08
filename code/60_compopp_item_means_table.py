"""
Comparative Opportunity Assessment — Item Means by Group × Year
==================================================================
Builds (a) a Word doc table and (b) a publication-quality slope chart
showing the mean response to each Q44 / K3X1 item for Estonian and
Russian respondents in 2020 and 2023.

Q44 scale: 1 = "much better for Estonians," 3 = "equal," 5 = "much better
for other nationalities." Distance from 3.0 indicates perceived inequality
direction; sign indicates which group is favored.

Outputs:
  reports/Comparative_Opportunity_Item_Means_by_Group_Year.docx
  viz/fig_compopp_item_means_grid.jpg
"""

from pathlib import Path
import pandas as pd
import pyreadstat
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Pt, Cm, RGBColor

ROOT = Path(__file__).parent.parent

# ---------- Item labels (Q44_1 ... Q44_12 / K3X1_1 ... K3X1_12) ------------
ITEM_LABELS = [
    "Material well-being",
    "Cultural participation",
    "Education",
    "Social / political rights",
    "Entrepreneurship",
    "Career & jobs",
    "Medical care",
    "Housing",
    "Leisure & holidays",
    "Children & youth opportunities",
    "Sports & exercise",
    "State benefits / services",
]


# ---------- Load data ------------------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy()

df20, _ = pyreadstat.read_sav(str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"),
                              encoding="latin1")
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy()


def to_num(s, dk=9):
    s = pd.to_numeric(s, errors="coerce")
    return s.where(s != dk).dropna().astype(float)


def stats_for(df, eth, var):
    sub = to_num(df[df["ethnicity_binary"] == eth][var])
    return {"M": sub.mean(), "SD": sub.std(ddof=1), "N": len(sub)}


# ---------- Compute the 12 × 4 grid ---------------------------------------
rows = []
for i, label in enumerate(ITEM_LABELS, start=1):
    v23 = f"Q44_{i}"
    v20 = f"K3X1_{i}"
    e20 = stats_for(df20, 0, v20)
    e23 = stats_for(df23, 0, v23)
    r20 = stats_for(df20, 1, v20)
    r23 = stats_for(df23, 1, v23)
    rows.append({
        "idx": i,
        "label": label,
        "e20_M": e20["M"], "e20_SD": e20["SD"], "e20_N": e20["N"],
        "e23_M": e23["M"], "e23_SD": e23["SD"], "e23_N": e23["N"],
        "r20_M": r20["M"], "r20_SD": r20["SD"], "r20_N": r20["N"],
        "r23_M": r23["M"], "r23_SD": r23["SD"], "r23_N": r23["N"],
    })


# ---------- Console diagnostics -------------------------------------------
print("=" * 100)
print("Q44 / K3X1 — Comparative Opportunity Assessment item means by group × year")
print("Scale: 1 = much better for Estonians, 3 = equal, 5 = much better for other nationalities")
print("=" * 100)
print(f"{'Item':<33s}  {'Est 2020':>14s}  {'Est 2023':>14s}  {'Rus 2020':>14s}  {'Rus 2023':>14s}")
for r in rows:
    print(f"  Q44_{r['idx']:<2d}  {r['label']:<25s}  "
          f"{r['e20_M']:5.2f} ({r['e20_SD']:4.2f})  "
          f"{r['e23_M']:5.2f} ({r['e23_SD']:4.2f})  "
          f"{r['r20_M']:5.2f} ({r['r20_SD']:4.2f})  "
          f"{r['r23_M']:5.2f} ({r['r23_SD']:4.2f})")


# ============================================================================
#  WORD DOC — formatted table
# ============================================================================
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin = sec.bottom_margin = Cm(1.8)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)


def H(text, size=14, color=None, bold=True):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color


def body(text, italic=False, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size)
    if italic: r.italic = True


H("Comparative Opportunity Assessment — Item Means by Group × Year",
  size=14, color=RGBColor(0x1f, 0x29, 0x37))
body(
    "Mean response to each Q44 / K3X1 item, broken down by ethnic group and survey wave. "
    "Scale: 1 = much better for Estonians; 3 = equal; 5 = much better for other nationalities. "
    "Cell format: M (SD), n. Distance from 3.0 indicates the perceived direction of inequality.",
    italic=True
)
doc.add_paragraph()

# Table: header + 12 item rows
header = ["Item", "Estonian 2020", "Estonian 2023", "Russian 2020", "Russian 2023"]
table = doc.add_table(rows=1, cols=len(header))
table.style = "Light Grid Accent 1"

for i, h in enumerate(header):
    cell = table.rows[0].cells[i]
    cell.text = ""
    rr = cell.paragraphs[0].add_run(h)
    rr.bold = True; rr.font.size = Pt(9)


def fmt_cell(M, SD, N):
    return f"{M:.2f} ({SD:.2f})\nn = {N}"


for r in rows:
    cells = table.add_row().cells
    cells[0].text = ""
    p = cells[0].paragraphs[0]
    rr = p.add_run(f"Q44_{r['idx']}\n"); rr.bold = True; rr.font.size = Pt(9)
    rr = p.add_run(r["label"]); rr.font.size = Pt(9)

    cells[1].text = fmt_cell(r["e20_M"], r["e20_SD"], r["e20_N"])
    cells[2].text = fmt_cell(r["e23_M"], r["e23_SD"], r["e23_N"])
    cells[3].text = fmt_cell(r["r20_M"], r["r20_SD"], r["r20_N"])
    cells[4].text = fmt_cell(r["r23_M"], r["r23_SD"], r["r23_N"])
    for c in cells[1:]:
        for para in c.paragraphs:
            for run in para.runs:
                run.font.size = Pt(9)

doc.add_paragraph()
body(
    "Note. 2020 items use codes K3X1_1..K3X1_12; 2023 items use Q44_1..Q44_12. "
    "Don't-know responses (code 9) recoded to missing before mean computation. "
    "Sample-size n varies by item due to item-level missingness. Group ns at the composite "
    "level (per canonical TSV): 2020 Estonian = 679, Russian = 603; 2023 Estonian = 844, "
    "Russian = 518.",
    italic=True, size=8
)

out_doc = ROOT / "reports" / "Comparative_Opportunity_Item_Means_by_Group_Year.docx"
doc.save(out_doc)
print(f"\nSaved Word doc: {out_doc}")


# ============================================================================
#  CHART — slope-style dot plot, 12 items × 4 means
# ============================================================================
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size":   9,
    "axes.edgecolor": "#333",
    "axes.linewidth": 0.7,
})

EST = "#2563eb"
RUS = "#d97706"
EST_LIGHT = "#93c5fd"
RUS_LIGHT = "#fcd34d"

fig, ax = plt.subplots(figsize=(11.0, 8.0), dpi=300)

n = len(rows)
y_pos = np.arange(n)[::-1]   # top row = first item

for i, r in enumerate(rows):
    y = y_pos[i]
    # Estonian: 2020 → 2023 line + dots
    ax.plot([r["e20_M"], r["e23_M"]], [y + 0.18, y + 0.18],
            color=EST, linewidth=1.4, alpha=0.85, zorder=2)
    ax.scatter([r["e20_M"]], [y + 0.18], s=46, color=EST_LIGHT,
               edgecolor=EST, linewidth=1.0, zorder=3,
               label="Estonian 2020" if i == 0 else None)
    ax.scatter([r["e23_M"]], [y + 0.18], s=58, color=EST,
               edgecolor=EST, linewidth=0, zorder=3,
               label="Estonian 2023" if i == 0 else None)
    # Russian: 2020 → 2023 line + dots
    ax.plot([r["r20_M"], r["r23_M"]], [y - 0.18, y - 0.18],
            color=RUS, linewidth=1.4, alpha=0.85, zorder=2)
    ax.scatter([r["r20_M"]], [y - 0.18], s=46, color=RUS_LIGHT,
               edgecolor=RUS, linewidth=1.0, zorder=3,
               label="Russian 2020" if i == 0 else None)
    ax.scatter([r["r23_M"]], [y - 0.18], s=58, color=RUS,
               edgecolor=RUS, linewidth=0, zorder=3,
               label="Russian 2023" if i == 0 else None)

# Equality reference line
ax.axvline(3.0, color="#666", linewidth=1.0, linestyle="--",
           zorder=1, label="3.0 = equal opportunity")

# Endpoint annotations: scale meanings
ax.text(1.0, n + 0.4, "1\nmuch better\nfor Estonians",
        ha="center", va="bottom", fontsize=7.5, color="#444",
        linespacing=1.2)
ax.text(3.0, n + 0.4, "3\nequal",
        ha="center", va="bottom", fontsize=7.5, color="#444",
        linespacing=1.2)
ax.text(5.0, n + 0.4, "5\nmuch better\nfor other nationalities",
        ha="center", va="bottom", fontsize=7.5, color="#444",
        linespacing=1.2)

# Axes
ax.set_yticks(y_pos)
ax.set_yticklabels([f"Q44_{r['idx']}  {r['label']}" for r in rows], fontsize=9)
ax.set_xlim(0.7, 5.3)
ax.set_xticks([1, 1.5, 2, 2.5, 3, 3.5, 4, 4.5, 5])
ax.set_ylim(-0.9, n + 1.3)
ax.set_xlabel("Mean response (1–5 scale)", fontsize=10, labelpad=8)

for s in ("top", "right"):
    ax.spines[s].set_visible(False)
ax.spines["left"].set_color("#666")
ax.spines["bottom"].set_color("#666")
ax.tick_params(axis="x", length=3, color="#666")
ax.tick_params(axis="y", length=0)

# Title and subtitle
fig.text(0.04, 0.965,
         "Comparative Opportunity Assessment — Item Means by Group × Year",
         fontsize=13, fontweight="bold", ha="left")
fig.text(0.04, 0.940,
         "Each item shows two lines: Estonian respondents (blue) and Russian respondents (orange).  "
         "Open dot = 2020; filled dot = 2023.  Dashed line at 3.0 = perceived equal opportunity.",
         fontsize=9, color="#444", ha="left")

# Legend
ax.legend(loc="upper left", bbox_to_anchor=(0.005, -0.06),
          frameon=False, fontsize=8, ncol=5, handlelength=1.4)

# Footer
fig.text(0.04, 0.020,
         "Means below 3.0 indicate the respondent group perceives the listed domain as more accessible to Estonians; above 3.0, more accessible to other nationalities.",
         fontsize=7, color="#555", ha="left")

plt.subplots_adjust(left=0.32, right=0.985, top=0.88, bottom=0.13)

out_jpg = ROOT / "viz" / "fig_compopp_item_means_grid.jpg"
plt.savefig(out_jpg, dpi=300, format="jpg",
            facecolor="white", pil_kwargs={"quality": 95})
print(f"Saved chart: {out_jpg}")
