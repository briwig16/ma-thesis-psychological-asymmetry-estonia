"""
Out-group Language Ability → SD: Primary Out-group Regression Check
====================================================================
Targeted, hypothesis-driven test: does out-group language ability predict
reduced social distance from that out-group? This is the linguistic-contact
extension of the classical contact-theory test in 63_contact_theory_regression.

Four regressions, one per group × year cell:
  - Russian 2020:    SD from Estonian-speakers ~ Estonian language ability
  - Russian 2023:    SD from Estonian-speakers ~ Estonian language ability
  - Estonian 2020:   SD from Russian-speakers  ~ Russian language ability
  - Estonian 2023:   SD from Russian-speakers  ~ Russian language ability

DV: SD: Primary Out-group, PC1 score.
    Fit per group pooled across years (items differ by group; PC1 must
    therefore be group-specific). Sign-oriented so HIGHER = MORE distance.

IV: Out-group language ability (Q71_*/K5_*).
    Raw scale: 1 = native, 2 = fluent, 3 = u/s/w, 4 = u/s a little,
               5 = understand a little / cannot speak, 6 = none, 9 = DK.
    Recoded as (7 − raw) with 9 → NaN, so HIGHER = MORE proficient.

Each regression is a single-equation OLS. No covariates. Bivariate linguistic-
contact association cross-sectionally within each group × year cell.

Outputs:
  code/_language_sd_primary_regression.tsv — machine-readable
  reports/Language_SD_Primary_Regression.docx — formatted Word doc
  viz/fig_language_sd_primary_regression.jpg — 2×2 scatter+fit panels (300 DPI)
"""

from pathlib import Path

import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm
import matplotlib.pyplot as plt
from sklearn.decomposition import PCA
from docx import Document
from docx.shared import Pt, Cm, RGBColor

ROOT = Path(__file__).parent.parent

# ---------- Load -----------------------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)
df23["__year__"] = 2023

df20, _ = pyreadstat.read_sav(str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"),
                              encoding="latin1")
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)
df20["__year__"] = 2020


def to_num(s, dk=9):
    s = pd.to_numeric(s, errors="coerce")
    return s.where(s != dk).astype(float)


def clean_items(df, items, dk=9):
    sub = df[items].apply(pd.to_numeric, errors="coerce")
    return sub.where(sub != dk)


def recode_lang(s):
    s = to_num(s, dk=9)
    return 7 - s   # 1=native..6=none → 1=none..6=native


def fit_pc1_oriented(items_df):
    """Fit PCA on listwise-complete rows; return PC1 oriented to track item-
    mean direction (higher = more distance)."""
    mask = items_df.notna().all(axis=1)
    X = items_df[mask].to_numpy()
    pca = PCA(n_components=1).fit(X)
    scores = pca.transform(X)[:, 0]
    pc1 = pd.Series(index=items_df.index, dtype=float)
    pc1.loc[mask] = scores
    if pc1.corr(items_df.mean(axis=1)) < 0:
        pc1 = -pc1
    return pc1, float(pca.explained_variance_ratio_[0])


# ---------- Cell specifications -------------------------------------------
SPECS = {
    "Estonian": {
        "color": "#2563eb",
        "items_23": ["Q57_1", "Q58_1", "Q59_1"],
        "items_20": ["K4X7_1", "K4X8_1", "K4X9_1"],
        "lang_23": "Q71_2",
        "lang_20": "K5_2",
        "sd_label": "SD from Russian-speakers",
        "lang_label": "Russian language ability",
    },
    "Russian": {
        "color": "#d97706",
        "items_23": ["Q57_2", "Q58_2", "Q59_2"],
        "items_20": ["K4X7_2", "K4X8_2", "K4X9_2"],
        "lang_23": "Q71_1",
        "lang_20": "K5_1",
        "sd_label": "SD from Estonian-speakers",
        "lang_label": "Estonian language ability",
    },
}


# ---------- Build PC1 per group (pooled years) + per-cell data ------------
cells = []  # (group, year, df_pair, spec)
for grp_name, spec in SPECS.items():
    grp_code = 0 if grp_name == "Estonian" else 1
    cols = ["i1", "i2", "i3"]

    sub23 = df23[df23["ethnicity_binary"] == grp_code].copy().reset_index(drop=True)
    sub20 = df20[df20["ethnicity_binary"] == grp_code].copy().reset_index(drop=True)

    items_23 = clean_items(sub23, spec["items_23"]).rename(columns=dict(zip(spec["items_23"], cols)))
    items_20 = clean_items(sub20, spec["items_20"]).rename(columns=dict(zip(spec["items_20"], cols)))
    pooled = pd.concat([items_23, items_20], ignore_index=True)
    pc1, ve = fit_pc1_oriented(pooled)

    spec["ve"] = ve   # PC1 variance explained (per group)

    pooled_meta = pd.concat([
        sub23[["__year__"]].assign(lang=recode_lang(sub23[spec["lang_23"]])),
        sub20[["__year__"]].assign(lang=recode_lang(sub20[spec["lang_20"]])),
    ], ignore_index=True)
    pooled_meta["pc1"] = pc1.values

    for yr in (2020, 2023):
        df_pair = (pooled_meta[pooled_meta["__year__"] == yr]
                   [["pc1", "lang"]].rename(columns={"pc1": "SD", "lang": "Lang"})
                   .dropna())
        cells.append((grp_name, yr, df_pair, spec))


# ---------- Regressions ----------------------------------------------------
def stars(p):
    return "***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else "ns"


results = []
for grp, yr, df_pair, spec in cells:
    X = sm.add_constant(df_pair["Lang"])
    y = df_pair["SD"]
    m = sm.OLS(y, X).fit(cov_type="HC3")
    b0 = m.params["const"]; b1 = m.params["Lang"]
    se1 = m.bse["Lang"]
    ci_lo, ci_hi = m.conf_int().loc["Lang"]
    t = m.tvalues["Lang"]; p = m.pvalues["Lang"]
    # Standardized beta (using sample SDs of x and y)
    sx = df_pair["Lang"].std(ddof=0); sy = df_pair["SD"].std(ddof=0)
    beta = b1 * sx / sy if sy > 0 else np.nan
    results.append({
        "label": f"{grp} {yr}",
        "group": grp, "year": yr,
        "n": int(m.nobs),
        "b0": b0, "b1": b1, "se1": se1,
        "ci_lo": ci_lo, "ci_hi": ci_hi,
        "t": t, "p": p, "beta": beta, "r2": m.rsquared,
        "lang_label": spec["lang_label"], "sd_label": spec["sd_label"],
        "color": spec["color"], "ve": spec["ve"],
        "model": m, "data": df_pair,
    })


# ---------- Console + TSV --------------------------------------------------
print("=" * 100)
print("OUT-GROUP LANGUAGE → SD: PRIMARY OUT-GROUP REGRESSION")
print("DV: SD Primary PC1 (higher = more distance).  IV: out-group language ability (higher = more proficient).")
print("Contact theory predicts NEGATIVE slope: more out-group language → less out-group distance.")
print("=" * 100)
for r in results:
    print(f"\n  {r['label']}  ({r['lang_label']} → {r['sd_label']})")
    print(f"    n = {r['n']}    PC1 var explained = {r['ve']*100:.1f}%")
    print(f"    b₁ = {r['b1']:+.4f}   SE = {r['se1']:.4f}   95% CI [{r['ci_lo']:+.4f}, {r['ci_hi']:+.4f}]")
    print(f"    β  = {r['beta']:+.3f}   t = {r['t']:+.3f}   p = {r['p']:.4g} {stars(r['p'])}   R² = {r['r2']:.4f}")

# TSV
out_rows = [
    {"label": r["label"], "group": r["group"], "year": r["year"], "n": r["n"],
     "b1": r["b1"], "SE": r["se1"], "CI_low": r["ci_lo"], "CI_high": r["ci_hi"],
     "t": r["t"], "p": r["p"], "beta": r["beta"], "R2": r["r2"]}
    for r in results
]
tsv = ROOT / "code" / "_language_sd_primary_regression.tsv"
pd.DataFrame(out_rows).to_csv(tsv, sep="\t", index=False, float_format="%.4f")
print(f"\nSaved TSV: {tsv}")


# ============================================================================
#  Word doc
# ============================================================================
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(2.0)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)


def H(t, sz=14):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(sz)


def body(t, italic=False, sz=10):
    p = doc.add_paragraph(); r = p.add_run(t); r.font.size = Pt(sz); r.italic = italic


H("Out-group Language Ability → SD: Primary Out-group Regression")
body(
    "Targeted bivariate regression of social distance from the primary out-group on "
    "self-reported ability in that out-group's language, run separately for each "
    "group × year cell. Linguistic-contact extension of the classical contact-theory "
    "test (Allport 1954; Pettigrew & Tropp 2006).",
    italic=True,
)
doc.add_paragraph()

# Table
header = ["Group × Year", "n", "b₁ (slope)", "β (std.)", "SE", "95% CI", "t", "p", "R²"]
table = doc.add_table(rows=1, cols=len(header))
table.style = "Light Grid Accent 1"
for i, h in enumerate(header):
    c = table.rows[0].cells[i]; c.text = ""
    rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)


def fmt_p(p):
    if p < 0.001: return "< .001"
    return f"{p:.3f}".lstrip("0")


def fmt_signed(x, dec=4):
    return f"{x:+.{dec}f}".replace("-", "−")


for r in results:
    cells_r = table.add_row().cells
    cells_r[0].text = r["label"]
    cells_r[1].text = str(r["n"])
    cells_r[2].text = f"{fmt_signed(r['b1'])} {stars(r['p'])}"
    cells_r[3].text = fmt_signed(r['beta'], dec=3)
    cells_r[4].text = f"{r['se1']:.4f}"
    cells_r[5].text = f"[{fmt_signed(r['ci_lo'])}, {fmt_signed(r['ci_hi'])}]"
    cells_r[6].text = fmt_signed(r["t"], dec=2)
    cells_r[7].text = fmt_p(r["p"])
    cells_r[8].text = f"{r['r2']:.3f}"
    for c in cells_r:
        for para in c.paragraphs:
            for run in para.runs:
                run.font.size = Pt(9)

doc.add_paragraph()
body(
    "Note. Each row reports an OLS regression of SD: Primary Out-group PC1 (higher = more "
    "social distance) on out-group language ability (1 = none … 6 = native). PC1 is fit on the "
    "three SD items, separately per group, pooled across years (items differ by group so a "
    "pooled-across-groups PC1 is not defined; this means cross-group PC1 metrics are similar "
    "but not strictly identical). HC3 robust standard errors. b₁ = unstandardised slope per "
    "1-unit increase in language ability. β = standardised slope. Contact theory predicts a "
    "negative slope. Listwise deletion. Cross-sectional design — no causal interpretation. "
    "*** p < .001, ** p < .01, * p < .05, ns = not significant.",
    italic=True, sz=8,
)

out_doc = ROOT / "reports" / "Language_SD_Primary_Regression.docx"
doc.save(out_doc)
print(f"Saved Word doc: {out_doc}")


# ============================================================================
#  Figure — 2×2 scatter+fit panels
# ============================================================================
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 9,
    "axes.edgecolor": "#333",
    "axes.linewidth": 0.6,
})

fig, axes = plt.subplots(2, 2, figsize=(11.0, 9.0), dpi=300, sharex=True, sharey=True)

# Layout: top row Russian (orange), bottom row Estonian (blue)
order = [
    ("Russian", 2020,  axes[0, 0]),
    ("Russian", 2023,  axes[0, 1]),
    ("Estonian", 2020, axes[1, 0]),
    ("Estonian", 2023, axes[1, 1]),
]

# Determine PC1 (y) plotting range from observed data
all_pc1 = pd.concat([r["data"]["SD"] for r in results])
y_min, y_max = all_pc1.min() - 0.3, all_pc1.max() + 0.3

for grp, yr, ax in order:
    r = next(x for x in results if x["group"] == grp and x["year"] == yr)
    df_pair = r["data"]
    rng = np.random.default_rng(42)
    jx = rng.uniform(-0.08, 0.08, size=len(df_pair))
    jy = rng.uniform(-0.05, 0.05, size=len(df_pair))
    ax.scatter(df_pair["Lang"] + jx, df_pair["SD"] + jy,
               s=12, color=r["color"], alpha=0.30, edgecolor="none", zorder=2)

    # Fit line + CI band
    x_grid = np.linspace(1, 6, 60)
    pred = r["model"].get_prediction(
        sm.add_constant(pd.Series(x_grid, name="Lang"))
    )
    ci = pred.summary_frame(alpha=0.05)
    ax.plot(x_grid, ci["mean"], color="#111", linewidth=1.6, zorder=3)
    ax.fill_between(x_grid, ci["mean_ci_lower"], ci["mean_ci_upper"],
                    color="#111", alpha=0.10, zorder=1)

    ax.set_title(r["label"], fontsize=11, fontweight="bold", loc="left", pad=8)
    ax.text(0.02, 0.97,
            f"b₁ = {r['b1']:+.3f}  ({stars(r['p'])})\n"
            f"β  = {r['beta']:+.3f}\n"
            f"95% CI [{r['ci_lo']:+.3f}, {r['ci_hi']:+.3f}]\n"
            f"R² = {r['r2']:.3f}   n = {r['n']}",
            transform=ax.transAxes, ha="left", va="top",
            fontsize=8.5, color="#222", family="monospace",
            bbox=dict(boxstyle="round,pad=0.4", facecolor="white",
                      edgecolor="#888", linewidth=0.5, alpha=0.92))

    ax.set_xlim(0.5, 6.5)
    ax.set_ylim(y_min, y_max)
    ax.set_xticks([1, 2, 3, 4, 5, 6])
    ax.set_xticklabels(["1\nnone", "2", "3", "4", "5", "6\nnative"], fontsize=8)
    for s_ in ("top", "right"):
        ax.spines[s_].set_visible(False)
    ax.tick_params(axis="both", labelsize=8, color="#888")
    ax.axhline(0, color="#999", linewidth=0.5, linestyle="--", zorder=0)

axes[1, 0].set_xlabel("Out-group language ability  (1 = none … 6 = native)", fontsize=9, color="#444")
axes[1, 1].set_xlabel("Out-group language ability  (1 = none … 6 = native)", fontsize=9, color="#444")
axes[0, 0].set_ylabel("SD: Primary Out-group PC1\n(higher = more distance)", fontsize=9, color="#444")
axes[1, 0].set_ylabel("SD: Primary Out-group PC1\n(higher = more distance)", fontsize=9, color="#444")

fig.text(0.04, 0.965,
         "Out-group Language → SD: Primary Out-group Regression",
         fontsize=13, fontweight="bold", ha="left")
fig.text(0.04, 0.940,
         "Top row: Russian respondents rating Estonian-speakers (orange).  "
         "Bottom row: Estonian respondents rating Russian-speakers (blue).",
         fontsize=9, color="#444")
fig.text(0.04, 0.922,
         "Contact theory predicts a NEGATIVE slope: more out-group language proficiency → less social distance.",
         fontsize=9, color="#444")
fig.text(0.04, 0.018,
         "Each dot = one respondent (jittered ±0.08 on x, ±0.05 on y for legibility). Black line = OLS fit. Shaded band = 95% CI on the conditional mean.",
         fontsize=7, color="#555")
fig.text(0.04, 0.005,
         "Cross-sectional associations within each cell; no causal interpretation.",
         fontsize=7, color="#555")

plt.subplots_adjust(left=0.07, right=0.985, top=0.88, bottom=0.085,
                    wspace=0.10, hspace=0.30)

out_jpg = ROOT / "viz" / "fig_language_sd_primary_regression.jpg"
plt.savefig(out_jpg, dpi=300, format="jpg",
            facecolor="white", pil_kwargs={"quality": 95})
print(f"Saved chart: {out_jpg}")
