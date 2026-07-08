"""
Build Word document: multi-group CFA fit indices + per-cell PC1 variance
for all 8 composites. Includes both between-group (Estonian vs Russian
within each year) and within-group (2020 vs 2023 within each ethnicity)
multi-group CFAs.

Reads:
  code/_latent_means_fit.tsv      (between-group MG-CFA, by year)
  code/_within_group_mg_cfa.tsv   (within-group MG-CFA, by ethnicity)
  code/_latent_means_alpha.tsv    (per-cell N and α)
  code/_per_cell_pc1.tsv          (per-cell PC1 % variance)

Output: reports/MG_CFA_Fit_and_PCA_Summary.docx

Numbers only. No interpretive verdict on fit.
"""

from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches

ROOT = Path("/Users/brianwiggins/Desktop/Claude Code/EIM2")
FIT     = pd.read_csv(ROOT / "code" / "_latent_means_fit.tsv",    sep="\t")
WIN     = pd.read_csv(ROOT / "code" / "_within_group_mg_cfa.tsv", sep="\t")
ALPHA   = pd.read_csv(ROOT / "code" / "_latent_means_alpha.tsv",  sep="\t")
PCA     = pd.read_csv(ROOT / "code" / "_per_cell_pc1.tsv",        sep="\t")

OUT = ROOT / "reports" / "MG_CFA_Fit_and_PCA_Summary.docx"

COMPOSITES = [
    "Superordinate Identity",
    "SD: Primary Out-group",
    "SD: General Out-group",
    "Comparative Opportunity Assessment",
    "Belief in Inevitable Conflict",
    "Minority Support Inclusion",
    "Contact: Estonian Speakers",
    "Contact: Russian Speakers",
]

ITEMS_2023 = {
    "Superordinate Identity":            "Q67_2, Q67_4 (rev), Q67_5",
    "SD: Primary Out-group":             "Q57_1/2, Q58_1/2, Q59_1/2 (group-specific referent)",
    "SD: General Out-group":             "Q57_4, Q57_5, Q58_4, Q58_5, Q59_4, Q59_5",
    "Comparative Opportunity Assessment": "Q44_1 – Q44_12",
    "Belief in Inevitable Conflict":     "Q63_1 (rev), Q63_2 (rev), Q63_3, Q63_4",
    "Minority Support Inclusion":        "Q68_1, Q68_2, Q68_3",
    "Contact: Estonian Speakers":        "Q51_1 – Q51_6",
    "Contact: Russian Speakers":         "Q52_1 – Q52_6",
}
ITEMS_2020 = {
    "Superordinate Identity":            "K6X5_2, K6X5_3 (rev), K6X5_4",
    "SD: Primary Out-group":             "K4X7_1/2, K4X8_1/2, K4X9_1/2 (group-specific referent)",
    "SD: General Out-group":             "K4X7_3, K4X8_3, K4X9_3",
    "Comparative Opportunity Assessment": "K3X1_1 – K3X1_12",
    "Belief in Inevitable Conflict":     "K6X1_1 (rev), K6X1_2 (rev), K6X1_3, K6X1_4",
    "Minority Support Inclusion":        "K6X6_1, K6X6_2, K6X6_3",
    "Contact: Estonian Speakers":        "K4X1_1 – K4X1_6",
    "Contact: Russian Speakers":         "K4X2_1 – K4X2_6",
}


def fmt(x, d=3):
    if pd.isna(x) or x == "":
        return "—"
    try:
        return f"{float(x):.{d}f}"
    except (TypeError, ValueError):
        return str(x)


def fmt_p(x):
    if pd.isna(x) or x == "":
        return "—"
    try:
        v = float(x)
    except (TypeError, ValueError):
        return str(x)
    return "<.001" if v < 0.001 else f"{v:.3f}"


def fmt_int(x):
    if pd.isna(x):
        return "—"
    try:
        return str(int(float(x)))
    except (TypeError, ValueError):
        return str(x)


doc = Document()
section = doc.sections[0]
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)


def add_table(headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        c = tbl.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.rows[i + 1].cells[j]
            c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            r.font.size = Pt(9)
    return tbl


# Title
doc.add_heading("Multi-Group CFA Fit Indices and PC1 Variance — All Composites", level=0)

p = doc.add_paragraph()
p.add_run(
    "Per-composite multi-group confirmatory factor analysis at three invariance levels "
    "(configural, metric, scalar). Two grouping structures are reported:"
).font.size = Pt(10)

p = doc.add_paragraph()
p.add_run("    • Between-group: Estonian vs. Russian, fit separately within each year.\n").font.size = Pt(10)
p.add_run("    • Within-group: 2020 vs. 2023, fit separately within each ethnic group.").font.size = Pt(10)

p = doc.add_paragraph()
p.add_run(
    "Estimator: MLR. Missing data: FIML. Fit indices: scaled χ², df, p, CFI, TLI, "
    "RMSEA, SRMR. Principal component analysis run separately on each cell "
    "(listwise deletion); PC1 % variance reported."
).font.size = Pt(10)

p = doc.add_paragraph()
p.add_run(
    "SD: Primary Out-group uses group-specific items (Estonians rate Russian-speakers; "
    "Russians rate Estonian-speakers). SD: General Out-group has different item sets in "
    "2020 (3 items) and 2023 (6 items); between-group MG-CFA fit per year, "
    "within-group MG-CFA not defined."
).font.size = Pt(10)

# -----------------------------------------------------------------------------
# Per-composite sections
# -----------------------------------------------------------------------------
for comp in COMPOSITES:
    doc.add_heading(comp, level=1)

    # Items
    p = doc.add_paragraph()
    r = p.add_run("Items (2023): "); r.bold = True; r.font.size = Pt(9)
    p.add_run(ITEMS_2023[comp]).font.size = Pt(9)
    p = doc.add_paragraph()
    r = p.add_run("Items (2020): "); r.bold = True; r.font.size = Pt(9)
    p.add_run(ITEMS_2020[comp]).font.size = Pt(9)

    # Per-cell descriptive table (N, α, PC1 % variance)
    doc.add_heading("Per-cell sample, reliability, PC1 variance", level=3)
    rows = []
    for year in (2020, 2023):
        for group in ("Estonian", "Russian"):
            a = ALPHA[(ALPHA.composite == comp) & (ALPHA.year == year) & (ALPHA.group == group)]
            pca = PCA[(PCA.composite == comp) & (PCA.year == year) & (PCA.group == group)]
            n_items = a["n_items"].iloc[0] if len(a) else "—"
            alpha = a["alpha"].iloc[0] if len(a) else None
            n_alpha = a["n"].iloc[0] if len(a) else None
            n_pca = pca["n"].iloc[0] if len(pca) else None
            eig = pca["pc1_eigenvalue"].iloc[0] if len(pca) else None
            pct = pca["pc1_pct_variance"].iloc[0] if len(pca) else None
            rows.append([
                year, group, fmt_int(n_items),
                fmt_int(n_alpha), fmt(alpha, 3),
                fmt_int(n_pca), fmt(eig, 3), fmt(pct, 2),
            ])
    add_table(
        ["Year", "Group", "k items", "N (α)", "Cronbach α",
         "N (PCA)", "PC1 eigenvalue", "PC1 % variance"],
        rows,
    )

    # Multi-group CFA — between-group, 2020
    doc.add_heading("Between-group MG-CFA — 2020 (Estonian vs Russian)", level=3)
    rows = []
    for level in ("configural", "metric", "scalar"):
        f = FIT[(FIT.composite == comp) & (FIT.year == 2020) & (FIT.level == level)]
        if len(f) == 0:
            continue
        f = f.iloc[0]
        rows.append([
            level, fmt(f.chisq, 3), fmt_int(f.df), fmt_p(f.p),
            fmt(f.cfi, 4), fmt(f.tli, 4), fmt(f.rmsea, 4), fmt(f.srmr, 4),
        ])
    add_table(
        ["Level", "χ²", "df", "p", "CFI", "TLI", "RMSEA", "SRMR"],
        rows,
    )

    # Multi-group CFA — between-group, 2023
    doc.add_heading("Between-group MG-CFA — 2023 (Estonian vs Russian)", level=3)
    rows = []
    for level in ("configural", "metric", "scalar"):
        f = FIT[(FIT.composite == comp) & (FIT.year == 2023) & (FIT.level == level)]
        if len(f) == 0:
            continue
        f = f.iloc[0]
        rows.append([
            level, fmt(f.chisq, 3), fmt_int(f.df), fmt_p(f.p),
            fmt(f.cfi, 4), fmt(f.tli, 4), fmt(f.rmsea, 4), fmt(f.srmr, 4),
        ])
    add_table(
        ["Level", "χ²", "df", "p", "CFI", "TLI", "RMSEA", "SRMR"],
        rows,
    )

    # Multi-group CFA — within-group (2020 vs 2023, by ethnicity)
    win_comp = WIN[WIN.composite == comp]
    if len(win_comp) > 0:
        for grp in ("Estonian", "Russian"):
            doc.add_heading(f"Within-group MG-CFA — {grp} (2020 vs 2023)", level=3)
            rows = []
            for level in ("configural", "metric", "scalar"):
                f = win_comp[(win_comp.group == grp) & (win_comp.level == level)]
                if len(f) == 0:
                    continue
                f = f.iloc[0]
                rows.append([
                    level, fmt(f.chisq, 3), fmt_int(f.df), fmt_p(f.p),
                    fmt(f.cfi, 4), fmt(f.tli, 4), fmt(f.rmsea, 4), fmt(f.srmr, 4),
                ])
            add_table(
                ["Level", "χ²", "df", "p", "CFI", "TLI", "RMSEA", "SRMR"],
                rows,
            )
    else:
        p = doc.add_paragraph()
        r = p.add_run(
            "Within-group MG-CFA not defined for this composite "
            "(item set differs across waves)."
        )
        r.italic = True
        r.font.size = Pt(9)

    doc.add_paragraph()

# Footer
doc.add_paragraph()
foot = doc.add_paragraph()
r = foot.add_run(
    f"Sources: code/_latent_means_fit.tsv, code/_within_group_mg_cfa.tsv, "
    f"code/_latent_means_alpha.tsv, code/_per_cell_pc1.tsv. "
    f"Generated by code/92_mg_cfa_summary_docx.py."
)
r.italic = True
r.font.size = Pt(8)

doc.save(OUT)
print(f"Wrote: {OUT.relative_to(ROOT)}")
