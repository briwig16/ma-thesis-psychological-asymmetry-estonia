"""
Build reports/Question_Wording_EN_EE.docx — exact wording (English + Estonian)
for every item used in every variable in the thesis.

Sources:
  - English: codebook/EIM23datamap.xlsx (Variable label column, full text)
  - Estonian: 2020 SPSS file variable labels (latin1 → utf-8 re-decode);
    SPSS labels are truncated at ~80 chars, so the Estonian text is
    sometimes partial. Where this happens, a "[truncated in SPSS]" marker
    is appended.

Structure: one heading per variable; under each heading, a table with
columns [Item code (2023) | Item code (2020) | English | Estonian].
"""

from pathlib import Path
import pandas as pd
import pyreadstat
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent.parent

# --- Load 2023 English codebook ---------------------------------------------
cb = pd.read_excel(ROOT / "codebook" / "EIM23datamap.xlsx",
                   sheet_name="Sheet1", header=1)
cb.columns = [str(c).strip() for c in cb.columns]
EN = dict(zip(cb["Variable name"].astype(str), cb["Variable label"].astype(str)))

# --- Load 2020 SPSS labels (Estonian, with latin1→utf-8 fix) ----------------
df20, meta = pyreadstat.read_sav(
    str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"), encoding="latin1"
)
EE = {}
for name, lbl in zip(meta.column_names, meta.column_labels):
    if lbl is None:
        EE[name] = ""
        continue
    try:
        lbl = lbl.encode("latin1").decode("utf-8")
    except Exception:
        pass
    EE[name] = lbl


def fmt_ee(text):
    if not text:
        return "—"
    if len(text) >= 76:
        return text + "  [truncated in SPSS]"
    return text


# --- Variable structure -----------------------------------------------------
# Each entry: (variable_name, intro_text, [(item_2023, item_2020, gloss), ...])
VARIABLES = [
    ("1. Superordinate Identity", None,
     [("Q67_2", "K6X5_2", "Pride in Estonian flag"),
      ("Q67_4", "K6X5_3", "Feel like a second-class citizen"),
      ("Q67_5", "K6X5_4", "Feel part of Estonian society")]),

    ("2. SD: Primary Out-group",
     "Estonians rate Russian-speakers; Russians rate Estonian-speakers. "
     "Items _1 and _2 are the two language groups; items presented in three "
     "contexts (neighbours / work-study / marriage).",
     [("Q57_1", "K4X7_1", "Neighbours — Russian-speakers (asked of Estonians)"),
      ("Q57_2", "K4X7_2", "Neighbours — Estonian-speakers (asked of Russians)"),
      ("Q58_1", "K4X8_1", "Work / study — Russian-speakers (asked of Estonians)"),
      ("Q58_2", "K4X8_2", "Work / study — Estonian-speakers (asked of Russians)"),
      ("Q59_1", "K4X9_1", "Marriage in family — Russian-speakers (asked of Estonians)"),
      ("Q59_2", "K4X9_2", "Marriage in family — Estonian-speakers (asked of Russians)")]),

    ("3. SD: General Out-group",
     "2023: six items (other Europeans + non-Europeans, in three contexts). "
     "2020: three items (new immigrants — past 5 years — in three contexts). "
     "Item set differs across waves; within-group cross-year comparison is "
     "treated as a methodological caveat.",
     [("Q57_4", "K4X7_3", "Neighbours — other Europeans (2023) / new immigrants (2020)"),
      ("Q57_5", None,     "Neighbours — people from outside Europe (2023 only)"),
      ("Q58_4", "K4X8_3", "Work / study — other Europeans (2023) / new immigrants (2020)"),
      ("Q58_5", None,     "Work / study — people from outside Europe (2023 only)"),
      ("Q59_4", "K4X9_3", "Marriage in family — other Europeans (2023) / new immigrants (2020)"),
      ("Q59_5", None,     "Marriage in family — people from outside Europe (2023 only)")]),

    ("4. Comparative Opportunity Assessment",
     "Q44 / K3X1 — 12 items rating life-opportunities. 1 = much better for "
     "Estonians, 3 = equal, 5 = much better for other nationalities.",
     [("Q44_1",  "K3X1_1",  "Material well-being"),
      ("Q44_2",  "K3X1_2",  "Cultural participation"),
      ("Q44_3",  "K3X1_3",  "Education"),
      ("Q44_4",  "K3X1_4",  "Social / political rights"),
      ("Q44_5",  "K3X1_5",  "Entrepreneurship"),
      ("Q44_6",  "K3X1_6",  "Career & jobs"),
      ("Q44_7",  "K3X1_7",  "Medical care"),
      ("Q44_8",  "K3X1_8",  "Housing"),
      ("Q44_9",  "K3X1_9",  "Leisure & holidays"),
      ("Q44_10", "K3X1_10", "Children & youth"),
      ("Q44_11", "K3X1_11", "Sports & exercise"),
      ("Q44_12", "K3X1_12", "State benefits / services")]),

    ("5. Belief in Inevitable Conflict",
     "Higher composite = stronger belief that ethnic conflict is inevitable "
     "(post-2026-04-30 Option B recoding).",
     [("Q63_1", "K6X1_1", "Ethnic conflict is inevitable (rev.)"),
      ("Q63_2", "K6X1_2", "Differences divide Estonian society (rev.)"),
      ("Q63_3", "K6X1_3", "Different ethnic groups can cooperate"),
      ("Q63_4", "K6X1_4", "Immigration enriches Estonia")]),

    ("6. Minority Inclusion Support",
     "Higher composite (inverted scale) = more supportive of minority inclusion.",
     [("Q68_1", "K6X6_1", "Involve non-Estonians in economic life"),
      ("Q68_2", "K6X6_2", "Involve non-Estonians in governance"),
      ("Q68_3", "K6X6_3", "Understand the opinions of non-Estonians")]),

    ("7. Contact: Estonian-speakers",
     "Q51 / K4X1 — six contexts. Scale 1 = almost every day … 5 = have not "
     "communicated. (Inverted in composites so higher = more contact.)",
     [("Q51_1", "K4X1_1", "Work / school"),
      ("Q51_2", "K4X1_2", "Neighbours"),
      ("Q51_3", "K4X1_3", "Internet / social media"),
      ("Q51_4", "K4X1_4", "Leisure"),
      ("Q51_5", "K4X1_5", "Family"),
      ("Q51_6", "K4X1_6", "Friends")]),

    ("8. Contact: Russian-speakers",
     "Q52 / K4X2 — six contexts; same scale as Contact: Estonian-speakers.",
     [("Q52_1", "K4X2_1", "Work / school"),
      ("Q52_2", "K4X2_2", "Neighbours"),
      ("Q52_3", "K4X2_3", "Internet / social media"),
      ("Q52_4", "K4X2_4", "Leisure"),
      ("Q52_5", "K4X2_5", "Family"),
      ("Q52_6", "K4X2_6", "Friends")]),

    ("9. Group Identity Patterns (single-item)",
     "Single-item measure of dual vs. exclusive identity (5-point continuum).",
     [("Q66", "K6X4", "Identification with own nationality vs. Estonian people")]),

    ("10. Territorial Attachment (single-item)",
     "Single-item measure: \"feel at home in Estonia\" (dropped from "
     "Superordinate Identity composite but reported standalone).",
     [("Q67_1", "K6X5_1", "Feel at home in Estonia")]),
]


# --- Build the Word document ------------------------------------------------
doc = Document()

# Page setup
for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

FONT = "Times New Roman"
SIZE = Pt(12)
BLACK = RGBColor(0x00, 0x00, 0x00)

# Default style
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = SIZE
style.font.color.rgb = BLACK

def setfont(run, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = SIZE
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = BLACK

# Title block
title = doc.add_paragraph()
setfont(title.add_run("Question Wording — English and Estonian"), bold=True)

sub = doc.add_paragraph()
setfont(sub.add_run("Items used in each composite and single-item variable of the EIM2 analysis"),
        italic=True, color=RGBColor(0x55, 0x55, 0x55))

# Sources / notes
note = doc.add_paragraph()
setfont(note.add_run("Sources: "), bold=True)
setfont(note.add_run(
    "English wording taken from the 2023 EIM codebook "
    "(codebook/EIM23datamap.xlsx, “Variable label” column). "
    "Estonian wording taken from the 2020 EIM SPSS file "
    "(data/EIM 2020_20.10.25.sav, variable labels, re-decoded from latin1 → utf-8). "
    "SPSS truncates variable labels at ≈80 characters, so some Estonian "
    "items below are marked [truncated in SPSS]. For the full Estonian "
    "instrument, consult the Estonian-language survey document."
))
note.paragraph_format.space_after = Pt(8)

# Direction reminder
note2 = doc.add_paragraph()
setfont(note2.add_run("Coding direction: "), bold=True)
setfont(note2.add_run(
    "see CLAUDE.md “Composite Variables — Quick Reference” for the original "
    "scale direction of each item battery, and for the items that are "
    "reverse-coded or inverted at the composite level. The wording below is "
    "the raw item as administered; coding transformations are applied later "
    "in the analysis pipeline."
))
note2.paragraph_format.space_after = Pt(14)


def add_table(items):
    """items: list of (item_2023, item_2020, gloss)."""
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    # Header
    hdr = table.rows[0].cells
    for i, h in enumerate(["2023 code", "2020 code", "English (codebook)", "Estonian (SPSS label)"]):
        p = hdr[i].paragraphs[0]
        setfont(p.add_run(h), bold=True)
    # Column widths
    widths_cm = [1.6, 1.6, 7.5, 6.5]
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)

    for code23, code20, gloss in items:
        row = table.add_row().cells
        setfont(row[0].paragraphs[0].add_run(code23 or "—"))
        setfont(row[1].paragraphs[0].add_run(code20 or "—"))
        # English text
        en_text = EN.get(code23, "")
        if not en_text and code20:
            en_text = EN.get(code20, "")
        en_para = row[2].paragraphs[0]
        if gloss:
            setfont(en_para.add_run(f"[{gloss}] "),
                    italic=True, color=RGBColor(0x66, 0x66, 0x66))
        setfont(en_para.add_run(en_text or "—"))
        # Estonian text (from 2020 SPSS file; some labels are truncated)
        ee_text = fmt_ee(EE.get(code20, "")) if code20 else "—"
        setfont(row[3].paragraphs[0].add_run(ee_text))


for var_name, intro, items in VARIABLES:
    h = doc.add_heading(var_name, level=1)
    for run in h.runs:
        run.font.name = FONT
        run.font.size = SIZE
        run.bold = True
    if intro:
        p = doc.add_paragraph()
        setfont(p.add_run(intro), italic=True, color=RGBColor(0x55, 0x55, 0x55))
    add_table(items)
    doc.add_paragraph()  # spacing

# Save
out = ROOT / "reports" / "Question_Wording_EN_EE.docx"
out.parent.mkdir(exist_ok=True)
doc.save(out)
print(f"Saved: {out}")
