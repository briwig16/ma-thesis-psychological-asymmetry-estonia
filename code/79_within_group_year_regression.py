"""
Within-group OLS regression of each composite on year (0 = 2020, 1 = 2023),
run separately for Estonian and Russian respondents.

This is the exact decomposition method used in Bilali, Çelik & Ok (2014) to
unpack significant Stage × Ethnicity interactions from their ANOVA.

For each ethnic group:
    composite = B0 + B1 * year_2023_indicator + ε

B1 represents the unstandardized mean shift on the original Likert metric
from 2020 to 2023 for that group. When the predictor is binary (0/1), B1 is
mathematically identical to the raw mean difference (M_2023 - M_2020) for
that group, and the t-test on B1 is equivalent to the within-group
independent-samples t-test.

Outputs:
    code/_within_group_year_regression.tsv
    reports/Within_Group_Year_Regression.docx  (narrative + backup table)
"""

from pathlib import Path

import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent.parent

# ---------- Load -----------------------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)
df23["__year__"] = 2023

df20, _ = pyreadstat.read_sav(
    str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"), encoding="latin1"
)
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)
df20["__year__"] = 2020


def clean(df, items, reverse_items=None, scale_max=None, dk_code=9):
    sub = df[items].apply(pd.to_numeric, errors="coerce")
    sub = sub.where(sub != dk_code)
    if reverse_items:
        for it in reverse_items:
            sub[it] = (scale_max + 1) - sub[it]
    return sub.mean(axis=1, skipna=True)


def build_group_year(spec, grp_code, year):
    df = df23 if year == 2023 else df20
    sub_df = df[df["ethnicity_binary"] == grp_code]
    if isinstance(spec[f"items_{year}"], dict):
        # group-specific items (e.g., SD: Primary)
        grp_key = "E" if grp_code == 0 else "R"
        items = spec[f"items_{year}"][grp_key]
        rev = (spec[f"rev_{year}"][grp_key]
               if spec.get(f"rev_{year}") else None)
    else:
        items = spec[f"items_{year}"]
        rev = spec.get(f"rev_{year}")
    vals = clean(sub_df, items, rev, spec["smax"])
    if spec["inv"]:
        vals = (spec["smax"] + 1) - vals
    return vals.dropna().reset_index(drop=True)


SPECS = [
    {"name": "Superordinate Identity", "smax": 4, "inv": True,
     "items_2023": ["Q67_2","Q67_4","Q67_5"], "rev_2023": ["Q67_4"],
     "items_2020": ["K6X5_2","K6X5_3","K6X5_4"], "rev_2020": ["K6X5_3"]},
    {"name": "SD: Primary Out-group", "smax": 5, "inv": False,
     "items_2023": {"E": ["Q57_1","Q58_1","Q59_1"], "R": ["Q57_2","Q58_2","Q59_2"]},
     "items_2020": {"E": ["K4X7_1","K4X8_1","K4X9_1"], "R": ["K4X7_2","K4X8_2","K4X9_2"]}},
    {"name": "SD: General Out-group", "smax": 5, "inv": False,
     "items_2023": ["Q57_4","Q57_5","Q58_4","Q58_5","Q59_4","Q59_5"],
     "items_2020": ["K4X7_3","K4X8_3","K4X9_3"]},
    {"name": "Comparative Opportunity Assessment", "smax": 5, "inv": True,
     "items_2023": [f"Q44_{i}" for i in range(1,13)],
     "items_2020": [f"K3X1_{i}" for i in range(1,13)]},
    {"name": "Belief in Inevitable Conflict", "smax": 4, "inv": False,
     "items_2023": ["Q63_1","Q63_2","Q63_3","Q63_4"], "rev_2023": ["Q63_1","Q63_2"],
     "items_2020": ["K6X1_1","K6X1_2","K6X1_3","K6X1_4"], "rev_2020": ["K6X1_1","K6X1_2"]},
    {"name": "Minority Inclusion Support", "smax": 4, "inv": True,
     "items_2023": ["Q68_1","Q68_2","Q68_3"],
     "items_2020": ["K6X6_1","K6X6_2","K6X6_3"]},
]


# ---------- Run regressions per group ------------------------------------
rows = []
for spec in SPECS:
    for grp_code, grp_name in [(0, "Estonian"), (1, "Russian")]:
        s20 = build_group_year(spec, grp_code, 2020)
        s23 = build_group_year(spec, grp_code, 2023)
        long = pd.DataFrame({
            "value": pd.concat([s20, s23], ignore_index=True),
            "year_2023": [0]*len(s20) + [1]*len(s23),
        }).dropna()
        X = sm.add_constant(long["year_2023"])
        model = sm.OLS(long["value"], X).fit()
        B0   = model.params["const"]
        B1   = model.params["year_2023"]
        se   = model.bse["year_2023"]
        t    = model.tvalues["year_2023"]
        p    = model.pvalues["year_2023"]
        ll, ul = model.conf_int().loc["year_2023"]
        rows.append({
            "variable": spec["name"],
            "group": grp_name,
            "n_2020": len(s20), "n_2023": len(s23),
            "B0_intercept": B0, "B1_year2023": B1,
            "SE": se, "CI_low": ll, "CI_high": ul,
            "t": t, "df": int(model.df_resid), "p": p,
            "R2": model.rsquared,
        })


out = pd.DataFrame(rows)
tsv_path = ROOT / "code" / "_within_group_year_regression.tsv"
out.to_csv(tsv_path, sep="\t", index=False, float_format="%.4f")
print(f"Saved TSV: {tsv_path}\n")


def stars(p):
    return "***" if p < .001 else "**" if p < .01 else "*" if p < .05 else ("⁺" if p < .10 else "ns")


def fmt_p(p):
    if p < .001:  return "< .001"
    if p < .0001: return "< .0001"
    return f"{p:.3f}".lstrip("0") or "0.000"


# ---------- Print summary -------------------------------------------------
print(f"{'Composite':<37}{'Group':<10}{'B1':>9}{'SE':>8}{'95% CI':>22}{'t':>10}{'p':>10}")
print("=" * 106)
for r in rows:
    print(f"{r['variable']:<37}{r['group']:<10}"
          f"{r['B1_year2023']:>+9.3f}{r['SE']:>8.3f}"
          f"  [{r['CI_low']:>+6.3f}, {r['CI_high']:>+6.3f}]"
          f"{r['t']:>+10.2f}{fmt_p(r['p']):>10} {stars(r['p'])}")


# ============================================================================
#                               WORD DOC
# ============================================================================
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def H(text, size=14, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    if color: r.font.color.rgb = color


def body(text, italic=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic
    return p


# ----- Title -----
H("Within-Group OLS Regression of Composite on Year",
  size=15, color=RGBColor(0x1f, 0x29, 0x37))

body(
    "Following the decomposition method used in Bilali, Çelik & Ok (2014, p. 7), "
    "a separate ordinary least squares regression was fitted within each ethnic "
    "group, with year (0 = 2020, 1 = 2023) as the sole predictor of the composite "
    "score:",
    italic=True,
)

eq = doc.add_paragraph()
eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = eq.add_run("composite = B₀ + B₁ × Year_2023 + ε")
r.italic = True; r.font.size = Pt(12)

body(
    "B₁ represents the unstandardized mean shift on the original Likert metric "
    "from 2020 to 2023 for that group. Because year is binary (0/1), B₁ is "
    "mathematically identical to the raw mean difference (M₂₀₂₃ − M₂₀₂₀) for "
    "that group, and its t-test is equivalent to an independent-samples t-test. "
    "This decomposition unpacks the significant Year × Ethnicity interactions "
    "from the 2 × 2 ANOVA (Table 1 of ANOVA_2x2_Bilali_Style.docx) by reporting "
    "which group changed, in which direction, by how much.",
    italic=True,
)

doc.add_paragraph()


# ----- Narrative section (Bilali style) ----------------------------------
H("Narrative decomposition", size=12)

def make_inline(r, group_label):
    """Build the inline-statistic string like 'B = +0.42, SE = 0.15, p = .005'."""
    p_val = r["p"]
    if p_val < .001:
        p_str = "p < .001"
    else:
        p_str = f"p = {p_val:.3f}".replace("0.", ".")
    return (f"B = {r['B1_year2023']:+.3f}, SE = {r['SE']:.3f}, {p_str}")

# Estonian respondents
sig_est = [r for r in rows if r["group"] == "Estonian" and r["p"] < .05]
ns_est  = [r for r in rows if r["group"] == "Estonian" and r["p"] >= .05]

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
run = p.add_run("Estonian respondents. ")
run.bold = True; run.font.size = Pt(11)

if sig_est:
    pieces = []
    for r in sig_est:
        sign_word = "higher" if r["B1_year2023"] > 0 else "lower"
        pieces.append(f"{sign_word} on {r['variable']} ({make_inline(r, 'Estonian')})")
    run = p.add_run("During 2023, Estonian respondents scored " + "; ".join(pieces) + " compared to 2020. ")
    run.font.size = Pt(11)

if ns_est:
    ns_pieces = []
    for r in ns_est:
        ns_pieces.append(f"{r['variable']} ({make_inline(r, 'Estonian')})")
    run = p.add_run("Year was not a significant predictor for: " + "; ".join(ns_pieces) + ".")
    run.font.size = Pt(11)


# Russian respondents
sig_rus = [r for r in rows if r["group"] == "Russian" and r["p"] < .05]
ns_rus  = [r for r in rows if r["group"] == "Russian" and r["p"] >= .05]

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
run = p.add_run("Russian respondents. ")
run.bold = True; run.font.size = Pt(11)

if sig_rus:
    pieces = []
    for r in sig_rus:
        sign_word = "higher" if r["B1_year2023"] > 0 else "lower"
        pieces.append(f"{sign_word} on {r['variable']} ({make_inline(r, 'Russian')})")
    run = p.add_run("During 2023, Russian respondents scored " + "; ".join(pieces) + " compared to 2020. ")
    run.font.size = Pt(11)

if ns_rus:
    ns_pieces = []
    for r in ns_rus:
        ns_pieces.append(f"{r['variable']} ({make_inline(r, 'Russian')})")
    run = p.add_run("Year was not a significant predictor for: " + "; ".join(ns_pieces) + ".")
    run.font.size = Pt(11)


doc.add_paragraph()


# ----- Backup table ------------------------------------------------------
H("Backup table — full regression statistics", size=12)
body(
    "Complete regression output for each composite × ethnic group. Bilali et al. "
    "did not include this table in their paper, reporting these statistics only "
    "inline; included here for reference.",
    italic=True, size=10,
)
doc.add_paragraph()

header = ["Variable", "Group", "n (2020 / 2023)", "B₁", "SE",
          "95% CI on B₁", "t (df)", "p", "Sig"]
table = doc.add_table(rows=1, cols=len(header))
table.style = "Light Grid Accent 1"
for i, h in enumerate(header):
    c = table.rows[0].cells[i]; c.text = ""
    rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)

for r in rows:
    row = table.add_row().cells
    cells_data = [
        r["variable"],
        r["group"],
        f"{r['n_2020']} / {r['n_2023']}",
        f"{r['B1_year2023']:+.3f}",
        f"{r['SE']:.3f}",
        f"[{r['CI_low']:+.3f}, {r['CI_high']:+.3f}]",
        f"{r['t']:+.2f} ({r['df']})",
        fmt_p(r["p"]),
        stars(r["p"]),
    ]
    for i, val in enumerate(cells_data):
        c = row[i]; c.text = ""
        rr = c.paragraphs[0].add_run(val); rr.font.size = Pt(9)
        if i in (3, 4, 5, 6, 7, 8):
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


doc.add_paragraph()
body(
    "Note. B₀ = intercept (mean composite score at 2020). B₁ = regression "
    "coefficient for year (mean shift from 2020 to 2023, on the original "
    "Likert metric after any scale-inversion noted in CLAUDE.md). SE = "
    "standard error of B₁. t and df: from the OLS regression. p = two-sided "
    "p-value testing B₁ ≠ 0. Significance: *** p < .001, ** p < .01, "
    "* p < .05, ⁺ p < .10, ns = not significant. Following Bilali et al. "
    "(2014), demographic controls were not added; their footnote 3 notes "
    "that adding age, gender, education, and income did not change their "
    "main conclusions.",
    italic=True, size=8,
)

out_doc = ROOT / "reports" / "Within_Group_Year_Regression.docx"
doc.save(out_doc)
print(f"\nSaved Word doc: {out_doc}")
