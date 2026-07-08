"""
Contact-Theory Regression Check
================================
Targeted, hypothesis-driven test of classical contact theory (Allport 1954;
Pettigrew & Tropp 2006): does within-group reported contact with the primary
out-group predict reduced social distance from that out-group?

Four regressions, one per group × year cell:
  - Russian 2020:    SD from Estonian-speakers ~ Contact with Estonian-speakers
  - Russian 2023:    SD from Estonian-speakers ~ Contact with Estonian-speakers
  - Estonian 2020:   SD from Russian-speakers  ~ Contact with Russian-speakers
  - Estonian 2023:   SD from Russian-speakers  ~ Contact with Russian-speakers

Each regression is a single-equation linear model. No covariates. The point
is to assess the bivariate contact-theory association cross-sectionally
within each group × year cell.

Direction handling:
  - Contact scale (Q51 / Q52, K4X1 / K4X2): raw 1 = almost every day,
    5 = have not communicated. Inverted via (6 − raw) so HIGHER = MORE contact.
  - SD scale: raw 1 = no distance, 5 = maximal distance. NOT inverted —
    higher = more distance.
  - Contact theory predicts NEGATIVE slope: more contact → less distance.

Outputs:
  Console: coefficient, SE, 95% CI, t, p, R², N for each regression
  reports/Contact_Theory_Regression.docx — formatted Word doc with table
  viz/fig_contact_theory_regression.jpg — 2×2 scatter+fit panels (300 DPI)
"""

from pathlib import Path
import math
import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, RGBColor

ROOT = Path(__file__).parent.parent

# ---------- Load data --------------------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy()

df20, _ = pyreadstat.read_sav(str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"),
                              encoding="latin1")
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy()


def to_num(s, dk=9):
    s = pd.to_numeric(s, errors="coerce")
    return s.where(s != dk).astype(float)


def composite_pairwise(df, items, dk=9):
    """Pairwise-deletion composite (skipna=True mean across items)."""
    sub = df[items].apply(pd.to_numeric, errors="coerce").where(lambda x: x != dk)
    return sub.mean(axis=1, skipna=True)


# ---------- Build SD: Primary + Contact composites for each cell -----------
# Russians: contact with / SD from Estonian-speakers
# Estonians: contact with / SD from Russian-speakers

# 2023: Q-codes
SD_RUS_2023      = ["Q57_2", "Q58_2", "Q59_2"]
SD_EST_2023      = ["Q57_1", "Q58_1", "Q59_1"]
CONTACT_RUS_2023 = [f"Q51_{i}" for i in range(1, 7)]   # contact with Estonian-speakers
CONTACT_EST_2023 = [f"Q52_{i}" for i in range(1, 7)]   # contact with Russian-speakers

# 2020: K-codes
SD_RUS_2020      = ["K4X7_2", "K4X8_2", "K4X9_2"]
SD_EST_2020      = ["K4X7_1", "K4X8_1", "K4X9_1"]
CONTACT_RUS_2020 = [f"K4X1_{i}" for i in range(1, 7)]
CONTACT_EST_2020 = [f"K4X2_{i}" for i in range(1, 7)]


def build_pair(df, eth_val, sd_items, contact_items):
    sub = df[df["ethnicity_binary"] == eth_val]
    sd      = composite_pairwise(sub, sd_items)
    contact = composite_pairwise(sub, contact_items)
    # Invert contact so higher = more contact
    contact_inv = 6 - contact
    out = pd.DataFrame({"SD": sd, "Contact": contact_inv}).dropna()
    return out


CELLS = [
    ("Russian 2020",   build_pair(df20, 1, SD_RUS_2020, CONTACT_RUS_2020),
                       "SD from Estonian-speakers", "Contact with Estonian-speakers"),
    ("Russian 2023",   build_pair(df23, 1, SD_RUS_2023, CONTACT_RUS_2023),
                       "SD from Estonian-speakers", "Contact with Estonian-speakers"),
    ("Estonian 2020",  build_pair(df20, 0, SD_EST_2020, CONTACT_EST_2020),
                       "SD from Russian-speakers",  "Contact with Russian-speakers"),
    ("Estonian 2023",  build_pair(df23, 0, SD_EST_2023, CONTACT_EST_2023),
                       "SD from Russian-speakers",  "Contact with Russian-speakers"),
]


# ---------- Run regressions -------------------------------------------------
def run_regression(df_pair):
    X = sm.add_constant(df_pair["Contact"])
    y = df_pair["SD"]
    model = sm.OLS(y, X).fit()
    b0, b1 = model.params["const"], model.params["Contact"]
    se1 = model.bse["Contact"]
    ci_lo, ci_hi = model.conf_int().loc["Contact"]
    t = model.tvalues["Contact"]
    p = model.pvalues["Contact"]
    r2 = model.rsquared
    n = int(model.nobs)
    return {"b0": b0, "b1": b1, "se1": se1, "ci_lo": ci_lo, "ci_hi": ci_hi,
            "t": t, "p": p, "r2": r2, "n": n, "model": model}


def stars(p):
    if p < 0.001: return "***"
    if p < 0.01:  return "**"
    if p < 0.05:  return "*"
    return "ns"


print("=" * 92)
print("CONTACT-THEORY REGRESSION CHECK")
print("Model: SD: Primary Out-group ~ Contact (inverted: higher = more contact)")
print("Contact theory predicts NEGATIVE slope: more contact → less distance")
print("=" * 92)

results = []
for label, df_pair, sd_label, contact_label in CELLS:
    r = run_regression(df_pair)
    r["label"] = label
    r["sd_label"] = sd_label
    r["contact_label"] = contact_label
    results.append(r)
    print(f"\n  {label} ({contact_label} → {sd_label})")
    print(f"    n = {r['n']}")
    print(f"    Intercept (b0)  = {r['b0']:+.3f}")
    print(f"    Slope (b1)      = {r['b1']:+.4f}   SE = {r['se1']:.4f}   "
          f"95% CI [{r['ci_lo']:+.4f}, {r['ci_hi']:+.4f}]")
    print(f"    t = {r['t']:+.3f}   p = {r['p']:.4g}   {stars(r['p'])}")
    print(f"    R² = {r['r2']:.4f}")
    print(f"    Interpretation: each +1 unit on contact (1 unit ≈ 'more frequent contact category')")
    print(f"                     associated with {r['b1']:+.4f} units on SD scale (1–5)")


# ============================================================================
#  Word doc — formatted regression table
# ============================================================================
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(2.0)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)


def H(text, size=14, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    if color: r.font.color.rgb = color


def body(text, italic=False, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size)
    if italic: r.italic = True


H("Contact-Theory Regression Check", size=14,
  color=RGBColor(0x1f, 0x29, 0x37))

body(
    "Targeted bivariate regression of social distance from the primary out-group on "
    "frequency of contact with that out-group, run separately for each group × year cell. "
    "Tests the classical contact-theory prediction (Allport, 1954; Pettigrew & Tropp, 2006) "
    "that more contact is associated with less social distance.",
    italic=True
)
doc.add_paragraph()

# Table
header = ["Group × Year", "n", "b₁ (slope)", "SE", "95% CI", "t", "p", "R²"]
table = doc.add_table(rows=1, cols=len(header))
table.style = "Light Grid Accent 1"
for i, h in enumerate(header):
    cell = table.rows[0].cells[i]; cell.text = ""
    rr = cell.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)


def fmt_p(p):
    if p < 0.001: return "< .001"
    return f"{p:.3f}".lstrip("0") if p < 1 else f"{p:.3f}"


def fmt_signed(x, dec=4):
    s = f"{x:+.{dec}f}"
    return s.replace("-", "−")


for r in results:
    cells = table.add_row().cells
    cells[0].text = r["label"]
    cells[1].text = str(r["n"])
    cells[2].text = f"{fmt_signed(r['b1'])} {stars(r['p'])}"
    cells[3].text = f"{r['se1']:.4f}"
    cells[4].text = f"[{fmt_signed(r['ci_lo'])}, {fmt_signed(r['ci_hi'])}]"
    cells[5].text = fmt_signed(r["t"], dec=2)
    cells[6].text = fmt_p(r["p"])
    cells[7].text = f"{r['r2']:.3f}"
    for c in cells:
        for para in c.paragraphs:
            for run in para.runs:
                run.font.size = Pt(9)


doc.add_paragraph()
body(
    "Note. Each row reports an OLS regression of the SD: Primary Out-group composite "
    "(higher = more social distance) on the Contact composite (inverted via 6 − raw, so "
    "higher = more frequent contact). For Russian respondents, the primary out-group is "
    "Estonian-speakers; for Estonian respondents, Russian-speakers. b₁ is the unstandardised "
    "slope coefficient: a one-unit increase in contact frequency is associated with the "
    "indicated change on the 1–5 SD scale. Contact theory predicts a negative slope (more "
    "contact → less distance). Listwise deletion: respondents missing on either composite "
    "are excluded. Cross-sectional design: associations are not causal. Significance: "
    "*** p < .001, ** p < .01, * p < .05, ns = not significant.",
    italic=True, size=8
)

out_doc = ROOT / "reports" / "Contact_Theory_Regression.docx"
doc.save(out_doc)
print(f"\nSaved Word doc: {out_doc}")


# ============================================================================
#  Figure — 2×2 scatter+fit panels
# ============================================================================
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size":   9,
    "axes.edgecolor": "#333",
    "axes.linewidth": 0.6,
})

RUS = "#d97706"; EST = "#2563eb"

fig, axes = plt.subplots(2, 2, figsize=(11.0, 9.0), dpi=300, sharex=True, sharey=True)

panel_config = [
    (results[0], CELLS[0][1], RUS, axes[0, 0]),
    (results[1], CELLS[1][1], RUS, axes[0, 1]),
    (results[2], CELLS[2][1], EST, axes[1, 0]),
    (results[3], CELLS[3][1], EST, axes[1, 1]),
]

for r, df_pair, color, ax in panel_config:
    # Scatter with jitter (responses are integer-mean composites, so jitter clarifies density)
    rng = np.random.default_rng(42)
    jx = rng.uniform(-0.04, 0.04, size=len(df_pair))
    jy = rng.uniform(-0.04, 0.04, size=len(df_pair))
    ax.scatter(df_pair["Contact"] + jx, df_pair["SD"] + jy,
               s=12, color=color, alpha=0.30, edgecolor="none", zorder=2)

    # Regression line
    x_grid = np.linspace(1, 5, 50)
    y_pred = r["b0"] + r["b1"] * x_grid
    ax.plot(x_grid, y_pred, color="#111", linewidth=1.6, zorder=3)

    # 95% CI band on the line
    pred = r["model"].get_prediction(sm.add_constant(pd.Series(x_grid, name="Contact")))
    ci = pred.summary_frame(alpha=0.05)
    ax.fill_between(x_grid, ci["mean_ci_lower"], ci["mean_ci_upper"],
                    color="#111", alpha=0.10, zorder=1)

    ax.set_title(r["label"], fontsize=11, fontweight="bold", loc="left", pad=8)
    ax.text(0.02, 0.97,
            f"b₁ = {r['b1']:+.3f}  ({stars(r['p'])})\n"
            f"95% CI [{r['ci_lo']:+.3f}, {r['ci_hi']:+.3f}]\n"
            f"R² = {r['r2']:.3f}   n = {r['n']}",
            transform=ax.transAxes,
            ha="left", va="top",
            fontsize=8.5, color="#222",
            family="monospace",
            bbox=dict(boxstyle="round,pad=0.4", facecolor="white",
                      edgecolor="#888", linewidth=0.5, alpha=0.92))

    ax.set_xlim(0.7, 5.3)
    ax.set_ylim(0.7, 5.3)
    ax.set_xticks([1, 2, 3, 4, 5])
    ax.set_yticks([1, 2, 3, 4, 5])
    for s_ in ("top", "right"):
        ax.spines[s_].set_visible(False)
    ax.tick_params(axis="both", labelsize=8, color="#888")

# Axis labels — only on edges
axes[1, 0].set_xlabel("Contact composite (inverted: higher = more frequent contact)",
                       fontsize=9, color="#444")
axes[1, 1].set_xlabel("Contact composite (inverted: higher = more frequent contact)",
                       fontsize=9, color="#444")
axes[0, 0].set_ylabel("SD: Primary Out-group composite (higher = more distance)",
                       fontsize=9, color="#444")
axes[1, 0].set_ylabel("SD: Primary Out-group composite (higher = more distance)",
                       fontsize=9, color="#444")

# Title and footer
fig.text(0.04, 0.965,
         "Contact-Theory Regression Check — SD: Primary Out-group ~ Contact",
         fontsize=13, fontweight="bold", ha="left")
fig.text(0.04, 0.940,
         "Top row: Russian respondents rating Estonian-speakers (orange).  "
         "Bottom row: Estonian respondents rating Russian-speakers (blue).",
         fontsize=9, color="#444")
fig.text(0.04, 0.922,
         "Contact theory predicts a NEGATIVE slope: more contact → less social distance.",
         fontsize=9, color="#444")

fig.text(0.04, 0.018,
         "Each dot = one respondent (jittered ±0.04 for legibility). Black line = OLS regression fit. Shaded band = 95% CI on the conditional mean.",
         fontsize=7, color="#555")
fig.text(0.04, 0.005,
         "Cross-sectional associations; no causal interpretation. EIM is a repeated cross-section, not a panel.",
         fontsize=7, color="#555")

plt.subplots_adjust(left=0.07, right=0.985, top=0.88, bottom=0.085,
                    wspace=0.10, hspace=0.30)

out_jpg = ROOT / "viz" / "fig_contact_theory_regression.jpg"
plt.savefig(out_jpg, dpi=300, format="jpg",
            facecolor="white", pil_kwargs={"quality": 95})
print(f"Saved chart: {out_jpg}")
