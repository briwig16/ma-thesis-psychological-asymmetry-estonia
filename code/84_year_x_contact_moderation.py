"""
Option B — Contact × Year moderation analysis.

For each composite × ethnic group, fit:
    composite = B0 + B1 * year_2023 + B2 * contact_centered + B3 * (year_2023 × contact_centered) + ε

where:
  - year_2023: binary 0/1 (0 = 2020, 1 = 2023)
  - contact: out-group contact composite, sign-inverted so higher = more contact
    (Estonian respondents → Contact: Russian Speakers; Russian respondents → Contact: Estonian Speakers)
  - contact_centered: contact minus the within-group sample mean. Centering means:
      B1 = year effect at AVERAGE contact level (a meaningful main effect)
      B2 = contact slope at 2020 (or pooled if we want — the interpretation is local)
      B3 = how much the year effect changes per +1 unit of contact

Tests the classical Pettigrew–Tropp moderation hypothesis: does the 2020→2023
shift depend on respondents' level of out-group contact? A NEGATIVE B3 on
"distance-direction" composites (e.g., SD: Primary) would mean high-contact
respondents shifted LESS than low-contact respondents — contact buffered the
post-2022 hardening. A POSITIVE B3 on "belonging-direction" composites would
mean high-contact respondents protected belonging better.

Outputs:
  code/_year_x_contact_moderation.tsv
  reports/Year_x_Contact_Moderation.docx
  viz/fig_year_x_contact_moderation.jpg
"""

import math
from pathlib import Path

import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm
import matplotlib.pyplot as plt
import matplotlib.lines as mlines
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent.parent

# ---------- Load -----------------------------------------------------------
df23 = pd.read_csv(ROOT / "data" / "EIM23.csv")
df23 = df23[df23["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)

df20, _ = pyreadstat.read_sav(
    str(ROOT / "data" / "EIM 2020_20.10.25.sav copy"), encoding="latin1"
)
df20["ethnicity_binary"] = df20.apply(
    lambda r: 0 if r.get("T9_1") == 1 else (1 if r.get("T9_2") == 1 else None),
    axis=1,
)
df20 = df20[df20["ethnicity_binary"].isin([0, 1])].copy().reset_index(drop=True)


def clean(df, items, reverse_items=None, scale_max=None, dk_code=9):
    sub = df[items].apply(pd.to_numeric, errors="coerce")
    sub = sub.where(sub != dk_code)
    if reverse_items:
        for it in reverse_items:
            sub[it] = (scale_max + 1) - sub[it]
    return sub.mean(axis=1, skipna=True)


# Out-group contact: Estonians' contact = Q52/K4X2 (with Russian speakers, the out-group)
# Russians' contact = Q51/K4X1 (with Estonian speakers, the out-group)
CONTACT_2023 = {0: [f"Q52_{i}" for i in range(1, 7)], 1: [f"Q51_{i}" for i in range(1, 7)]}
CONTACT_2020 = {0: [f"K4X2_{i}" for i in range(1, 7)], 1: [f"K4X1_{i}" for i in range(1, 7)]}


def build_cell(spec, df, year, grp_code):
    """Return a DataFrame with composite value, contact (inverted: higher = more contact),
    and year dummy for this cell."""
    if isinstance(spec[f"items_{year}"], dict):
        grp_key = "E" if grp_code == 0 else "R"
        items = spec[f"items_{year}"][grp_key]
        rev = (spec[f"rev_{year}"][grp_key]
               if spec.get(f"rev_{year}") else None)
    else:
        items = spec[f"items_{year}"]
        rev = spec.get(f"rev_{year}")

    sub_df = df[df["ethnicity_binary"] == grp_code]
    composite_vals = clean(sub_df, items, rev, spec["smax"])
    if spec["inv"]:
        composite_vals = (spec["smax"] + 1) - composite_vals

    # Out-group contact (raw scale: 1 = almost every day, 5 = no communication)
    contact_items = (CONTACT_2023[grp_code] if year == 2023 else CONTACT_2020[grp_code])
    contact_vals = clean(sub_df, contact_items)
    # Invert so higher = more contact
    contact_inv = 6 - contact_vals

    out = pd.DataFrame({
        "value":    composite_vals.values,
        "contact":  contact_inv.values,
        "year_2023": 1 if year == 2023 else 0,
    })
    return out


SPECS = [
    {"name": "Superordinate Identity", "smax": 4, "inv": True,
     "items_2023": ["Q67_2","Q67_4","Q67_5"], "rev_2023": ["Q67_4"],
     "items_2020": ["K6X5_2","K6X5_3","K6X5_4"], "rev_2020": ["K6X5_3"]},
    {"name": "SD: Primary Out-group", "smax": 5, "inv": False,
     "items_2023": {"E": ["Q57_1","Q58_1","Q59_1"], "R": ["Q57_2","Q58_2","Q59_2"]},
     "items_2020": {"E": ["K4X7_1","K4X8_1","K4X9_1"], "R": ["K4X7_2","K4X8_2","K4X9_2"]}},
    {"name": "SD: General Out-group", "smax": 5, "inv": False,
     "items_2023": ["Q57_4","Q57_5","Q58_4","Q58_5","Q59_4","Q59_5"],
     "items_2020": ["K4X7_3","K4X8_3","K4X9_3"]},
    {"name": "Comparative Opportunity Assessment", "smax": 5, "inv": True,
     "items_2023": [f"Q44_{i}" for i in range(1,13)],
     "items_2020": [f"K3X1_{i}" for i in range(1,13)]},
    {"name": "Belief in Inevitable Conflict", "smax": 4, "inv": False,
     "items_2023": ["Q63_1","Q63_2","Q63_3","Q63_4"], "rev_2023": ["Q63_1","Q63_2"],
     "items_2020": ["K6X1_1","K6X1_2","K6X1_3","K6X1_4"], "rev_2020": ["K6X1_1","K6X1_2"]},
    {"name": "Minority Inclusion Support", "smax": 4, "inv": True,
     "items_2023": ["Q68_1","Q68_2","Q68_3"],
     "items_2020": ["K6X6_1","K6X6_2","K6X6_3"]},
]

# Reference: B1 from year-only model (script 79)
year_only = pd.read_csv(ROOT / "code" / "_within_group_year_regression.tsv", sep="\t")


def stars(p):
    return "***" if p < .001 else "**" if p < .01 else "*" if p < .05 else ("⁺" if p < .10 else "ns")


def fmt_p(p):
    if pd.isna(p): return "—"
    if p < .001: return "< .001"
    return f"{p:.3f}".lstrip("0")


rows = []
for spec in SPECS:
    for grp_code, grp_name in [(0, "Estonian"), (1, "Russian")]:
        s20 = build_cell(spec, df20, 2020, grp_code)
        s23 = build_cell(spec, df23, 2023, grp_code)
        long = pd.concat([s20, s23], ignore_index=True).dropna()

        # Mean-center contact within group
        long["contact_c"] = long["contact"] - long["contact"].mean()
        long["yearXcontact"] = long["year_2023"] * long["contact_c"]

        X = sm.add_constant(long[["year_2023", "contact_c", "yearXcontact"]])
        m = sm.OLS(long["value"], X).fit(cov_type="HC3")

        # Reference unadjusted B1
        u = year_only[(year_only["variable"] == spec["name"]) &
                      (year_only["group"] == grp_name)].iloc[0]

        rows.append({
            "variable": spec["name"], "group": grp_name,
            "N": int(m.nobs),
            "contact_mean": long["contact"].mean(),
            "contact_sd":   long["contact"].std(ddof=1),
            "B1_unadj":   u["B1_year2023"], "p_unadj": u["p"],
            "B1_year":    m.params["year_2023"],
            "SE_year":    m.bse["year_2023"],
            "p_year":     m.pvalues["year_2023"],
            "B2_contact": m.params["contact_c"],
            "SE_contact": m.bse["contact_c"],
            "p_contact":  m.pvalues["contact_c"],
            "B3_interaction": m.params["yearXcontact"],
            "SE_interaction": m.bse["yearXcontact"],
            "CI_low_interaction":  m.conf_int().loc["yearXcontact"][0],
            "CI_high_interaction": m.conf_int().loc["yearXcontact"][1],
            "p_interaction": m.pvalues["yearXcontact"],
            "R2": m.rsquared,
        })

out = pd.DataFrame(rows)
tsv = ROOT / "code" / "_year_x_contact_moderation.tsv"
out.to_csv(tsv, sep="\t", index=False, float_format="%.4f")
print(f"Saved TSV: {tsv}\n")

# ---- Print summary ----
print(f"{'Composite':<37}{'Group':<10}{'B1 year':>10}{'B2 contact':>12}{'B3 inter':>12}{'p inter':>10}")
print("=" * 95)
for r in rows:
    print(f"{r['variable']:<37}{r['group']:<10}"
          f"{r['B1_year']:>+10.3f}{r['B2_contact']:>+12.3f}"
          f"{r['B3_interaction']:>+12.4f}{fmt_p(r['p_interaction']):>10} {stars(r['p_interaction'])}")


# ============================================================================
#                               WORD DOC
# ============================================================================
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(2.0)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)


def H(text, size=14, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    if color: r.font.color.rgb = color


def body(text, italic=False, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic


H("Year × Contact Moderation — Within-Group Regression",
  size=15, color=RGBColor(0x1f, 0x29, 0x37))

body(
    "For each composite × ethnic group, the within-group regression is extended "
    "to include out-group contact and a Year × Contact interaction. This tests "
    "the classical Pettigrew–Tropp moderation hypothesis: does the 2020→2023 "
    "shift depend on respondents' level of contact with the other ethnic group? "
    "Out-group contact is sign-inverted (6 − raw) so higher = more frequent "
    "contact. Contact is mean-centered within each group, so B₁ (year main "
    "effect) is interpreted at average contact level.",
    italic=True
)

eq = doc.add_paragraph()
eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = eq.add_run("composite = B₀ + B₁ × Year_2023 + B₂ × Contact_centered "
                "+ B₃ × (Year × Contact_c) + ε")
r.italic = True; r.font.size = Pt(11)

body(
    "Interpretation of B₃. A NEGATIVE B₃ on a 'higher = more negative attitude' "
    "composite (e.g., SD: Primary, Belief in Conflict) means high-contact "
    "respondents shifted LESS post-2022 — contact buffered the hardening. A "
    "POSITIVE B₃ on a 'higher = positive attitude' composite (e.g., Superordinate "
    "Identity, Minority Inclusion Support) means high-contact respondents "
    "PROTECTED the positive attitude over time. The opposite signs indicate "
    "contact AMPLIFIED rather than buffered the shift.",
    italic=True, size=9,
)

doc.add_paragraph()


H("Table 6. Year × Contact moderation regressions", size=12)

header = ["Variable", "Group", "N",
          "B₁ (year)", "B₂ (contact)", "B₃ (year × contact)",
          "95% CI on B₃", "p (B₃)", "Sig", "R²"]
table = doc.add_table(rows=1, cols=len(header))
table.style = "Light Grid Accent 1"
for i, h in enumerate(header):
    c = table.rows[0].cells[i]; c.text = ""
    rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)

for r in rows:
    row = table.add_row().cells
    cells = [
        r["variable"], r["group"], str(r["N"]),
        f"{r['B1_year']:+.3f} ({r['SE_year']:.3f}) {stars(r['p_year'])}",
        f"{r['B2_contact']:+.3f} ({r['SE_contact']:.3f}) {stars(r['p_contact'])}",
        f"{r['B3_interaction']:+.4f} ({r['SE_interaction']:.4f})",
        f"[{r['CI_low_interaction']:+.4f}, {r['CI_high_interaction']:+.4f}]",
        fmt_p(r["p_interaction"]),
        stars(r["p_interaction"]),
        f"{r['R2']:.3f}",
    ]
    for i, val in enumerate(cells):
        c = row[i]; c.text = ""
        rr = c.paragraphs[0].add_run(val); rr.font.size = Pt(8.5)
        if i >= 3:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
body(
    "Note. Out-group contact = mean of 6 contact items (Q51/Q52 or K4X1/K4X2), "
    "inverted so higher = more frequent contact. Contact is mean-centered within "
    "each group's sample. Year_2023 = 1 for 2023 respondents, 0 for 2020. HC3 "
    "robust standard errors. *** p < .001, ** p < .01, * p < .05, ⁺ p < .10, "
    "ns = not significant.",
    italic=True, size=8,
)

# Narrative of significant interactions
doc.add_paragraph()
H("Significant moderation findings", size=12)

sig_rows = [r for r in rows if r["p_interaction"] < .10]
if not sig_rows:
    body(
        "No Year × Contact interaction reached even marginal significance (p < .10) "
        "in any of the 12 regressions. Contact level did not moderate the 2020→2023 "
        "shift on any composite for either ethnic group. This means the asymmetric-"
        "divergence patterns documented in the unadjusted within-group regressions "
        "applied uniformly across high- and low-contact respondents within each "
        "group — heightened contact neither buffered nor amplified the shifts.",
        italic=True
    )
else:
    body(
        f"Of the 12 regressions, {len(sig_rows)} showed a Year × Contact interaction at p < .10:",
        italic=True
    )
    for r in sig_rows:
        b = r["B3_interaction"]; direction = "negative" if b < 0 else "positive"
        body(
            f"{r['variable']} × {r['group']}: B₃ = {b:+.4f}, "
            f"95% CI [{r['CI_low_interaction']:+.4f}, {r['CI_high_interaction']:+.4f}], "
            f"p = {fmt_p(r['p_interaction'])}. A {direction} interaction means "
            f"high-contact respondents shifted {'less' if (b < 0) == (r['B1_year'] > 0) else 'more'} "
            f"than low-contact respondents over 2020 → 2023.",
            size=10
        )

out_doc = ROOT / "reports" / "Year_x_Contact_Moderation.docx"
doc.save(out_doc)
print(f"\nSaved Word doc: {out_doc}")


# ============================================================================
#                             FOREST PLOT
# ============================================================================
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 9,
    "axes.edgecolor": "#333",
    "axes.linewidth": 0.7,
})

RUS = "#d97706"
EST = "#2563eb"

COMPOSITES = [s["name"] for s in SPECS]
PRETTY = {
    "Superordinate Identity":             "Superordinate Identity",
    "SD: Primary Out-group":              "SD: Primary Out-group",
    "SD: General Out-group":              "SD: General Out-group ⁱ",
    "Comparative Opportunity Assessment": "Comparative Opportunity",
    "Belief in Inevitable Conflict":      "Belief in Inevitable Conflict",
    "Minority Inclusion Support":         "Minority Inclusion Support",
}

fig, ax = plt.subplots(figsize=(15.0, 8.0), dpi=300)

# Y-positions: composite blocks of 2 rows (Estonian above Russian)
y_positions = {}
y = 13.0
for comp in COMPOSITES:
    y_positions[(comp, "Estonian")] = y; y -= 1.0
    y_positions[(comp, "Russian")]  = y; y -= 1.6

y_min, y_max = min(y_positions.values()) - 0.6, max(y_positions.values()) + 0.7
for i, comp in enumerate(COMPOSITES):
    y_top = y_positions[(comp, "Estonian")] + 0.6
    y_bot = y_positions[(comp, "Russian")]  - 0.6
    if i % 2 == 0:
        ax.axhspan(y_bot, y_top, facecolor="#f7f7f7", zorder=0)

for r in rows:
    y = y_positions[(r["variable"], r["group"])]
    color = EST if r["group"] == "Estonian" else RUS
    ll, ul = r["CI_low_interaction"], r["CI_high_interaction"]
    B3 = r["B3_interaction"]

    # CI bracket
    ax.plot([ll, ul], [y, y], color=color, linewidth=2.2, alpha=0.85,
            solid_capstyle="round", zorder=2)
    for xx in (ll, ul):
        ax.plot([xx, xx], [y - 0.10, y + 0.10],
                color=color, linewidth=1.6, alpha=0.9, zorder=2)
    # Dot
    sig = r["p_interaction"] < .10
    if sig:
        ax.scatter([B3], [y], s=170, color=color,
                   edgecolor="white", linewidth=1.0, zorder=3)
    else:
        ax.scatter([B3], [y], s=160, facecolor="white",
                   edgecolor=color, linewidth=2.0, zorder=3)

    # Annotation
    ax.text(max(ul, 0) + 0.012, y,
            f"{r['group']:<8}  B₃ = {B3:+.4f}  "
            f"95% CI [{ll:+.4f}, {ul:+.4f}]  "
            f"p = {fmt_p(r['p_interaction'])} {stars(r['p_interaction'])}",
            ha="left", va="center", fontsize=8.3,
            color="#222", family="monospace")

ax.axvline(0, color="#666", linewidth=1.0, linestyle="-", zorder=1)

# Y-axis labels at composite midpoints
yticks, yticklabels = [], []
for comp in COMPOSITES:
    y_E = y_positions[(comp, "Estonian")]
    y_R = y_positions[(comp, "Russian")]
    yticks.append((y_E + y_R) / 2)
    yticklabels.append(PRETTY[comp])
ax.set_yticks(yticks)
ax.set_yticklabels(yticklabels, fontsize=10, fontweight="bold")
ax.set_ylim(y_min, y_max)
ax.invert_yaxis()

# X-axis range
all_lo = min(r["CI_low_interaction"] for r in rows)
all_hi = max(r["CI_high_interaction"] for r in rows)
ax.set_xlim(all_lo - 0.02, all_hi + 0.20)
ax.set_xlabel("B₃ — interaction coefficient (Year × Out-group Contact)",
              fontsize=10, color="#444", labelpad=8)

for s in ("top", "right"):
    ax.spines[s].set_visible(False)
ax.spines["left"].set_color("#ccc")
ax.spines["bottom"].set_color("#666")
ax.tick_params(axis="x", length=2.5, color="#666", labelsize=9)
ax.tick_params(axis="y", length=0)

# Title
fig.text(0.04, 0.965,
         "Year × Out-group Contact Interaction — Forest Plot",
         fontsize=14, fontweight="bold", ha="left")
fig.text(0.04, 0.940,
         "Does the size of the 2020 → 2023 within-group shift depend on respondents' level of out-group contact?",
         fontsize=9, color="#444", ha="left")
fig.text(0.04, 0.918,
         "B₃ = change in the year coefficient per +1 unit increase in contact frequency. Filled dot = p < .10. Open dot = ns. Reference line at B₃ = 0 (no moderation).",
         fontsize=9, color="#444", ha="left")

# Legend
handles = [
    mlines.Line2D([], [], marker="o", color=EST, markerfacecolor=EST,
                  markeredgecolor="white", markersize=10, linestyle="None",
                  label="Estonian (sig.)"),
    mlines.Line2D([], [], marker="o", color=EST, markerfacecolor="white",
                  markeredgecolor=EST, markersize=10, markeredgewidth=1.8,
                  linestyle="None", label="Estonian (ns)"),
    mlines.Line2D([], [], marker="o", color=RUS, markerfacecolor=RUS,
                  markeredgecolor="white", markersize=10, linestyle="None",
                  label="Russian (sig.)"),
    mlines.Line2D([], [], marker="o", color=RUS, markerfacecolor="white",
                  markeredgecolor=RUS, markersize=10, markeredgewidth=1.8,
                  linestyle="None", label="Russian (ns)"),
]
fig.legend(handles=handles, loc="upper right", bbox_to_anchor=(0.97, 0.940),
           frameon=False, fontsize=8.5, ncol=2,
           handlelength=1.4, columnspacing=1.5)

fig.text(0.04, 0.022,
         "Out-group contact: Estonian respondents → contact with Russian-speakers (Q52/K4X2); Russian respondents → contact with Estonian-speakers (Q51/K4X1). Inverted so higher = more frequent contact; mean-centered within group.",
         fontsize=7.5, color="#555", ha="left")
fig.text(0.04, 0.008,
         "ⁱ SD: General Out-group uses different items in 2020 (3 items) vs 2023 (6 items). HC3 robust standard errors. Significance: *** p<.001, ** p<.01, * p<.05, ⁺ p<.10.",
         fontsize=7.5, color="#555", ha="left")

plt.subplots_adjust(left=0.18, right=0.985, top=0.88, bottom=0.075)

out_jpg = ROOT / "viz" / "fig_year_x_contact_moderation.jpg"
plt.savefig(out_jpg, dpi=300, format="jpg", facecolor="white",
            pil_kwargs={"quality": 95})
print(f"Saved chart: {out_jpg}")
